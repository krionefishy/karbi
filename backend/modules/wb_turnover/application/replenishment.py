import asyncio
import io
import re
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from backend.modules.wb_core.application import SellerNotFoundError
from backend.modules.wb_core.infrastructure.postgres import SellerRepository
from backend.modules.wb_turnover.infrastructure.postgres import TurnoverRepository

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MOSCOW = ZoneInfo("Europe/Moscow")

# Восемь округов в том же порядке, в каком их держит у себя менеджер. Список
# зашит намеренно: колонка, которая появляется и исчезает вместе с заказами,
# ломает формулы в таблице, куда отчёт вставляют.
DISTRICTS = (
    "Дальневосточный федеральный округ",
    "Приволжский федеральный округ",
    "Северо-Западный федеральный округ",
    "Северо-Кавказский федеральный округ",
    "Сибирский федеральный округ",
    "Уральский федеральный округ",
    "Центральный федеральный округ",
    "Южный федеральный округ",
)
# Куда попадает всё остальное: заказы из-за границы и строки, у которых WB округ
# не назвал. Без этой колонки «Итого» не сходилось бы с суммой по строке, и
# расхождение списывали бы на ошибку выгрузки.
OTHER_DISTRICT = "Прочее"

_LEAD_COLUMNS = ("Баркод", "Артикул WB", "Артикул продавца", "Наименование")


@dataclass(frozen=True, slots=True)
class ReplenishmentReportFile:
    seller_name: str
    day: date
    content: bytes

    @property
    def filename(self) -> str:
        slug = re.sub(r"[^\w-]+", "-", self.seller_name, flags=re.UNICODE).strip("-") or "seller"
        return f"podsort_{slug}_{self.day.isoformat()}.xlsx"


@dataclass(frozen=True, slots=True)
class _Row:
    """Одна строка отчёта: карточка и три её измерения."""

    barcode: str
    article: str
    vendor_code: str
    name: str
    # None — метрику по этому артикулу не считали. Ноль сказал бы «не заказывают».
    rate: float | None
    regions: dict[str, int]
    stocks: dict[int, int]


@dataclass(frozen=True, slots=True)
class _Data:
    """Всё, что нужно книге, снятое одним чтением базы."""

    seller_name: str
    rows: list[_Row]
    warehouses: list[tuple[int, str]]
    window_days: int
    region_from: date
    region_to: date
    regions_filled_from: date | None
    regions_filled_to: date | None
    fbs_stocks_at: datetime | None
    turnover_date: date | None


class ReplenishmentReportService:
    """Отчёт для подсорта: темп, спрос по округам и остатки по складам FBS.

    Три блока, которые менеджер сводил руками из трёх выгрузок WB. Ничего
    нового у Wildberries для него не спрашивается — всё это автоматизация
    оборачиваемости уже собрала.
    """

    def __init__(self, database_url: str, *, window_days: int, region_history_days: int) -> None:
        self._database_url = database_url
        self._window_days = window_days
        self._region_history_days = region_history_days

    async def build(self, seller_id: uuid.UUID, day: date | None = None) -> ReplenishmentReportFile:
        return await asyncio.to_thread(self._build, seller_id, day or datetime.now(MOSCOW).date())

    def _build(self, seller_id: uuid.UUID, day: date) -> ReplenishmentReportFile:
        # Как и отчёт по отзывам: чтение и сборка книги живут в этом треде, а
        # соединения общего движка привязаны к event loop приложения.
        data = asyncio.run(self._load(seller_id, day))
        return ReplenishmentReportFile(data.seller_name, day, self._render(data))

    async def _load(self, seller_id: uuid.UUID, day: date) -> _Data:
        engine = create_async_engine(self._database_url, poolclass=NullPool)
        try:
            async with async_sessionmaker(engine, expire_on_commit=False)() as session:
                sellers = SellerRepository(session)
                seller = await sellers.get(seller_id)
                if seller is None:
                    raise SellerNotFoundError
                turnover = TurnoverRepository(session)
                tracked = await turnover.tracked(seller_id)
                # Окно заканчивается вчера — база собирается прожитыми сутками,
                # и сегодняшние три часа в ней не лежат.
                region_to = day - timedelta(days=1)
                region_from = region_to - timedelta(days=self._region_history_days - 1)
                articles = [item for item in await sellers.list_articles(seller_id) if item.state == "active"]
                barcodes = await sellers.list_article_barcodes(seller_id)
                turnover_date = await turnover.latest_turnover_date(seller_id)
                rates = {
                    row.article: float(row.avg_daily_orders)
                    for row in (await turnover.turnover_on(seller_id, turnover_date) if turnover_date else [])
                }
                regions = await turnover.region_orders(seller_id, region_from, region_to)
                stocks = await turnover.warehouse_stocks(seller_id)
                warehouses = sorted(await turnover.warehouses(seller_id), key=lambda item: item[1].lower())
                rows = [
                    _Row(
                        barcode=barcodes.get(item.article, ""),
                        article=item.article,
                        vendor_code=item.vendor_code,
                        name=item.name,
                        rate=rates.get(item.article),
                        regions=regions.get(item.article, {}),
                        stocks=stocks.get(item.article, {}),
                    )
                    for item in articles
                ]
                # Сверху то, что уходит быстрее всего: разговор о подсорте
                # начинается с него, а не с алфавита.
                rows.sort(key=lambda row: (-(row.rate or 0.0), row.name, row.article))
                return _Data(
                    seller_name=seller.name,
                    rows=rows,
                    warehouses=warehouses,
                    window_days=self._window_days,
                    region_from=region_from,
                    region_to=region_to,
                    regions_filled_from=tracked.regions_filled_from if tracked else None,
                    regions_filled_to=tracked.regions_filled_to if tracked else None,
                    fbs_stocks_at=tracked.fbs_stocks_at if tracked else None,
                    turnover_date=turnover_date,
                )
        finally:
            await engine.dispose()

    def _render(self, data: _Data) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Подсорт"

        header_fill = PatternFill("solid", fgColor="1F2A44")
        group_fill = PatternFill("solid", fgColor="35486F")
        header_font = Font(bold=True, color="FFFFFF")
        centered = Alignment(horizontal="center", vertical="center", wrap_text=True)

        districts = [*DISTRICTS, OTHER_DISTRICT]
        columns = [*_LEAD_COLUMNS, f"Ср. темп {data.window_days} дн., шт/день", *districts, "Итого"]
        columns += [name or f"Склад {warehouse_id}" for warehouse_id, name in data.warehouses]
        first_region = len(_LEAD_COLUMNS) + 2
        first_stock = first_region + len(districts) + 1

        sheet.append([None] * len(columns))
        sheet.append(columns)
        self._group(sheet, first_region, first_stock - 1, f"Заказы по регионам за {self._region_history_days} дн., шт")
        if data.warehouses:
            self._group(sheet, first_stock, len(columns), "Остатки FBS по складам, шт")
        for column in range(1, len(columns) + 1):
            cell = sheet.cell(row=2, column=column)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = centered
        for cell in sheet[1]:
            cell.fill = group_fill

        for row in data.rows:
            counts = [row.regions.get(district, 0) for district in DISTRICTS]
            counts.append(sum(count for name, count in row.regions.items() if name not in DISTRICTS))
            # Пустой слой остатков — это «не собирали», и нули на его месте
            # прочли бы как пустые склады по всему кабинету.
            stocks = [
                row.stocks.get(warehouse_id, 0) if data.fbs_stocks_at else None for warehouse_id, _ in data.warehouses
            ]
            sheet.append(
                [
                    row.barcode,
                    row.article,
                    row.vendor_code,
                    row.name,
                    round(row.rate, 2) if row.rate is not None else None,
                    *counts,
                    sum(counts),
                    *stocks,
                ]
            )

        for index, width in enumerate((16, 14, 26, 44, 14), start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for index in range(first_region, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 13
        sheet.freeze_panes = sheet.cell(row=3, column=first_region)

        self._notes(workbook, data)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _group(sheet, first: int, last: int, title: str) -> None:
        if last < first:
            return
        sheet.cell(row=1, column=first, value=title)
        sheet.cell(row=1, column=first).font = Font(bold=True, color="FFFFFF")
        sheet.cell(row=1, column=first).alignment = Alignment(horizontal="center")
        if last > first:
            sheet.merge_cells(start_row=1, start_column=first, end_row=1, end_column=last)

    @staticmethod
    def _depth(data: _Data) -> int:
        """Сколько суток окна действительно собрано.

        Отрезок собранных суток пересекается с окном отчёта, а не считается от
        его начала: пока история догружается, окно неполное, и читатель должен
        видеть это числом, а не догадываться по нулям в дальних округах.
        """
        if data.regions_filled_from is None or data.regions_filled_to is None:
            return 0
        first = max(data.regions_filled_from, data.region_from)
        last = min(data.regions_filled_to, data.region_to)
        return max(0, (last - first).days + 1)

    def _notes(self, workbook: Workbook, data: _Data) -> None:
        """Лист «Как читать»: без него нельзя отличить ноль от «не собрали»."""
        sheet = workbook.create_sheet("Как читать")
        depth = self._depth(data)
        notes = [
            ("Кабинет", data.seller_name),
            ("Отчёт собран", datetime.now(MOSCOW).strftime("%d.%m.%Y %H:%M МСК")),
            (
                f"Ср. темп {data.window_days} дн.",
                "Заказы минус отмены за окно, делённые на дни, когда товар был в продаже. "
                f"Расчёт от {data.turnover_date.strftime('%d.%m.%Y') if data.turnover_date else '—'}. "
                "Пусто — метрику по этому артикулу ещё не считали.",
            ),
            (
                "Заказы по регионам",
                f"Период {data.region_from.strftime('%d.%m.%Y')} — {data.region_to.strftime('%d.%m.%Y')}. "
                f"Собрано суток: {depth} из {self._region_history_days}. "
                "База набирается по суткам назад, поэтому первое время после подключения окно неполное. "
                "Сегодняшний день в него не входит: он ещё не прожит.",
            ),
            (
                "Остатки FBS",
                (
                    "Объявленный остаток на складах продавца на "
                    f"{data.fbs_stocks_at.astimezone(MOSCOW).strftime('%d.%m.%Y %H:%M МСК')}. "
                    "Это цифра, которую задекларировал селлер, а не пересчёт полки."
                    if data.fbs_stocks_at
                    else "Ни разу не собраны — колонки пустые, а не нулевые."
                ),
            ),
            (
                "Чего здесь нет",
                "Остатков на складах WB (FBO), остатков 1С и товара в пути: подсорт про то, "
                "что лежит на наших складах и как быстро оно уходит.",
            ),
        ]
        for title, value in notes:
            sheet.append([title, value])
        for row in sheet.iter_rows(min_col=1, max_col=1):
            row[0].font = Font(bold=True)
        for row in sheet.iter_rows(min_col=2, max_col=2):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 96
