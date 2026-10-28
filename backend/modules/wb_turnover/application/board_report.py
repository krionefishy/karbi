import io
import re
from dataclasses import dataclass
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.modules.wb_turnover.application.board import FBS_NO, FBS_YES, BoardView

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

COLUMNS = (
    ("Комментарии", 34),
    ("Кабинет", 20),
    ("Артикул WB", 14),
    ("Баркод", 16),
    ("Артикул продавца", 30),
    ("Номенклатура 1С", 40),
    ("Остатки ФБО (Склад ВБ РФ)", 16),
    ("Продажи за 7 дней", 14),
    ("Средние продажи в день", 14),
    ("Оборачиваемость", 16),
    ("Остатки ФБС", 12),
)

FBS_LABELS = {FBS_YES: "Есть", FBS_NO: "Нет"}


@dataclass(frozen=True, slots=True)
class BoardReportFile:
    day: date
    content: bytes

    @property
    def filename(self) -> str:
        return f"oborachivaemost_{self.day.isoformat()}.xlsx"


def render_board(view: BoardView) -> BoardReportFile:
    """Книга в тех же столбцах, что и таблица аналитика.

    Среднее и оборачиваемость — формулы, подсветка — условный формат: если она
    поправит число руками, Excel пересчитает и перекрасит сам.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Оборачиваемость"
    columns = [(re.sub(r"за \d+ дней", f"за {view.window_days} дней", title), width) for title, width in COLUMNS]
    sheet.append([title for title, _ in columns])
    header_fill = PatternFill("solid", fgColor="1F2A44")
    for index, (_, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 32

    for number, row in enumerate(view.rows, start=2):
        sheet.append(
            [
                row.note,
                row.seller_name,
                int(row.article) if row.article.isdigit() else row.article,
                row.barcode,
                row.vendor_code,
                row.nomenclature,
                row.stock,
                row.orders,
                f"=H{number}/{view.window_days}",
                f'=IF(H{number}=0,"нет продаж",IF(G{number}="","нет данных",ROUNDDOWN(G{number}/I{number},0)))',
                FBS_LABELS.get(row.fbs),
            ]
        )
        sheet.cell(row=number, column=4).number_format = "@"
        sheet.cell(row=number, column=9).number_format = "0.0"

    last = max(len(view.rows) + 1, 2)
    scope = f"J2:J{last}"
    sheet.conditional_formatting.add(
        scope,
        FormulaRule(
            formula=[f"AND(ISNUMBER($J2),$J2<{view.red_days})"],
            fill=PatternFill("solid", bgColor="FFC7CE"),
            font=Font(color="9C0006"),
        ),
    )
    sheet.conditional_formatting.add(
        scope,
        FormulaRule(
            formula=[f"AND(ISNUMBER($J2),$J2>={view.red_days},$J2<{view.yellow_days})"],
            fill=PatternFill("solid", bgColor="FFEB9C"),
            font=Font(color="9C5700"),
        ),
    )
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return BoardReportFile(day=view.day, content=buffer.getvalue())
