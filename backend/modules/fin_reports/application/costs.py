"""Файл себестоимости: выгрузка, которую аналитик ведёт в Excel по кабинетам.

Формат один для обоих маркетплейсов — строка шапки с колонками «Кабинет»,
«Артикул пост.» и «Себес, руб»; артикул маркетплейса лежит в «АртикулВБ» у
Wildberries и в «SKU» у Ozon, по этой колонке файл и опознаётся. Остальные
колонки («Проч.затр» и другие) не читаются: в отчёт идёт только себестоимость.
"""

import io
import re
import uuid
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from backend.modules.fin_reports.domain import MARKETPLACE_OZON, MARKETPLACE_WB

CABINET_COLUMN = "кабинет"
COST_COLUMN = "себес, руб"
VENDOR_COLUMN = "артикул пост."
ARTICLE_COLUMNS = {"артикулвб": MARKETPLACE_WB, "sku": MARKETPLACE_OZON}
# Шапка — вторая строка выгрузки, над ней подписи групп колонок; ищем с запасом.
HEADER_SEARCH_ROWS = 10
MAX_PROBLEMS = 50


class CostFileError(Exception):
    """Файл не читается или это не файл себестоимости; текст — для человека."""


@dataclass(frozen=True, slots=True)
class CostRow:
    cabinet: str
    article: str
    vendor_code: str
    cost: Decimal


@dataclass(frozen=True, slots=True)
class CostFile:
    marketplace: str
    rows: list[CostRow]
    problems: list[str] = field(default_factory=list)


def _header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _article(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value or "").strip()


def _cost(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value is None:
        return None
    text = str(value).replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        cost = Decimal(text).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None
    return cost if cost >= 0 else None


def read_cost_file(content: bytes) -> CostFile:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise CostFileError("Файл не открывается как Excel (.xlsx)") from error
    try:
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()
    for index, raw in enumerate(rows[:HEADER_SEARCH_ROWS]):
        names = {_header(value): column for column, value in enumerate(raw) if value is not None}
        marketplaces = [(names[name], marketplace) for name, marketplace in ARTICLE_COLUMNS.items() if name in names]
        if CABINET_COLUMN in names and COST_COLUMN in names and marketplaces:
            article_column, marketplace = marketplaces[0]
            return _read_rows(
                rows[index + 1 :],
                first_row=index + 2,
                marketplace=marketplace,
                cabinet=names[CABINET_COLUMN],
                article=article_column,
                vendor=names.get(VENDOR_COLUMN),
                cost=names[COST_COLUMN],
            )
    raise CostFileError(
        "В файле нет шапки себестоимости: нужны колонки «Кабинет», «Себес, руб» и «АртикулВБ» либо «SKU»"
    )


def _read_rows(
    rows: Iterable[tuple[Any, ...]],
    *,
    first_row: int,
    marketplace: str,
    cabinet: int,
    article: int,
    vendor: int | None,
    cost: int,
) -> CostFile:
    parsed: dict[tuple[str, str], CostRow] = {}
    problems: list[str] = []

    def cell(raw: tuple[Any, ...], column: int | None) -> Any:
        return raw[column] if column is not None and column < len(raw) else None

    for number, raw in enumerate(rows, start=first_row):
        name = str(cell(raw, cabinet) or "").strip()
        code = _article(cell(raw, article))
        if not name and not code:
            continue
        value = _cost(cell(raw, cost))
        if not name or not code:
            problems.append(f"Строка {number}: нет кабинета или артикула")
        elif value is None:
            problems.append(f"Строка {number}: себестоимость «{cell(raw, cost)}» — не число")
        else:
            # Артикул, повторённый в файле, — последняя строка: так его увидел бы и человек.
            parsed[(name.lower(), code)] = CostRow(name, code, str(cell(raw, vendor) or "").strip(), value)
    return CostFile(marketplace, list(parsed.values()), problems[:MAX_PROBLEMS])


def _cabinet_key(name: str) -> tuple[str, ...]:
    """Значимые слова названия: «Ип Саниев», «ИП Саниев Д.Т.» и «Саниев» — один кабинет.

    Отбрасываются только «ИП» и инициалы. Остальное сверяется целиком: по
    одному первому слову «ООО Альфа» и «ООО Бета» оказались бы одним кабинетом,
    и себестоимость одного записалась бы другому.
    """
    words = re.split(r"[^\w]+", name.lower().replace("ё", "е"))
    return tuple(word for word in words if len(word) > 1 and word != "ип")


def match_cabinets(cabinets: Iterable[str], sellers: Mapping[uuid.UUID, str]) -> dict[str, uuid.UUID]:
    """Кабинет файла -> селлер реестра по названию. Нет пары или она не одна — кабинета в ответе нет."""
    by_key: dict[tuple[str, ...], list[uuid.UUID]] = {}
    for seller_id, name in sellers.items():
        by_key.setdefault(_cabinet_key(name), []).append(seller_id)
    matched: dict[str, uuid.UUID] = {}
    for cabinet in cabinets:
        key = _cabinet_key(cabinet)
        found = by_key.get(key, []) if key else []
        if len(found) == 1:
            matched[cabinet] = found[0]
    return matched
