import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.modules.fin_reports.application.view import ArticlesView, Figures, PeriodColumn, PnlView, SkusView
from backend.modules.fin_reports.domain import (
    ARTICLE_COLUMNS,
    LINES,
    MARKETPLACE_WB,
    NO_SALES,
    OZON_LINES,
    OZON_SKU_COLUMNS,
    UNVERIFIED_COLUMNS,
    Line,
    OzonStatement,
    Statement,
)
from backend.modules.fin_reports.domain.pnl import LEVEL_INFO, LEVEL_SUBTOTAL, LEVEL_TOTAL

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOTAL_FILL = PatternFill("solid", fgColor="7030A0")
TOTAL_FONT = Font(bold=True, color="FFFFFF")
SUBTOTAL_FILL = PatternFill("solid", fgColor="E4D4F4")
BOLD = Font(bold=True)
MUTED = Font(italic=True, color="7F7F7F")
MONEY_FORMAT = "#,##0.00;-#,##0.00;-"
TITLE_WIDTH = 46
MONEY_WIDTH = 18
# Excel не принимает в имени листа больше 31 знака и эти символы.
SHEET_FORBIDDEN = str.maketrans({char: " " for char in "[]:*?/\\"})


@dataclass(frozen=True, slots=True)
class FinReportFile:
    content: bytes
    filename: str
    ascii_filename: str


def _sheet_title(text: str) -> str:
    return text.translate(SHEET_FORBIDDEN).strip()[:31] or "Лист"


def _statement_sheet(sheet: Worksheet, columns: list[tuple[str, Figures]], lines: tuple[Line, ...]) -> None:
    """Строки ОПиУ сверху вниз, по колонке на каждый столбец отчёта."""
    sheet.append(["Статья", *(title for title, _ in columns)])
    for cell in sheet[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGNMENT
    sheet.row_dimensions[1].height = 32
    for line in lines:
        sheet.append([line.title, *(float(figures.values[line.key]) for _, figures in columns)])
        row = sheet[sheet.max_row]
        for cell in row[1:]:
            cell.number_format = MONEY_FORMAT
        if line.level == LEVEL_TOTAL:
            for cell in row:
                cell.fill, cell.font = TOTAL_FILL, TOTAL_FONT
        elif line.level == LEVEL_SUBTOTAL:
            for cell in row:
                cell.fill, cell.font = SUBTOTAL_FILL, BOLD
        elif line.level == LEVEL_INFO:
            for cell in row:
                cell.font = MUTED
    sheet.column_dimensions["A"].width = TITLE_WIDTH
    for index in range(2, len(columns) + 2):
        sheet.column_dimensions[get_column_letter(index)].width = MONEY_WIDTH
    sheet.freeze_panes = "B2"


def _dynamics(view: PnlView) -> list[tuple[str, Figures]]:
    return [(f"{view.year} ИТОГО", view.total), *((column.label, column.total) for column in view.periods)]


def build_workbook(
    weeks: PnlView,
    months: PnlView,
    chosen: PeriodColumn | None,
    articles: ArticlesView | None,
    today: date,
    *,
    skus: SkusView | None = None,
) -> FinReportFile:
    workbook = Workbook()
    wb = weeks.marketplace == MARKETPLACE_WB
    lines = LINES if wb else OZON_LINES
    tag = "WB" if wb else "Ozon"
    source = "Отчёт WB ещё не дочитан" if wb else "Начисления Ozon ещё не дочитаны"
    names = {seller.seller_id: seller.name for seller in weeks.sellers}
    first = workbook.active
    assert first is not None
    if chosen is not None:
        first.title = _sheet_title(f"ОПиУ {tag} {chosen.label}")
        blank: Figures = Statement() if wb else OzonStatement()
        by_seller = [(names[seller_id], chosen.by_seller.get(seller_id, blank)) for seller_id in names]
        _statement_sheet(first, [(f"{tag} итого", chosen.total), *by_seller], lines)
        if chosen.pending:
            waiting = ", ".join(names[seller_id] for seller_id in names if seller_id in chosen.pending)
            first.append([])
            first.append([f"{source}: {waiting}. Цифры по этим кабинетам неполные."])
            first[first.max_row][0].font = MUTED
    else:
        first.title = f"ОПиУ {tag}"
        first.append(["Отчётов реализации за год пока нет" if wb else "Начислений Ozon за год пока нет"])
    _statement_sheet(workbook.create_sheet(_sheet_title(f"Недели {tag} {weeks.year}")), _dynamics(weeks), lines)
    _statement_sheet(workbook.create_sheet(_sheet_title(f"Месяцы {tag} {months.year}")), _dynamics(months), lines)
    if articles is not None:
        _articles_sheet(workbook.create_sheet(_sheet_title(f"По артикулам {articles.period.number} нед.")), articles)
        _stocks_sheet(workbook.create_sheet(_sheet_title(f"Остатки {articles.period.number} нед.")), articles)
    if skus is not None:
        _skus_sheet(workbook.create_sheet(_sheet_title(f"ЮНИТ Ozon {skus.period.number} нед.")), skus)
    uncosted = workbook.create_sheet("Без себестоимости")
    article_title = "Артикул ВБ" if wb else "SKU"
    revenue_title = f"Реализация до СПП за {weeks.year}" if wb else f"Реализация с баллами за {weeks.year}"
    uncosted.append(["Кабинет", article_title, "Артикул продавца", revenue_title])
    for cell in uncosted[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGNMENT
    for item in weeks.uncosted:
        uncosted.append([item.seller_name, str(item.nm_id), item.vendor_code, float(item.revenue)])
        uncosted[uncosted.max_row][3].number_format = MONEY_FORMAT
    for letter, width in zip("ABCD", (24, 14, 44, 22), strict=True):
        uncosted.column_dimensions[letter].width = width
    uncosted.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    stamp = chosen.period.key if chosen is not None else str(weeks.year)
    return FinReportFile(
        content=buffer.getvalue(),
        filename=f"ОПиУ_{tag}_{stamp}_{today:%d.%m.%Y}.xlsx",
        ascii_filename=f"pnl_{tag.lower()}_{stamp}_{today:%Y%m%d}.xlsx",
    )


PERCENT_FORMAT = "0.0%"
COUNT_FORMAT = "#,##0"
STOCK_COLUMNS = (
    "Артикул",
    "Артикул продавца",
    "Размер",
    "Остатки на складах ВБ",
    "Себестоимость остатка",
    "Цена реализации остатка",
    "Средняя сумма продаж в день",
    "Оборачиваемость остатка по продажам",
    "Категория",
)


def _is_percent(title: str) -> bool:
    return title.startswith("%") or title.startswith("Доля")


def _cell_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    return value


def _articles_sheet(sheet: Worksheet, view: ArticlesView) -> None:
    """Лист «По артикулам» в колонках старого отчёта; кабинет — первая колонка."""
    titles = ["ИП", *(title for title, _ in ARTICLE_COLUMNS)]
    sheet.append(titles)
    for cell in sheet[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGNMENT
    sheet.row_dimensions[1].height = 60
    for seller in view.sellers:
        for row in seller.rows:
            sheet.append([seller.name, *(_cell_value(take(row)) for _, take in ARTICLE_COLUMNS)])
            cells = sheet[sheet.max_row]
            for index, (title, _) in enumerate(ARTICLE_COLUMNS, start=1):
                value = cells[index].value
                if isinstance(value, float):
                    cells[index].number_format = PERCENT_FORMAT if _is_percent(title) else MONEY_FORMAT
                elif isinstance(value, int):
                    cells[index].number_format = COUNT_FORMAT
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 13
    sheet.column_dimensions["C"].width = 40
    for index in range(4, len(titles) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 14
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{max(sheet.max_row, 1)}"
    notes = []
    if view.stock_is_live:
        stamp = view.stock_taken_at.strftime("%d.%m.%Y %H:%M") if view.stock_taken_at else "—"
        notes.append(f"Остатки — из зеркала на момент выгрузки ({stamp}), снимок на конец недели ещё не снят.")
    pending = [seller.name for seller in view.sellers if seller.pending]
    if pending:
        notes.append(f"Отчёт WB ещё не дочитан: {', '.join(pending)}. Строки по этим кабинетам неполные.")
    untaxed = [seller.name for seller in view.sellers if seller.tax_rate is None]
    if untaxed:
        notes.append(f"Ставка налога не задана: {', '.join(untaxed)} — колонка «Налог» нулевая.")
    behind = [
        f"{seller.name} (по {seller.ads_through:%d.%m})" if seller.ads_through else seller.name
        for seller in view.sellers
        if seller.ads_through is None or seller.ads_through < view.date_to
    ]
    if behind:
        notes.append(f"Реклама дочитана не до конца недели: {', '.join(behind)}.")
    notes.append(
        "Реклама по артикулам: списания кампаний поделены по статистике артикулов за неделю; "
        "кампании, которых нет в зеркале, — в строке «(пусто)»."
    )
    notes.append("Колонки без источника в WB, оставлены для раскладки: " + ", ".join(UNVERIFIED_COLUMNS) + ".")
    for note in notes:
        sheet.append([])
        sheet.append([note])
        sheet[sheet.max_row][0].font = MUTED


def _stocks_sheet(sheet: Worksheet, view: ArticlesView) -> None:
    sheet.append(["ИП", *STOCK_COLUMNS])
    for cell in sheet[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGNMENT
    sheet.row_dimensions[1].height = 32
    for seller in view.sellers:
        for row in seller.stocks:
            days = row.turnover_days
            sheet.append(
                [
                    seller.name,
                    str(row.nm_id),
                    row.vendor_code,
                    row.tech_size,
                    row.quantity,
                    float(row.stock_cost),
                    float(row.stock_price),
                    float(row.daily_sales),
                    float(days) if days is not None else NO_SALES,
                    row.category,
                ]
            )
            cells = sheet[sheet.max_row]
            cells[4].number_format = COUNT_FORMAT
            for index in (5, 6, 7, 8):
                if isinstance(cells[index].value, float):
                    cells[index].number_format = MONEY_FORMAT
    for letter, width in zip("ABCDEFGHIJ", (16, 13, 40, 8, 12, 16, 18, 18, 18, 12), strict=True):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "D2"


def _skus_sheet(sheet: Worksheet, view: SkusView) -> None:
    """Лист «ЮНИТ Ozon»: деньги SKU за неделю из начислений, кабинет — первая колонка."""
    titles = ["Кабинет", *(title for title, _ in OZON_SKU_COLUMNS)]
    sheet.append(titles)
    for cell in sheet[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGNMENT
    sheet.row_dimensions[1].height = 48
    for seller in view.sellers:
        for row in seller.rows:
            sheet.append([seller.name, *(_cell_value(take(row)) for _, take in OZON_SKU_COLUMNS)])
            cells = sheet[sheet.max_row]
            for index in range(1, len(titles)):
                value = cells[index].value
                if isinstance(value, float):
                    cells[index].number_format = MONEY_FORMAT
                elif isinstance(value, int):
                    cells[index].number_format = COUNT_FORMAT
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 14
    for index in range(3, len(titles) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 15
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{max(sheet.max_row, 1)}"
    notes = []
    pending = [
        f"{seller.name} (по {seller.collected_through:%d.%m})" if seller.collected_through else seller.name
        for seller in view.sellers
        if seller.pending
    ]
    if pending:
        notes.append(f"Начисления Ozon дочитаны не до конца недели: {', '.join(pending)}. Строки неполные.")
    notes.append(
        "Начисления без SKU — реклама, подписки, услуги кабинета — в строке «(пусто)»; "
        "«К перечислению» — продажи с баллами минус всё, что Ozon удержал по SKU."
    )
    for note in notes:
        sheet.append([])
        sheet.append([note])
        sheet[sheet.max_row][0].font = MUTED
