import io
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.modules.wb_fbs_penalties.application.view import GroupTotal, PenaltyRowView
from backend.modules.wb_fbs_penalties.domain import TRACE_FOUND, TRACE_NO_ORDER, TRACE_NO_SUPPLY

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MOSCOW = ZoneInfo("Europe/Moscow")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOTAL_FONT = Font(bold=True)
MUTED_FONT = Font(color="808080")

# Её колонки — первыми и в её порядке; дальше то, что она заполняла руками.
COLUMNS: tuple[tuple[str, int], ...] = (
    ("Кабинет", 16),
    ("Неделя отчёта", 22),
    ("Баркод", 16),
    ("Артикул WB", 13),
    ("Название", 34),
    ("Вид удержания", 34),
    ("Группа", 20),
    ("Сумма, ₽", 12),
    ("Стикер МП", 15),
    ("Номер заказа (srid)", 40),
    ("Сборочное задание", 16),
    ("Дата заказа", 17),
    ("Склад продавца (отгрузил)", 26),
    ("Поставка", 18),
    ("Поставка создана", 17),
    ("QR отсканирован", 17),
    ("Склад WB назначения", 26),
)
TEXT_COLUMNS = {3, 9, 10, 11}


@dataclass(frozen=True, slots=True)
class PenaltiesReportFile:
    seller_name: str
    date_from: date
    date_to: date
    content: bytes
    group: str | None = None

    @property
    def filename(self) -> str:
        """Имя с названием кабинета — как она увидит файл в загрузках."""
        name = re.sub(r"[^\w-]+", "_", self.seller_name, flags=re.UNICODE).strip("_") or "seller"
        return self._filename(name)

    @property
    def ascii_filename(self) -> str:
        """То же для `filename=` в заголовке: HTTP-заголовок не переносит кириллицу, а
        кабинеты называются по-русски — иначе выгрузка падает пятисоткой."""
        name = re.sub(r"[^A-Za-z0-9-]+", "_", self.seller_name).strip("_") or "seller"
        return self._filename(name)

    def _filename(self, name: str) -> str:
        group = f"_{self.group}" if self.group else ""
        return f"fbs_penalties_{name}{group}_{self.date_from.isoformat()}_{self.date_to.isoformat()}.xlsx"


class WorkbookWriter:
    """Книга на кабинет, которая пишется потоком: итоги сверху, дальше строки порциями.

    openpyxl в режиме write-only не держит лист в памяти — строка уходит на диск
    сразу, и книга на сто тысяч строк собирается за секунды.
    """

    def __init__(self, seller_name: str, date_from: date, date_to: date, totals: Sequence[GroupTotal]) -> None:
        self.seller_name = seller_name
        self.workbook = Workbook(write_only=True)
        self.sheet = self.workbook.create_sheet("Штрафы")
        for column, (_, width) in enumerate(COLUMNS, start=1):
            self.sheet.column_dimensions[get_column_letter(column)].width = width
        title = self._cell(f"{seller_name}: удержания {date_from:%d.%m.%Y} — {date_to:%d.%m.%Y}")
        title.font = TOTAL_FONT
        self.sheet.append([title])
        for total in totals:
            amount = self._cell(total.amount)
            amount.number_format = "#,##0.00"
            self.sheet.append([total.title, total.count, amount])
        self.sheet.append([])
        header = []
        for name, _ in COLUMNS:
            cell = self._cell(name)
            cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGNMENT
            header.append(cell)
        self.header_row = len(totals) + 3
        self.sheet.append(header)
        self.sheet.freeze_panes = f"D{self.header_row + 1}"
        self.rows = 0

    def add(self, items: Sequence[PenaltyRowView]) -> None:
        for item in items:
            row = item.row
            values: list[object] = [
                self.seller_name,
                f"{row.date_from:%d.%m.%Y} — {row.date_to:%d.%m.%Y}",
                row.barcode,
                row.nm_id or None,
                row.sa_name,
                row.bonus_type_name,
                item.group_title,
                row.amount,
                str(row.sticker_id) if row.sticker_id else "",
                row.srid,
                str(row.assembly_id) if row.assembly_id else "",
                _stamp(item.order_created_at or row.order_dt),
                item.warehouse_name or ("задание не найдено" if item.trace == TRACE_NO_ORDER else ""),
                item.supply_id or ("без поставки" if item.trace == TRACE_NO_SUPPLY else ""),
                _stamp(item.supply_created_at),
                _stamp(item.supply_scan_dt),
                item.destination_office_name or "",
            ]
            cells = []
            for column, value in enumerate(values, start=1):
                cell = self._cell(value)
                if column in TEXT_COLUMNS:
                    # Стикер и номера — текстом: иначе Excel покажет 5,74E+10.
                    cell.number_format = "@"
                if column == 8:
                    cell.number_format = "#,##0.00"
                if column in (13, 14) and item.trace != TRACE_FOUND:
                    cell.font = MUTED_FONT
                cells.append(cell)
            self.sheet.append(cells)
            self.rows += 1

    def finish(self) -> bytes:
        last = self.header_row + self.rows
        self.sheet.auto_filter.ref = f"A{self.header_row}:{get_column_letter(len(COLUMNS))}{last}"
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        return buffer.getvalue()

    def _cell(self, value: object) -> WriteOnlyCell:
        return WriteOnlyCell(self.sheet, value=value)


def _stamp(moment: datetime | None) -> str:
    return moment.astimezone(MOSCOW).strftime("%d.%m.%Y %H:%M") if moment else ""
