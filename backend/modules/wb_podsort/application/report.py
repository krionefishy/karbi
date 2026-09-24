import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.modules.wb_podsort.domain import REGIONS, PodsortRow, PodsortSettings
from backend.shared.exports import export_stem

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MOSCOW = ZoneInfo("Europe/Moscow")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
REGION_FILL = PatternFill("solid", fgColor="DDEBF7")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
MANUAL_FILL = PatternFill("solid", fgColor="F2F2F2")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
MONTHS = (
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)

# Столбцы сводного листа — как в «Подсорт 23.09.2026» образца; ручные остаются пустыми.
SUMMARY_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Кабинет", 20),
    ("Артикул ВБ", 13),
    ("Баркод", 16),
    ("Короба кол-во", 10),
    ("Артикул товара", 34),
    ("Расчёт системы", 11),
    ("Запрос логиста, шт", 11),
    ("Проверка Ильвира, шт", 11),
    ("Сборка склад, шт", 11),
    ("Примечание", 24),
    ("Заказы за окно, шт", 11),
    ("Ср. в день", 9),
    ("Остаток WB в регионе", 11),
    ("Дней покрытия", 10),
)
MANUAL_COLUMNS = {4, 7, 8, 9, 10}
CALC_COLUMN = 6

CALC_COMMENT = (
    "Расчёт системы = продажи за окно / дней в окне × дней покрытия − остаток на складах WB этого региона. "
    "Меньше нуля — 0. Остаток нашего склада не учитывается: наличие проверяет Ильвир."
)


@dataclass(frozen=True, slots=True)
class SellerSection:
    """Кабинет в книге: его строки и оговорки, которые надо сказать рядом с цифрами."""

    name: str
    rows: Sequence[PodsortRow]
    window_days_loaded: int
    remains_at: datetime | None


@dataclass(frozen=True, slots=True)
class PodsortReportFile:
    day: date
    content: bytes

    @property
    def filename(self) -> str:
        return f"podsort_{self.day.isoformat()}.xlsx"

    @property
    def ascii_filename(self) -> str:
        return self.filename


@dataclass(frozen=True, slots=True)
class ReportContext:
    today: date
    last_day: date
    window_start: date
    months: tuple[date, ...]
    settings: PodsortSettings
    unplaced_warehouses: tuple[str, ...]


def month_title(month: date) -> str:
    return f"{MONTHS[month.month - 1]} {month.year}"


def _stamp(moment: datetime | None) -> str:
    return moment.astimezone(MOSCOW).strftime("%d.%m.%Y %H:%M") if moment else "не собирались"


def _cover(value: float | None) -> float | None:
    return round(value, 1) if value is not None else None


class PodsortWorkbook:
    """Книга подсорта на все кабинеты: сводный лист по регионам, лист на кабинет, «Как читать»."""

    def __init__(self, context: ReportContext, sections: Sequence[SellerSection]) -> None:
        self.context = context
        self.sections = sections
        self.workbook = Workbook()

    def build(self) -> bytes:
        summary = self.workbook.active
        assert summary is not None
        summary.title = f"Подсорт {self.context.today:%d.%m.%Y}"
        self._summary(summary)
        taken = {summary.title}
        for section in self.sections:
            title = self._sheet_title(section.name, taken)
            taken.add(title)
            self._seller(self.workbook.create_sheet(title), section)
        self._how_to_read(self.workbook.create_sheet("Как читать"))
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        return buffer.getvalue()

    # --- сводный лист ---------------------------------------------------------

    def notes(self) -> list[str]:
        """Пометки над сводным листом: как посчитано и где цифрам нельзя верить."""
        context, settings = self.context, self.context.settings
        notes = [
            f"Подсорт на {context.today:%d.%m.%Y}. Продажи за {settings.window_days} дн. "
            f"({context.window_start:%d.%m}–{context.last_day:%d.%m}), везём на {settings.cover_days} дн.",
            f"Расчёт системы = продажи за {settings.window_days} дн. / {settings.window_days} × "
            f"{settings.cover_days} − остаток на складах WB в регионе. Меньше нуля — 0.",
            "Остаток WB в регионе вычтен: товар, который уже лежит на складах региона, система не везёт. "
            "Склады Центрального региона (Коледино, Тула, Электросталь) отгружают и в соседние регионы — "
            "их остаток частично работает на них, поэтому смотрите «Дней покрытия».",
            "Остаток нашего склада не учитывается — наличие проверяет Ильвир в столбце «Проверка Ильвира». "
            "Товар в пути и «Склад WB РФ» (WB не говорит, где он лежит) не вычитаются.",
        ]
        if context.unplaced_warehouses:
            notes.append(
                "Склады без региона, их остаток не вычтен: "
                + ", ".join(context.unplaced_warehouses)
                + ". Регион задаётся на странице подсорта."
            )
        for section in self.sections:
            if section.window_days_loaded < settings.window_days:
                notes.append(
                    f"{section.name}: заказы загружены за {section.window_days_loaded} из {settings.window_days} дн. "
                    "окна — расчёт занижен, догрузка идёт."
                )
            if section.remains_at is None:
                notes.append(f"{section.name}: остатки по складам WB ещё не собирались — остаток принят за 0.")
        return notes

    def _summary(self, sheet: Worksheet) -> None:
        width = len(SUMMARY_COLUMNS)
        for column, (_, size) in enumerate(SUMMARY_COLUMNS, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = size
        line = 1
        for note in self.notes():
            cell = sheet.cell(line, 1, note)
            cell.fill, cell.alignment = NOTE_FILL, WRAP
            sheet.merge_cells(start_row=line, start_column=1, end_row=line, end_column=width)
            sheet.row_dimensions[line].height = 30
            line += 1
        line += 1
        header_row = 0
        for region in self.context.settings.regions:
            title = sheet.cell(line, 1, f"Подсорт на {region}")
            title.font, title.fill = BOLD, REGION_FILL
            sheet.merge_cells(start_row=line, start_column=1, end_row=line, end_column=width)
            line += 1
            for column, (name, _) in enumerate(SUMMARY_COLUMNS, start=1):
                cell = sheet.cell(line, column, name)
                cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGNMENT
            sheet.cell(line, CALC_COLUMN).comment = Comment(CALC_COMMENT, "Подсорт")
            header_row = header_row or line
            line += 1
            count = 0
            for section in self.sections:
                for row in section.rows:
                    figures = row.targets.get(region)
                    if figures is None:
                        continue
                    need = figures.need(self.context.settings.window_days, self.context.settings.cover_days)
                    if need <= 0:
                        continue
                    values: list[object] = [
                        section.name,
                        row.info.nm_id or None,
                        row.info.barcode,
                        None,
                        row.info.vendor_code,
                        need,
                        None,
                        None,
                        None,
                        None,
                        figures.window_orders,
                        round(figures.average(self.context.settings.window_days), 1),
                        figures.stock,
                        _cover(figures.cover(self.context.settings.window_days)),
                    ]
                    for column, value in enumerate(values, start=1):
                        cell = sheet.cell(line, column, value)
                        if column == 3:
                            cell.number_format = "@"
                        if column in MANUAL_COLUMNS:
                            cell.fill = MANUAL_FILL
                        if column == CALC_COLUMN:
                            cell.font = BOLD
                    line += 1
                    count += 1
            if count == 0:
                sheet.cell(line, 1, "Везти нечего: остатка WB в регионе хватает на срок покрытия.")
                line += 1
            line += 1
        if header_row:
            sheet.freeze_panes = sheet.cell(header_row + 1, 3)

    # --- лист кабинета ---------------------------------------------------------

    def _seller(self, sheet: Worksheet, section: SellerSection) -> None:
        context, settings = self.context, self.context.settings
        current = month_title(context.months[-1])
        columns: list[tuple[str, int]] = [
            ("Номенклатура", 34),
            ("Предмет", 26),
            ("Баркод", 16),
            ("Артикул WB", 13),
            ("Размер", 8),
            *[(month_title(month), 11) for month in context.months],
            ("14 дней", 9),
            ("7 дней", 9),
            (f"Ср. в день ({settings.window_days} дн.)", 10),
            (f"Из заказов за {settings.window_days} дн. со склада продавца", 12),
            *[(f"{region}, {current}", 11) for region in REGIONS],
            ("Склад WB РФ и склады без региона", 12),
        ]
        for region in settings.regions:
            columns.extend(
                [
                    (f"{region}: заказы за {settings.window_days} дн.", 11),
                    (f"{region}: остаток WB", 11),
                    (f"{region}: подсорт", 11),
                ]
            )
        title = sheet.cell(1, 1, f"Подсорт {section.name}")
        title.font = BOLD
        sheet.cell(
            1,
            3,
            f"Заказы по {context.last_day:%d.%m.%Y}; остатки WB на {_stamp(section.remains_at)}. "
            "Месяц — с первого числа по вчера; ранние месяцы видны настолько, насколько WB хранит заказы.",
        )
        for column, (name, size) in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = size
            cell = sheet.cell(2, column, name)
            cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGNMENT
        sheet.row_dimensions[2].height = 45
        for line, row in enumerate(section.rows, start=3):
            info = row.info
            values: list[object] = [
                info.vendor_code,
                info.subject,
                info.barcode,
                info.nm_id or None,
                info.tech_size if info.tech_size not in ("", "0") else "-",
                *row.month_orders,
                row.orders_14,
                row.orders_7,
                round(row.average(settings.window_days), 1),
                row.window_fbs_orders,
                *[row.month_by_region.get(region, 0) for region in REGIONS],
                row.unplaced_stock,
            ]
            for region in settings.regions:
                figures = row.targets[region]
                values.extend(
                    [
                        figures.window_orders,
                        figures.stock,
                        figures.need(settings.window_days, settings.cover_days),
                    ]
                )
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(line, column, value)
                if column == 3:
                    cell.number_format = "@"
        sheet.freeze_panes = "F3"
        if section.rows:
            sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{len(section.rows) + 2}"

    # --- как читать -------------------------------------------------------------

    def _how_to_read(self, sheet: Worksheet) -> None:
        settings = self.context.settings
        sheet.column_dimensions["A"].width = 120
        lines = [
            "Откуда цифры",
            "Заказы — статистика WB по дню заказа, с отменами, как «Итого заказов» в отчёте «География заказов». "
            "Регион — округ покупателя; округа склеены как у WB: Юг с Северным Кавказом, Дальний Восток с Сибирью.",
            "Остатки — отчёт WB «Остатки на складах» по каждому складу. Регион склада — по городу в названии "
            "или как задали на странице подсорта.",
            "",
            "Как считается «Расчёт системы»",
            f"Среднее в день = заказы региона за {settings.window_days} полных дней до вчера / {settings.window_days}.",
            f"Потребность = среднее в день × {settings.cover_days} дн. покрытия.",
            "Расчёт системы = потребность − остаток на складах WB этого региона, округление до целого; "
            "меньше нуля — 0.",
            "«Дней покрытия» — на сколько дней хватит остатка региона при нынешних продажах.",
            "",
            "Что не учитывается",
            "Остаток нашего склада — его в WB нет; наличие проверяет Ильвир.",
            "Товар в пути к покупателю и возвраты в пути, «Склад WB РФ» и склады без региона.",
            "Заказы со склада продавца (FBS) входят в продажи, как в «Итого заказов»: отдельный столбец на листе "
            "кабинета показывает их долю.",
            "",
            "Ручные столбцы сводного листа (короба, запрос логиста, проверка Ильвира, сборка склад, примечание) "
            "система не заполняет.",
        ]
        headings = {"Откуда цифры", "Как считается «Расчёт системы»", "Что не учитывается"}
        for line, text in enumerate(lines, start=1):
            cell = sheet.cell(line, 1, text)
            cell.alignment = WRAP
            if text in headings:
                cell.font = BOLD

    @staticmethod
    def _sheet_title(name: str, taken: set[str]) -> str:
        """Имя листа Excel — не длиннее 31 знака, без запрещённых символов и без повторов."""
        base = f"Подсорт {export_stem(name).replace('_', ' ')}"[:31]
        title, suffix = base, 2
        while title in taken:
            title = f"{base[: 31 - len(str(suffix)) - 1]} {suffix}"
            suffix += 1
        return title
