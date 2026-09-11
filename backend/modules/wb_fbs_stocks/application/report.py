import io
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import Outline

from backend.modules.wb_fbs_stocks.application.view import BoardView

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COMPARISON_SHEET = "Сравнение"

# Первые две колонки как в таблице селлера: её заметка и баркод.
NOTE_COLUMN = 1
BARCODE_COLUMN = 2
FIRST_DATA_COLUMN = 3

GROUP_FILL = PatternFill("solid", fgColor="1F3864")
COLUMN_FILL = PatternFill("solid", fgColor="2E5F8A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTERED = Alignment(horizontal="center", vertical="center")
# Ноль — красным, по ТЗ. Условным форматом, а не заливкой: если она поправит
# число руками, цвет уйдёт сам.
ZERO_FILL = PatternFill("solid", fgColor="FFC7CE")
ZERO_FONT = Font(color="9C0006")
NOTE_FONT = Font(color="404040")


@dataclass(frozen=True, slots=True)
class BoardReportFile:
    day: date
    content: bytes

    @property
    def filename(self) -> str:
        return f"fbs_stocks_{self.day.isoformat()}.xlsx"


def render_workbook(views: Sequence[BoardView]) -> bytes:
    """Книга в формате таблицы селлера: лист на кабинет и общий лист сравнения.

    На листе кабинета группы столбцов сворачиваются тем же «плюсиком», что у неё:
    сводный столбец слева, склады группы — справа под ним. Суммы — формулами,
    чтобы Excel пересчитал их, если она подвинет число руками.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    taken: set[str] = set()
    for view in views:
        _seller_sheet(workbook.create_sheet(_sheet_title(view.seller_name, taken)), view)
    _comparison_sheet(workbook.create_sheet(COMPARISON_SHEET), views)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _seller_sheet(sheet, view: BoardView) -> None:
    sheet.sheet_properties.outlinePr = Outline(summaryRight=False, summaryBelow=False)
    sheet.cell(row=1, column=NOTE_COLUMN, value=view.seller_name)
    sheet.cell(row=1, column=BARCODE_COLUMN, value="Баркод")
    column = FIRST_DATA_COLUMN
    spans: list[tuple[int, int, int]] = []  # (сводный столбец, первый склад, последний склад)
    for group in view.groups:
        sheet.cell(row=1, column=column, value=group.title)
        first = column + 1
        for offset, member in enumerate(group.columns):
            sheet.cell(row=1, column=first + offset, value=member.name)
        spans.append((column, first, first + len(group.columns) - 1))
        column = first + len(group.columns)
    last_column = column - 1

    for index, row in enumerate(view.rows, start=2):
        sheet.cell(row=index, column=NOTE_COLUMN, value=row.note or None).font = NOTE_FONT
        barcode = sheet.cell(row=index, column=BARCODE_COLUMN, value=row.barcode)
        barcode.number_format = "@"
        for group, (summary, first, last) in zip(view.groups, spans, strict=True):
            total = _sum_formula(first, last, index) if group.columns else 0
            sheet.cell(row=index, column=summary, value=total).alignment = CENTERED
            for offset, member in enumerate(group.columns):
                sheet.cell(row=index, column=first + offset, value=row.amount(member.warehouse_id)).alignment = CENTERED

    _style_header(sheet, [summary for summary, _, _ in spans], last_column)
    for summary, first, last in spans:
        sheet.column_dimensions[get_column_letter(summary)].width = 16
        if last >= first:
            sheet.column_dimensions.group(get_column_letter(first), get_column_letter(last), outline_level=1)
    if view.rows:
        _paint_zeroes(sheet, 2, len(view.rows) + 1, FIRST_DATA_COLUMN, last_column)
    sheet.freeze_panes = sheet.cell(row=2, column=FIRST_DATA_COLUMN)


def _comparison_sheet(sheet, views: Sequence[BoardView]) -> None:
    """Все кабинеты друг под другом: свои склады раскрыты, округа — суммами."""
    sheet.sheet_properties.outlinePr = Outline(summaryRight=False, summaryBelow=False)
    own_width = max((len(view.own_group.columns) if view.own_group else 0 for view in views), default=0)
    district_start = FIRST_DATA_COLUMN + 1 + own_width
    row_index = 1
    last_column = FIRST_DATA_COLUMN
    for view in views:
        own = view.own_group
        districts = view.district_groups
        sheet.cell(row=row_index, column=NOTE_COLUMN, value=view.seller_name)
        sheet.cell(row=row_index, column=BARCODE_COLUMN, value="Баркод")
        sheet.cell(row=row_index, column=FIRST_DATA_COLUMN, value=own.title if own else "Наш склад")
        own_columns = own.columns if own else ()
        for offset, member in enumerate(own_columns):
            sheet.cell(row=row_index, column=FIRST_DATA_COLUMN + 1 + offset, value=member.name)
        for offset, group in enumerate(districts):
            sheet.cell(row=row_index, column=district_start + offset, value=group.title)
        last_column = max(last_column, district_start + len(districts) - 1)
        _style_header(
            sheet,
            [FIRST_DATA_COLUMN, *range(district_start, district_start + len(districts))],
            last_column,
            row=row_index,
        )
        first_row = row_index + 1
        for row in view.rows:
            row_index += 1
            sheet.cell(row=row_index, column=NOTE_COLUMN, value=row.note or None).font = NOTE_FONT
            barcode = sheet.cell(row=row_index, column=BARCODE_COLUMN, value=row.barcode)
            barcode.number_format = "@"
            total = (
                _sum_formula(FIRST_DATA_COLUMN + 1, FIRST_DATA_COLUMN + len(own_columns), row_index)
                if own_columns
                else 0
            )
            sheet.cell(row=row_index, column=FIRST_DATA_COLUMN, value=total).alignment = CENTERED
            for offset, member in enumerate(own_columns):
                cell = sheet.cell(
                    row=row_index, column=FIRST_DATA_COLUMN + 1 + offset, value=row.amount(member.warehouse_id)
                )
                cell.alignment = CENTERED
            for offset, group in enumerate(districts):
                cell = sheet.cell(row=row_index, column=district_start + offset, value=row.group_total(group))
                cell.alignment = CENTERED
        if view.rows:
            # Красятся только записанные ячейки: пустая клетка для Excel равна нулю,
            # и разрыв между блоками стал бы красным.
            _paint_zeroes(sheet, first_row, row_index, FIRST_DATA_COLUMN, FIRST_DATA_COLUMN + len(own_columns))
            if districts:
                _paint_zeroes(sheet, first_row, row_index, district_start, district_start + len(districts) - 1)
        # Пустая строка между кабинетами, как в её примере.
        row_index += 2
    if own_width:
        sheet.column_dimensions.group(
            get_column_letter(FIRST_DATA_COLUMN + 1), get_column_letter(FIRST_DATA_COLUMN + own_width), outline_level=1
        )
    sheet.column_dimensions[get_column_letter(FIRST_DATA_COLUMN)].width = 16


def _sum_formula(first: int, last: int, row: int) -> str:
    return f"=SUM({get_column_letter(first)}{row}:{get_column_letter(last)}{row})"


def _style_header(sheet, summary_columns: Sequence[int], last_column: int, *, row: int = 1) -> None:
    summaries = set(summary_columns)
    for column in range(1, last_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = GROUP_FILL if column in summaries or column < FIRST_DATA_COLUMN else COLUMN_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
    sheet.row_dimensions[row].height = 42
    sheet.column_dimensions[get_column_letter(NOTE_COLUMN)].width = 28
    sheet.column_dimensions[get_column_letter(BARCODE_COLUMN)].width = 18
    for column in range(FIRST_DATA_COLUMN, last_column + 1):
        letter = get_column_letter(column)
        if sheet.column_dimensions[letter].width is None or sheet.column_dimensions[letter].width < 14:
            sheet.column_dimensions[letter].width = 14


def _paint_zeroes(sheet, first_row: int, last_row: int, first_column: int, last_column: int) -> None:
    if last_column < first_column or last_row < first_row:
        return
    area = f"{get_column_letter(first_column)}{first_row}:{get_column_letter(last_column)}{last_row}"
    sheet.conditional_formatting.add(area, CellIsRule(operator="equal", formula=["0"], fill=ZERO_FILL, font=ZERO_FONT))


def _sheet_title(name: str, taken: set[str]) -> str:
    """Имя листа Excel: до 31 знака, без запрещённых символов, уникальное в книге."""
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Кабинет"
    title = base
    suffix = 2
    while title in taken or title == COMPARISON_SHEET:
        tail = f" {suffix}"
        title = base[: 31 - len(tail)] + tail
        suffix += 1
    taken.add(title)
    return title
