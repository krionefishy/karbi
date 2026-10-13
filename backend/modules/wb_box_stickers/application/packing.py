"""Раскладка товара из таблицы сборки по коробам шаблона коробовки WB.

Шаблон скачан из кабинета после создания коробов: строка на короб, «ШК короба»
и «ШК короба для печати» уже стоят, «Баркод товара» пустой, «Кол-во товаров» — 0.
Короба идут по порядку строк таблицы сборки и по порядку внутри схемы; первый
короб получает первую строку шаблона, второй — вторую. Так порядок строк файла
совпадает с таблицей сборки, и стикеры потом лягут по порядку без лишней работы.

Записываются только баркод и количество. Коды коробов, лишние колонки и формат
ячеек не трогаются: кабинет на изменённый формат колонки отвечает
«Не найден ШК короба».
"""

import io
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass, field

from openpyxl import load_workbook

from backend.modules.wb_box_stickers.application.workbook import (
    HEADER_BARCODE,
    HEADER_QUANTITY,
    HEADER_SHK,
    _digits,
    _header,
    read_workbook,
)
from backend.modules.wb_box_stickers.domain import Assembly, PackedBox, StickerInputError, WorkbookBox

HEADER_SEARCH_ROWS = 10


@dataclass(frozen=True, slots=True)
class PackingPlan:
    boxes: list[PackedBox]
    problems: list[str] = field(default_factory=list)

    @property
    def ready(self) -> bool:
        return not self.problems and bool(self.boxes)


def plan_packing(template: Sequence[WorkbookBox], assembly: Assembly) -> PackingPlan:
    """Кому какой короб: строки шаблона по порядку против коробов по схемам таблицы."""
    problems = list(assembly.problems)
    filled = [box.shk for box in template if box.lines]
    if filled:
        problems.append(
            f"В шаблоне уже заполнены короба {', '.join(filled[:10])}"
            + (f" и ещё {len(filled) - 10}" if len(filled) > 10 else "")
            + " — нужен пустой шаблон, скачанный из кабинета сразу после создания коробов"
        )
    needed = assembly.boxes
    if needed and len(template) != needed:
        difference = needed - len(template)
        action = f"добавьте в кабинете {difference}" if difference > 0 else f"удалите в кабинете {-difference}"
        problems.append(
            f"В шаблоне {len(template)} коробов, а по таблице сборки нужно {needed}: {action} и скачайте шаблон заново"
        )
    if problems:
        return PackingPlan(boxes=[], problems=problems)

    boxes: list[PackedBox] = []
    slots = iter(template)
    for row in assembly.rows:
        for quantity in row.boxes:
            slot = next(slots)
            boxes.append(
                PackedBox(
                    position=len(boxes) + 1,
                    shk=slot.shk,
                    package_code=slot.package_code,
                    barcode=row.barcode,
                    quantity=quantity,
                )
            )
    return PackingPlan(boxes=boxes)


def fill_template(data: bytes, boxes: Sequence[PackedBox]) -> bytes:
    """Тот же файл шаблона с баркодом и количеством в строке каждого короба."""
    try:
        book = load_workbook(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
        raise StickerInputError("Excel не открывается — нужен файл .xlsx из кабинета WB") from error
    sheet = book.active
    if sheet is None:
        raise StickerInputError("В Excel нет листов")

    columns: dict[str, int] = {}
    header_row = 0
    for row in sheet.iter_rows(min_row=1, max_row=HEADER_SEARCH_ROWS):
        titles = {_header(cell.value): cell.column for cell in row}
        if HEADER_SHK in titles:
            header_row = row[0].row
            columns = {"shk": titles[HEADER_SHK]}
            columns["barcode"] = titles.get(HEADER_BARCODE, 0)
            columns["quantity"] = titles.get(HEADER_QUANTITY, 0)
            break
    if not header_row or not columns["barcode"] or not columns["quantity"]:
        raise StickerInputError("В шаблоне нет колонок «Баркод товара» и «Кол-во товаров» — это не шаблон коробовки WB")

    by_shk = {box.shk: box for box in boxes}
    for row_number in range(header_row + 1, sheet.max_row + 1):
        shk_cell = sheet.cell(row=row_number, column=columns["shk"])
        shk = _digits(shk_cell.value)
        box = by_shk.get(shk)
        if box is None:
            continue
        # Баркод строкой: числом Excel покажет его как 2,05E+12, а кабинет не найдёт товар.
        barcode_cell = sheet.cell(row=row_number, column=columns["barcode"])
        barcode_cell.value = box.barcode
        barcode_cell.data_type = "s"
        sheet.cell(row=row_number, column=columns["quantity"]).value = box.quantity
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def pack(template_data: bytes, assembly: Assembly) -> tuple[PackingPlan, bytes | None]:
    """Разбор шаблона, раскладка и запись; файл только когда проблем нет.

    Готовый файл перечитывается тем же разбором, что и коробовка для стикеров:
    в каждом коробе должен оказаться ровно тот товар и то количество, что в плане.
    """
    template = read_workbook(template_data)
    plan = plan_packing(template, assembly)
    if not plan.ready:
        return plan, None
    data = fill_template(template_data, plan.boxes)
    written = {box.shk: box.lines for box in read_workbook(data)}
    expected = {box.shk: ((box.barcode, box.quantity),) for box in plan.boxes}
    wrong = [
        shk
        for shk, lines in expected.items()
        if tuple((line.barcode, line.quantity) for line in written.get(shk, ())) != lines
    ]
    if wrong or list(written) != [box.shk for box in plan.boxes]:
        return PackingPlan(boxes=plan.boxes, problems=["Записанный файл не прошёл проверку, отдавать его нельзя"]), None
    return plan, data
