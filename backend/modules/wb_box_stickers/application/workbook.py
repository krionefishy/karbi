"""Excel коробовки: какие короба, в каком порядке и что селлер в них положил."""

import io
import re
import zipfile
from typing import Any

from openpyxl import load_workbook

from backend.modules.wb_box_stickers.domain import BoxLine, StickerInputError, WorkbookBox

HEADER_SHK = "шк короба"
HEADER_CODE = "шк короба для печати"
HEADER_BARCODE = "баркод товара"
HEADER_QUANTITY = "кол-во товаров"
# Шапка у WB в первой строке; несколько строк сверху — запас на ручную правку.
HEADER_SEARCH_ROWS = 10


def _header(value: Any) -> str:
    return " ".join(str(value or "").split()).lower()


def _digits(value: Any) -> str:
    """ШК короба и баркод: в Excel это то строка, то число, то число с «.0»."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", "", str(value))
    return text if text.isdigit() else ""


def _quantity(value: Any) -> int:
    if isinstance(value, bool) or value is None:
        return 0
    if isinstance(value, int | float):
        return int(value) if float(value).is_integer() and value > 0 else 0
    text = str(value).strip()
    return int(text) if text.isdigit() else 0


def read_workbook(data: bytes) -> list[WorkbookBox]:
    """Короба по порядку первой строки. Короб в нескольких строках — несколько баркодов."""
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
        raise StickerInputError("Excel не открывается — нужен файл .xlsx из кабинета WB") from error
    try:
        sheet = book.active
        if sheet is None:
            raise StickerInputError("В Excel нет листов")
        # WB пишет в файл размер листа «A1», и read_only без сброса видит одну ячейку.
        sheet.reset_dimensions()  # type: ignore[union-attr]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()

    columns: dict[str, int] = {}
    header_row = -1
    for number, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        titles = [_header(cell) for cell in row]
        if HEADER_SHK in titles:
            header_row = number
            for index, title in enumerate(titles):
                if title == HEADER_SHK:
                    columns.setdefault("shk", index)
                elif title.startswith(HEADER_CODE):
                    columns.setdefault("code", index)
                elif title == HEADER_BARCODE:
                    columns.setdefault("barcode", index)
                elif title == HEADER_QUANTITY:
                    columns.setdefault("quantity", index)
            break
    if header_row < 0:
        raise StickerInputError("В Excel нет колонки «ШК короба» — нужен файл коробовки из кабинета WB")

    def cell(row: tuple[Any, ...], key: str) -> Any:
        index = columns.get(key)
        return row[index] if index is not None and index < len(row) else None

    order: list[str] = []
    codes: dict[str, str] = {}
    lines: dict[str, list[BoxLine]] = {}
    first_row: dict[str, int] = {}
    for number, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        shk = _digits(cell(row, "shk"))
        if not shk:
            continue
        code = str(cell(row, "code") or "").strip()
        if shk not in first_row:
            order.append(shk)
            first_row[shk] = number
            codes[shk] = code
            lines[shk] = []
        elif code and codes[shk] and code != codes[shk]:
            raise StickerInputError(
                f"У короба {shk} в Excel два разных кода для печати (строки {first_row[shk]} и {number})"
            )
        elif code:
            codes[shk] = code
        barcode = _digits(cell(row, "barcode"))
        quantity = _quantity(cell(row, "quantity"))
        if barcode and quantity:
            lines[shk].append(BoxLine(barcode=barcode, quantity=quantity))
    if not order:
        raise StickerInputError("В Excel нет ни одного короба в колонке «ШК короба»")
    return [WorkbookBox(shk=shk, package_code=codes[shk], lines=tuple(lines[shk]), row=first_row[shk]) for shk in order]
