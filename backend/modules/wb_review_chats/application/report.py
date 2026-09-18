import io
from dataclasses import dataclass
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.modules.wb_review_chats.application.view import ReviewChatsView
from backend.modules.wb_review_chats.domain import (
    GROUP_BARE,
    GROUP_FOLLOWED,
    OUTCOME_PENDING,
    OUTCOME_REPLIED,
    OUTCOME_SILENT,
    GroupSummary,
)
from backend.shared.exports import export_stem

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)
WRAP = Alignment(vertical="top", wrap_text=True)

GROUP_TITLES = {GROUP_FOLLOWED: "С нашим сообщением", GROUP_BARE: "Без нашего сообщения"}
OUTCOME_TITLES = {OUTCOME_REPLIED: "Ответил", OUTCOME_SILENT: "Не ответил", OUTCOME_PENDING: "Ждём ответа"}

SUMMARY_COLUMNS = ("Диалогов", "Ответили", "Не ответили", "Ждём ответа", "% ответивших", "% не ответивших")
DIALOG_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Автосообщение WB", 18),
    ("Артикул WB", 13),
    ("Товар", 36),
    ("Группа", 22),
    ("Исход", 14),
    ("Наше сообщение", 18),
    ("Ответ покупателя", 18),
    ("Текст ответа", 60),
    ("Чат", 40),
)


@dataclass(frozen=True, slots=True)
class ReviewChatsReportFile:
    seller_name: str
    date_from: date
    date_to: date
    content: bytes

    @property
    def filename(self) -> str:
        return self._filename(export_stem(self.seller_name))

    @property
    def ascii_filename(self) -> str:
        return self._filename(export_stem(self.seller_name, ascii_only=True))

    def _filename(self, name: str) -> str:
        return f"review_chats_{name}_{self.date_from.isoformat()}_{self.date_to.isoformat()}.xlsx"


def build_workbook(view: ReviewChatsView, timezone: ZoneInfo) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    _summary_sheet(summary, view)
    _dialogs_sheet(workbook.create_sheet("Диалоги"), view, timezone)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _summary_sheet(sheet: Worksheet, view: ReviewChatsView) -> None:
    sheet.append([f"{view.seller_name}: чаты после отзыва {view.date_from:%d.%m.%Y} — {view.date_to:%d.%m.%Y}"])
    sheet["A1"].font = BOLD
    sheet.append([f"Ответ засчитывается в течение {view.reply_window_hours} ч; проценты — без диалогов, где ещё ждём"])
    sheet.append([])
    _header(sheet, ("Группа", *SUMMARY_COLUMNS))
    sheet.append([GROUP_TITLES[GROUP_FOLLOWED], *_summary_cells(view.followed)])
    sheet.append([GROUP_TITLES[GROUP_BARE], *_summary_cells(view.bare)])
    sheet.append([])
    _header(sheet, ("День", *(f"{title}: {column}" for title in GROUP_TITLES.values() for column in SUMMARY_COLUMNS)))
    for day in view.days:
        sheet.append([day.day.strftime("%d.%m.%Y"), *_summary_cells(day.followed), *_summary_cells(day.bare)])
    sheet.column_dimensions["A"].width = 24
    for column in range(2, 2 + 2 * len(SUMMARY_COLUMNS)):
        sheet.column_dimensions[get_column_letter(column)].width = 16
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0%"


def _dialogs_sheet(sheet: Worksheet, view: ReviewChatsView, timezone: ZoneInfo) -> None:
    _header(sheet, tuple(name for name, _ in DIALOG_COLUMNS))
    for column, (_, width) in enumerate(DIALOG_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for item in view.dialogs:
        dialog = item.dialog
        reply = dialog.reply_text or ("вложение" if dialog.reply_has_attachments else "")
        sheet.append(
            [
                _stamp(dialog.prompt_at, timezone),
                dialog.nm_id,
                item.product_name,
                GROUP_TITLES[dialog.group],
                OUTCOME_TITLES[dialog.outcome],
                _stamp(dialog.follow_up_at, timezone) + (" (после ответа)" if dialog.follow_up_late else ""),
                _stamp(dialog.reply_at, timezone),
                reply,
                dialog.chat_id,
            ]
        )
        sheet.cell(row=sheet.max_row, column=8).alignment = WRAP
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(DIALOG_COLUMNS))}{max(sheet.max_row, 1)}"


def _summary_cells(summary: GroupSummary) -> list[object]:
    return [
        summary.total,
        summary.replied,
        summary.silent,
        summary.pending,
        summary.reply_rate if summary.reply_rate is not None else "—",
        summary.silent_rate if summary.silent_rate is not None else "—",
    ]


def _header(sheet: Worksheet, names: tuple[str, ...]) -> None:
    sheet.append(list(names))
    for cell in sheet[sheet.max_row]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGNMENT


def _stamp(moment: datetime | None, timezone: ZoneInfo) -> str:
    return moment.astimezone(timezone).strftime("%d.%m.%Y %H:%M") if moment else ""
