import io
import re
from dataclasses import dataclass
from datetime import date
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.modules.wb_card_checklist.application.view import ChecklistView
from backend.modules.wb_card_checklist.domain import COUNTED_ITEMS, ITEMS, ITEMS_BY_KEY

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

IDENTITY_HEADERS = ("Дата заведения", "Артикул WB", "Артикул продавца", "Баркод", "Наименование товара", "ИП")
IDENTITY_WIDTHS = (14, 13, 22, 16, 36, 20)
FIRST_ITEM_COLUMN = len(IDENTITY_HEADERS) + 1
LAST_ITEM_COLUMN = FIRST_ITEM_COLUMN + len(ITEMS) - 1
DONE_COLUMN = LAST_ITEM_COLUMN + 1
STATUS_COLUMN = DONE_COLUMN + 1
COMMENT_COLUMN = STATUS_COLUMN + 1

# Цвета — из таблицы, которую менеджеры вели руками: тёмная шапка для
# реквизитов, синяя для пунктов проверки.
IDENTITY_FILL = PatternFill("solid", fgColor="1F3864")
ITEM_FILL = PatternFill("solid", fgColor="2E5F8A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTERED = Alignment(horizontal="center", vertical="center")


@dataclass(frozen=True, slots=True)
class ChecklistReportFile:
    seller_name: str
    day: date
    content: bytes

    @property
    def filename(self) -> str:
        slug = re.sub(r"[^\w-]+", "-", self.seller_name, flags=re.UNICODE).strip("-") or "seller"
        return f"checklist_{slug}_{self.day.isoformat()}.xlsx"


def render_workbook(view: ChecklistView, timezone: ZoneInfo) -> bytes:
    """The checklist as a spreadsheet: identity columns, then what WB shows.

    Counted items are TRUE/FALSE (empty when WB gave no data), the reference
    item carries its figure as text, and «готово» and the status are
    formulas over the item columns, as in the managers' own table.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Чек-лист"
    total = len(COUNTED_ITEMS)
    done_letter = get_column_letter(DONE_COLUMN)
    first_letter = get_column_letter(FIRST_ITEM_COLUMN)
    last_letter = get_column_letter(LAST_ITEM_COLUMN)
    status_letter = get_column_letter(STATUS_COLUMN)

    headers = [*IDENTITY_HEADERS, *(item.title for item in ITEMS), f"Готово из {total}", "Статус", "Комментарий"]
    sheet.append(headers)
    for column, _ in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.fill = ITEM_FILL if FIRST_ITEM_COLUMN <= column <= LAST_ITEM_COLUMN else IDENTITY_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
    sheet.row_dimensions[1].height = 42

    for index, row in enumerate(view.rows, start=2):
        created = row.card_created_at.astimezone(timezone).date() if row.card_created_at else None
        sheet.append(
            [
                created,
                row.article,
                row.vendor_code,
                row.barcode,
                row.title,
                view.seller_name,
                # COUNTIF считает только TRUE: текст справочного пункта и
                # пустые ячейки «нет данных» в готовность не попадают.
                *(item.done if ITEMS_BY_KEY[item.key].counted else item.detail for item in row.items),
                f"=COUNTIF({first_letter}{index}:{last_letter}{index},TRUE)",
                f'=IF({done_letter}{index}={total},"ГОТОВ","НЕ ГОТОВ")',
                row.comment or None,
            ]
        )
        sheet.cell(row=index, column=1).number_format = "DD.MM.YYYY"
        # Артикул и баркод — текст: иначе Excel покажет баркод как 2,05E+12.
        sheet.cell(row=index, column=2).number_format = "@"
        sheet.cell(row=index, column=4).number_format = "@"
        for column in range(FIRST_ITEM_COLUMN, STATUS_COLUMN + 1):
            sheet.cell(row=index, column=column).alignment = CENTERED

    last_row = max(len(view.rows) + 1, 2)
    status_range = f"{status_letter}2:{status_letter}{last_row}"
    sheet.conditional_formatting.add(
        status_range,
        CellIsRule(
            operator="equal",
            formula=['"ГОТОВ"'],
            font=Font(bold=True, color="006100"),
            fill=PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE"),
        ),
    )
    sheet.conditional_formatting.add(
        status_range,
        CellIsRule(
            operator="equal",
            formula=['"НЕ ГОТОВ"'],
            font=Font(bold=True, color="9C0006"),
            fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE"),
        ),
    )

    for column, width in enumerate(IDENTITY_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(FIRST_ITEM_COLUMN, LAST_ITEM_COLUMN + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 14
    sheet.column_dimensions[done_letter].width = 12
    sheet.column_dimensions[status_letter].width = 13
    sheet.column_dimensions[get_column_letter(COMMENT_COLUMN)].width = 36
    sheet.freeze_panes = sheet.cell(row=2, column=FIRST_ITEM_COLUMN)

    _instruction(workbook.create_sheet("Инструкция"), view.min_stock)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _instruction(sheet, min_stock: int) -> None:
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    total = len(COUNTED_ITEMS)
    rows: list[tuple[str, str]] = [
        ("Чек-лист карточки товара — выгрузка из сервиса", ""),
        ("", ""),
        ("Какие товары в таблице", f"Все карточки кабинета с остатком от {min_stock} шт. (WB и свои склады)."),
        ("Откуда данные", "Всё берётся из данных WB, руками ничего не отмечается."),
        (f"Колонка «Готово из {total}»", f"Формула COUNTIF по пунктам строки. «ГОТОВ» — когда выполнены все {total}."),
        ("Пустая ячейка", "У WB нет данных по пункту: это «не знаем», а не «не выполнено»."),
        ("", ""),
        ("Пункт", "Что означает"),
    ]
    rows.extend((item.title, item.meaning) for item in ITEMS)
    for values in rows:
        sheet.append(list(values))
    sheet["A1"].font = Font(bold=True, size=13)
    for row in (3, 4, 5, 6, 8):
        sheet.cell(row=row, column=1).font = bold
    sheet.cell(row=8, column=2).font = bold
    for line in sheet.iter_rows(min_row=3):
        for cell in line:
            cell.alignment = wrap
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 90
