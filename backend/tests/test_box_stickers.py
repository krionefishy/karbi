"""Стикеры коробов FBO: Excel коробовки + PDF стикеров WB → PDF в порядке Excel.

Файлы здесь синтетические, но повторяют то, что отдаёт кабинет WB: стикер
75×120 мм с коротким номером текстом и двумя одинаковыми QR с полным кодом,
Excel с колонками коробовки и размером листа «A1» в разметке.
"""

import io
import re
import uuid
import zipfile
from collections.abc import Sequence

import pytest
import segno
from httpx import AsyncClient
from openpyxl import Workbook
from pypdf import PdfReader
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from backend.modules.wb_box_stickers.application import (
    BoxStickerService,
    StickerBuildError,
    StickerInputError,
    read_stickers,
    read_workbook,
)
from backend.modules.wb_box_stickers.application.stickers import _font
from backend.modules.wb_box_stickers.domain import BoxLine, Package
from backend.modules.wb_box_stickers.infrastructure.wb import WBSuppliesClient
from backend.modules.wb_core.domain import BarcodeCard, Seller
from backend.modules.wb_core.domain.entities import EGRESS_VERIFIED
from backend.modules.wb_core.infrastructure.postgres import SellerRepository
from backend.tests.egress_stub import make_gateway

SUPPLY = 40000001
SELLER_ID = uuid.UUID("11111111-1111-1111-1111-111111111111")
OTHER_ID = uuid.UUID("22222222-2222-2222-2222-222222222222")
DRILL, SAW = "2050000000011", "2050000000028"

# Короб: короткий номер, полный код, содержимое.
BOXES: list[tuple[str, str, list[tuple[str, int]]]] = [
    ("6180474", "$Ts;AAA;0001;1;Xk00001;TAS", [(DRILL, 20)]),
    ("6180488", "$Ts;BBB;0002;1;Xk00002;TAS", [(SAW, 36)]),
    ("6180553", "$Ts;CCC;0003;1;Xk00003;TAS", [(DRILL, 10), (SAW, 5)]),
    ("6180598", "$Ts;DDD;0004;1;Xk00004;TAS", [(SAW, 20)]),
]


def sticker_pdf(boxes: Sequence[tuple[str, str, list[tuple[str, int]]]], *, supply: int = SUPPLY) -> bytes:
    font = _font()
    buffer = io.BytesIO()
    sheet = canvas.Canvas(buffer, pagesize=(212.6, 340.16))
    for shk, code, lines in boxes:
        qr = io.BytesIO()
        segno.make_qr(code, error="m").save(qr, kind="png", scale=6, border=2)
        image = ImageReader(io.BytesIO(qr.getvalue()))
        sheet.setFont(font, 12)
        sheet.drawCentredString(106, 318, f"{shk[:3]} {shk[3:]}")
        sheet.setFont(font, 10)
        sheet.drawCentredString(106, 304, "Тип поставки: Короб")
        sheet.drawCentredString(106, 290, f"Кол-во товаров: {sum(q for _, q in lines)} шт")
        sheet.drawImage(image, 46, 160, 120, 120)
        sheet.drawString(10, 140, f"№ поставки: {supply}")
        sheet.drawString(10, 125, "Плановая дата: 04.10.26")
        sheet.drawString(10, 95, "Продавец: ИП Тестов Т. Т.")
        sheet.drawImage(image, 150, 6, 56, 56)
        sheet.showPage()
    sheet.save()
    return buffer.getvalue()


def workbook_xlsx(rows: Sequence[tuple[str, int, str, str]], *, broken_dimension: bool = True) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(
        ["Баркод товара", "Кол-во товаров", "ШК короба", "Срок годности", "ШК короба для печати в стороннем сервисе"]
    )
    for barcode, quantity, shk, code in rows:
        sheet.append([barcode, quantity, shk, None, code])
    buffer = io.BytesIO()
    book.save(buffer)
    if not broken_dimension:
        return buffer.getvalue()
    # Как у WB: размер листа «A1» — read_only без сброса видит одну ячейку.
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', data)
            target.writestr(item, data)
    return output.getvalue()


def excel_rows(order: Sequence[int], *, filled: bool = True) -> list[tuple[str, int, str, str]]:
    rows = []
    for index in order:
        shk, code, lines = BOXES[index]
        if not filled:
            rows.append(("", 0, shk, code))
            continue
        rows.extend((barcode, quantity, shk, code) for barcode, quantity in lines)
    return rows


def seller(seller_id: uuid.UUID, name: str) -> Seller:
    return Seller(seller_id, name, 2, "success", None, None, egress_status=EGRESS_VERIFIED)


class FakeSellers(SellerRepository):
    def __init__(self) -> None:
        self.sellers = [seller(OTHER_ID, "Другой кабинет"), seller(SELLER_ID, "ИП Тестов Т.Т.")]

    async def list_sellers(self, *, include_archived: bool = False) -> list[Seller]:
        return list(self.sellers)

    async def list_barcode_cards(self, seller_id: uuid.UUID) -> dict[str, BarcodeCard]:
        return {
            DRILL: BarcodeCard(DRILL, "1001", "KARBI-Дрель", "Дрели", "0"),
            SAW: BarcodeCard(SAW, "1002", "KARBI-Пила", "Пилы", "XL"),
        }


class FakeSupplies(WBSuppliesClient):
    def __init__(self, packages: list[Package] | None = None) -> None:
        super().__init__(make_gateway())
        self.asked: list[str] = []
        self.rows = packages if packages is not None else [package(box) for box in BOXES]

    async def owns(self, seller_id: str, supply_id: int) -> bool:
        self.asked.append(seller_id)
        return seller_id == str(SELLER_ID) and supply_id == SUPPLY

    async def packages(self, seller_id: str, supply_id: int) -> list[Package]:
        return list(self.rows)


def package(box: tuple[str, str, list[tuple[str, int]]]) -> Package:
    _, code, lines = box
    return Package(code, sum(q for _, q in lines), tuple(BoxLine(b, q) for b, q in lines))


def service(supplies: FakeSupplies | None = None) -> BoxStickerService:
    return BoxStickerService(FakeSellers(), supplies or FakeSupplies())


def test_workbook_keeps_row_order_despite_wb_dimension() -> None:
    boxes = read_workbook(workbook_xlsx(excel_rows([2, 0, 3, 1])))
    assert [box.shk for box in boxes] == ["6180553", "6180474", "6180598", "6180488"]
    assert boxes[0].lines == (BoxLine(DRILL, 10), BoxLine(SAW, 5))
    assert boxes[0].package_code == BOXES[2][1]


def test_workbook_accepts_numbers_and_rejects_foreign_files() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["ШК короба", "Баркод товара", "Кол-во товаров"])
    sheet.append([6180474.0, 2050000000011, 3])
    buffer = io.BytesIO()
    book.save(buffer)
    [box] = read_workbook(buffer.getvalue())
    assert (box.shk, box.package_code, box.lines) == ("6180474", "", (BoxLine(DRILL, 3),))

    with pytest.raises(StickerInputError, match="ШК короба"):
        read_workbook(_plain_book())
    with pytest.raises(StickerInputError, match="не открывается"):
        read_workbook(b"not an excel file")


def _plain_book() -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Артикул", "Количество"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_stickers_read_text_and_qr() -> None:
    pages = read_stickers(sticker_pdf(BOXES))
    assert [(page.shk, page.package_code, page.quantity) for page in pages] == [
        (shk, code, sum(q for _, q in lines)) for shk, code, lines in BOXES
    ]
    assert {page.supply_id for page in pages} == {SUPPLY}
    assert pages[0].seller_name == "ИП Тестов Т. Т."
    with pytest.raises(StickerInputError, match="PDF"):
        read_stickers(b"%PDF-broken")


async def test_build_follows_excel_order_and_passes_own_check() -> None:
    supplies = FakeSupplies()
    order = [3, 1, 0, 2]
    plan, data = await service(supplies).build(workbook_xlsx(excel_rows(order)), sticker_pdf(BOXES), stamp=True)
    # Кабинет с именем со стикера спрашивают первым — чужим поставку не показывали.
    assert supplies.asked == [str(SELLER_ID)]
    assert (plan.supply_id, plan.seller_id, plan.problems) == (SUPPLY, SELLER_ID, [])
    assert [box.shk for box in plan.boxes] == [BOXES[index][0] for index in order]
    assert plan.boxes[3].items[1].vendor_code == "KARBI-Пила"

    pages = read_stickers(data)
    assert [page.package_code for page in pages] == [BOXES[index][1] for index in order]
    first = PdfReader(io.BytesIO(data)).pages[0].extract_text()
    assert "№ 1 из 4 · 20 шт" in first
    assert "KARBI-Пила (XL) × 20" in first


async def test_build_without_stamp_and_without_filled_contents() -> None:
    plan, data = await service().build(
        workbook_xlsx(excel_rows([1, 0, 2, 3], filled=False)), sticker_pdf(BOXES), stamp=False
    )
    assert plan.ready
    assert [page.shk for page in read_stickers(data)] == ["6180488", "6180474", "6180553", "6180598"]


async def test_problems_block_the_build() -> None:
    wrong = [package(BOXES[0]), package(BOXES[1]), Package(BOXES[2][1], 15, (BoxLine(DRILL, 15),)), package(BOXES[3])]
    plan = await service(FakeSupplies(wrong)).plan(workbook_xlsx(excel_rows([0, 1, 2])), sticker_pdf(BOXES))
    assert not plan.ready
    problems = "\n".join(plan.problems)
    assert "618 0553: в Excel" in problems
    assert "В PDF есть стикеры коробов, которых нет в Excel: 618 0598" in problems

    with pytest.raises(StickerBuildError):
        await service(FakeSupplies(wrong)).build(workbook_xlsx(excel_rows([0, 1, 2])), sticker_pdf(BOXES), stamp=True)


async def test_quantity_on_sticker_must_match_wb() -> None:
    changed = [Package(BOXES[0][1], 25, (BoxLine(DRILL, 25),)), *[package(box) for box in BOXES[1:]]]
    plan = await service(FakeSupplies(changed)).plan(
        workbook_xlsx(excel_rows([0, 1, 2, 3], filled=False)), sticker_pdf(BOXES)
    )
    assert plan.problems == ["Короб 618 0474: на стикере 20 шт, а в WB 25 шт — скачайте стикеры заново"]


async def test_unknown_supply_names_the_cause() -> None:
    plan = await service().plan(workbook_xlsx(excel_rows([0])), sticker_pdf(BOXES[:1], supply=40000999))
    assert plan.seller_id is None
    assert any("Поставка 40000999 не нашлась" in problem for problem in plan.problems)


async def test_routes_need_login(client: AsyncClient) -> None:
    files = {"workbook": ("a.xlsx", b"x"), "stickers": ("a.pdf", b"y")}
    response = await client.post("/api/v1/wb/box-stickers/check", files=files)
    assert response.status_code == 401
