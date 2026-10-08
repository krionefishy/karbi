"""Финансовые отчёты: ОПиУ из зеркала отчётов реализации, себестоимость и выгрузка."""

import io
import uuid
from collections.abc import AsyncIterator
from dataclasses import fields, replace
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Any

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, update

from backend.app.application import Application
from backend.modules.fin_reports.application import (
    CostFileError,
    FactsBuilder,
    match_cabinets,
    read_cost_file,
    week_ads,
)
from backend.modules.fin_reports.domain import (
    ARTICLE_COLUMNS,
    GRANULARITY_MONTH,
    GRANULARITY_WEEK,
    AdSpend,
    CostBook,
    CostPrice,
    Period,
    deduction_kind,
    parse_period,
    period_of,
    statement,
)
from backend.modules.fin_reports.infrastructure.postgres import CostPriceModel, FinReportsRepository, WbReportModel
from backend.modules.wb_core.domain import (
    AdvertCampaign,
    AdvertNmStat,
    AdvertSpend,
    SalesReport,
    SalesReportRow,
    SalesReportTotals,
)
from backend.modules.wb_core.infrastructure.postgres import MirrorRepository
from backend.modules.wb_core.infrastructure.postgres.models import SellerModel
from backend.shared.settings import load_settings
from backend.workers.fin_reports.worker import FinReportsWorker

SETTINGS = load_settings("backend/shared/settings/config.test.yaml")
API = "/api/v1/fin-reports"
AUTOMATION = "/api/v1/automations/fin-reports/sellers"
DRILL, SAW = 1112466805, 1466514852


def money(value: str | int) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"))


def blank(kind: Any) -> Any:
    """Пустое значение поля по его типу; необязательные поля — `None`."""
    if kind is Decimal:
        return money(0)
    return {int: 0, str: "", bool: False}.get(kind)


def totals(**values: Any) -> SalesReportTotals:
    """Сложенные строки отчёта: всё нулевое, кроме названного."""
    base = {item.name: blank(item.type) for item in fields(SalesReportTotals)}
    base.update(report_id=1, nm_id=DRILL)
    base.update(values)
    return SalesReportTotals(**base)


# --- расчёт -----------------------------------------------------------------------


def test_the_statement_reproduces_the_financiers_week_to_the_kopeck() -> None:
    """Байбурин, неделя 21–27.09.2026: итоги WB на входе, строки отчёта финансиста на выходе."""
    figures = statement(
        [
            totals(
                doc_type_name="Продажа",
                quantity=1265,
                gross=money("9281775.92"),
                retail_amount=money("5524919.34"),
                for_pay=money("5614547.66"),
            ),
            totals(
                doc_type_name="Возврат",
                quantity=5,
                gross=money("29259"),
                retail_amount=money("17985"),
                for_pay=money("17985"),
            ),
            totals(seller_oper_name="Логистика", delivery_service=money("333462.07")),
            totals(seller_oper_name="Штраф", penalty=money("9676.94")),
            totals(seller_oper_name="Хранение", paid_storage=money("291.1")),
            totals(seller_oper_name="Обработка товара", paid_acceptance=money("480")),
            totals(
                seller_oper_name="Удержание",
                bonus_type_name="Оказание услуг «WB Продвижение», документ №316847084",
                deduction=money("1436770"),
            ),
            totals(
                seller_oper_name="Удержание",
                bonus_type_name='Возврат неиспользованного остатка аванса за услугу "Баллы за отзывы"',
                deduction=money("-520364.16"),
            ),
            # Живые строки недели: у двух возвратов лояльность положительная, WB её вычитает.
            totals(
                doc_type_name="Продажа", cashback_amount=money("199482"), cashback_commission_change=money("10761.16")
            ),
            totals(doc_type_name="Возврат", cashback_amount=money("942"), cashback_commission_change=money("94.2")),
        ],
        # 1260 проданных за вычетом возвратов по 1532,45 ₽ дают себестоимость недели из отчёта.
        lambda nm_id: money("1532.45"),
    )

    values = figures.values
    assert values["revenue_before_spp"] == money("9252516.92")
    assert values["planned_commission"] == money("-3655954.26")
    assert values["revenue_after_spp"] == money("5506934.34")
    assert values["for_pay"] == money("5596562.66")
    assert values["logistics"] == money("-333462.07")
    assert values["commission"] == money("89628.32")
    assert values["penalties"] == money("-9676.94")
    assert values["storage"] == money("-291.10")
    assert values["acceptance"] == money("-480")
    assert values["advertising"] == money("-1436770")
    # Возврат аванса за отзывы больше нового аванса — строка положительная.
    assert values["reviews"] == money("520364.16")
    assert values["loyalty_points"] == money("-198540")
    assert values["loyalty_program"] == money("-10666.96")
    assert values["cost"] == money("-1930887.00")
    assert values["direct_expenses"] == sum(
        (values[key] for key in ("cost", "logistics", "commission", "penalties", "storage", "acceptance")),
        values["advertising"] + values["reviews"] + values["loyalty_points"] + values["loyalty_program"],
    )
    assert values["gross_margin"] == values["revenue_after_spp"] + values["direct_expenses"]
    assert figures.uncosted == 0


def test_sales_without_a_cost_price_are_counted_apart() -> None:
    seen: list[tuple[int, Decimal]] = []
    figures = statement(
        [
            totals(doc_type_name="Продажа", quantity=2, gross=money(2000), retail_amount=money(1500)),
            totals(nm_id=SAW, doc_type_name="Продажа", quantity=1, gross=money(700), retail_amount=money(600)),
            totals(nm_id=SAW, doc_type_name="Возврат", quantity=1, gross=money(100), retail_amount=money(90)),
        ],
        lambda nm_id: money(400) if nm_id == DRILL else None,
        on_uncosted=lambda item, revenue: seen.append((item.nm_id, revenue)),
    )

    assert figures.values["cost"] == money(-800)
    assert figures.uncosted == money(600)
    assert seen == [(SAW, money(700)), (SAW, money(-100))]


@pytest.mark.parametrize(
    ("name", "line"),
    [
        ("Оказание услуг «WB Продвижение», документ №310258471", "advertising"),
        ('Аванс за услугу "Баллы за отзывы"', "reviews"),
        ('Возврат неиспользованного остатка аванса за услугу "Баллы за отзывы"', "reviews"),
        ("Предоставление услуг по подписке «Джем», документ №305140207", "other_deductions"),
        ("Перенос карточек товара", "other_deductions"),
        ("Отчет об утилизированном товаре (по складу) за август 2026, документ №315137878", "other_deductions"),
        ("Разовое изменение срока перечисления денежных средств", "payout_term"),
        ("", "other_deductions"),
    ],
)
def test_deductions_are_told_apart_by_the_name_wb_gives_them(name: str, line: str) -> None:
    assert deduction_kind(name) == line


def test_a_week_split_by_the_month_end_is_one_week_and_two_months() -> None:
    """WB закрывает отчёт 31 августа и открывает новый 1 сентября — неделя одна, месяца два."""
    last_august, first_september = date(2026, 8, 31), date(2026, 9, 1)

    assert period_of(last_august, GRANULARITY_WEEK) == period_of(first_september, GRANULARITY_WEEK)
    assert period_of(last_august, GRANULARITY_WEEK).key == "2026-W36"
    assert period_of(last_august, GRANULARITY_MONTH).key == "2026-M08"
    assert period_of(first_september, GRANULARITY_MONTH).key == "2026-M09"
    # Неделя с 29 декабря — уже первая неделя следующего года.
    assert period_of(date(2025, 12, 29), GRANULARITY_WEEK) == Period(2026, 1, GRANULARITY_WEEK)
    assert parse_period("2026-W39") == Period(2026, 39, GRANULARITY_WEEK)
    week = Period(2026, 39, GRANULARITY_WEEK)
    assert week.label(date(2026, 9, 21), date(2026, 9, 27)) == "39 (21.09-27.09)"
    assert Period(2026, 9, GRANULARITY_MONTH).label(date(2026, 9, 1), date(2026, 9, 27)) == "сентябрь"
    for junk in ("2026", "2026-X01", "2026-W60", "abc-W01"):
        with pytest.raises(ValueError):
            parse_period(junk)


def test_a_new_cost_price_starts_on_its_date_and_the_first_one_covers_the_past() -> None:
    seller = uuid.uuid4()
    book = CostBook(
        [
            CostPrice(seller, "wb", "1", "drill", money(900), date(2026, 10, 1)),
            CostPrice(seller, "wb", "1", "drill", money(1000), date(2026, 11, 1)),
        ]
    )

    # До первой загрузки файла — по ней же: иначе вся история осталась бы без себестоимости.
    assert book.cost("1", date(2026, 3, 1)) == money(900)
    assert book.cost("1", date(2026, 10, 31)) == money(900)
    assert book.cost("1", date(2026, 11, 1)) == money(1000)
    assert book.cost("2", date(2026, 11, 1)) is None


# --- файл себестоимости ---------------------------------------------------------------


def cost_workbook(header: list[str], rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Информация об артикуле", None, None, None, None, "Цены и себестоимость"])
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


WB_HEADER = ["Кабинет", "АртикулВБ", "Артикул пост.", "Иден-ор КТ", "Штрих Код", "Цена", "Себес, руб", "Проч.затр, руб"]
OZON_HEADER = ["Кабинет", "Наименование", "Артикул ОЗОН", "Артикул пост.", "SKU", "ШтрихКод", "Себес, руб"]


def test_the_wb_cost_file_is_read_by_its_header_and_other_costs_are_ignored() -> None:
    content = cost_workbook(
        WB_HEADER,
        [
            ["Ип Байбурин", DRILL, "KARBI - Шуруповерт", "867457092", "2051917005518", "6147", 1526, "50"],
            ["Ип Мунаева", 1281924316, "Щетка", "1", "2", "1868", "550,5", "50"],
            ["Ип Мунаева", 1270964388, "Дождевики", "1", "2", "953", "дорого", "50"],
            [None, None, None, None, None, None, None, None],
            ["Ип Мунаева", None, "Без артикула", "1", "2", "1", 10, "50"],
        ],
    )

    parsed = read_cost_file(content)

    assert parsed.marketplace == "wb"
    assert [(row.cabinet, row.article, row.cost) for row in parsed.rows] == [
        ("Ип Байбурин", str(DRILL), money(1526)),
        ("Ип Мунаева", "1281924316", money("550.50")),
    ]
    assert parsed.problems == [
        "Строка 5: себестоимость «дорого» — не число",
        "Строка 7: нет кабинета или артикула",
    ]


def test_the_ozon_cost_file_is_keyed_by_sku() -> None:
    parsed = read_cost_file(
        cost_workbook(OZON_HEADER, [["Ип Яночкина", "Пылесос", 5686528578, "Автопылесос", 5236200315, "OZN1", 730]])
    )

    assert parsed.marketplace == "ozon"
    assert [(row.article, row.vendor_code, row.cost) for row in parsed.rows] == [
        ("5236200315", "Автопылесос", money(730))
    ]


def test_a_file_that_is_not_a_cost_file_is_refused() -> None:
    with pytest.raises(CostFileError):
        read_cost_file(b"not an excel file")
    with pytest.raises(CostFileError):
        read_cost_file(cost_workbook(["Кабинет", "Артикул"], [["Ип Байбурин", 1]]))


def test_cabinets_are_matched_to_sellers_by_surname() -> None:
    baiburin, saniev, twin_a, twin_b, alpha = (uuid.uuid4() for _ in range(5))
    sellers = {
        baiburin: "Байбурин",
        saniev: "ИП Саниев Д.Т.",
        twin_a: "Иванов",
        twin_b: "ИП Иванов А.А.",
        alpha: "ООО Альфа",
    }

    matched = match_cabinets(
        ["Ип Байбурин", "Ип Саниев", "Ип Иванов", "Ип Неизвестный", "ООО Бета", "ооо альфа", "ИП"], sellers
    )

    # Фамилия у двух селлеров сразу — кабинет не угадывается, как и незнакомый. Общего
    # первого слова мало: себестоимость «ООО Бета» не должна лечь на «ООО Альфа».
    assert matched == {"Ип Байбурин": baiburin, "Ип Саниев": saniev, "ооо альфа": alpha}


# --- через API, на зеркале ------------------------------------------------------------


def report(report_id: int, date_from: date, date_to: date, report_type: int = 1) -> SalesReport:
    zero = money(0)
    return SalesReport(
        report_id=report_id,
        report_type=report_type,
        period="weekly",
        date_from=date_from,
        date_to=date_to,
        create_date=None,
        currency="RUB",
        seller_finance_name="ИП Тест",
        **{item.name: zero for item in fields(SalesReport) if item.name.endswith("_sum")},
    )


def row(rrd_id: int, report_id: int, **values: Any) -> SalesReportRow:
    base = {item.name: blank(item.type) for item in fields(SalesReportRow)}
    base.update(rrd_id=rrd_id, report_id=report_id, nm_id=DRILL, vendor_code="KARBI - Шуруповерт")
    return replace(SalesReportRow(**base), **values)


def sale(rrd_id: int, report_id: int, *, gross: int, after: int, pay: int, **values: Any) -> SalesReportRow:
    return row(
        rrd_id,
        report_id,
        doc_type_name="Продажа",
        seller_oper_name="Продажа",
        quantity=1,
        retail_price_withdisc=money(gross),
        retail_amount=money(after),
        for_pay=money(pay),
        **values,
    )


W38, W39_MAIN, W39_BUYOUT, AUG31, SEP01, PENDING, JUL_AUG = 3801, 3901, 3902, 3601, 3602, 4001, 3101


@pytest_asyncio.fixture
async def application() -> AsyncIterator[Application]:
    application = Application(SETTINGS)
    app = application.get_app()
    async with app.router.lifespan_context(app):
        yield application


@pytest_asyncio.fixture
async def seller(application: Application) -> AsyncIterator[uuid.UUID]:
    stamp = datetime(2026, 9, 29, 6, 0, tzinfo=UTC)
    async with application.database.session() as session:
        model = SellerModel(name="ИП Финтест Ф.Ф.", catalog_sync_status="success", egress_status="verified")
        session.add(model)
        await session.flush()
        seller_id = model.id
        mirror = MirrorRepository(session)
        await mirror.upsert_sales_reports(
            seller_id,
            [
                # Неделю на стыке июля и августа WB не поделил: один отчёт, операции в двух месяцах.
                report(JUL_AUG, date(2026, 7, 27), date(2026, 8, 2)),
                report(AUG31, date(2026, 8, 31), date(2026, 8, 31)),
                report(SEP01, date(2026, 9, 1), date(2026, 9, 6)),
                report(W38, date(2026, 9, 14), date(2026, 9, 20)),
                report(W39_MAIN, date(2026, 9, 21), date(2026, 9, 27)),
                report(W39_BUYOUT, date(2026, 9, 21), date(2026, 9, 27), report_type=2),
                report(PENDING, date(2026, 9, 28), date(2026, 10, 4)),
            ],
            now=stamp,
        )
        await mirror.insert_sales_report_rows(
            seller_id,
            [
                sale(11, JUL_AUG, gross=400, after=300, pay=280, rr_date=date(2026, 7, 31)),
                sale(12, JUL_AUG, gross=700, after=500, pay=480, rr_date=date(2026, 8, 1)),
                sale(1, AUG31, gross=1000, after=800, pay=700),
                sale(2, SEP01, gross=2000, after=1500, pay=1400),
                sale(3, W38, gross=3000, after=2400, pay=2500),
                sale(4, W39_MAIN, gross=5000, after=4000, pay=4200),
                sale(5, W39_MAIN, gross=900, after=700, pay=650, nm_id=SAW, vendor_code="KARBI - Пила"),
                row(
                    6,
                    W39_MAIN,
                    doc_type_name="Возврат",
                    quantity=1,
                    retail_price_withdisc=money(1000),
                    retail_amount=money(800),
                    for_pay=money(840),
                ),
                row(7, W39_MAIN, seller_oper_name="Логистика", delivery_service=money(300)),
                row(
                    8,
                    W39_MAIN,
                    seller_oper_name="Удержание",
                    deduction=money(500),
                    bonus_type_name="Оказание услуг «WB Продвижение», документ №1",
                ),
                sale(9, W39_BUYOUT, gross=600, after=500, pay=450),
                sale(10, PENDING, gross=9999, after=9999, pay=9999),
            ],
            now=stamp,
        )
        for report_id in (JUL_AUG, AUG31, SEP01, W38, W39_MAIN, W39_BUYOUT):
            await mirror.advance_sales_report(seller_id, report_id, cursor=10, loaded_at=stamp)
        await session.commit()
    # Отчёт на странице — из сумм, которые складывает воркер; здесь его проход сделан руками.
    outcome = await FactsBuilder(application.database).build(seller_id, limit=100, now=stamp)
    assert (outcome.built, outcome.left) == (6, 0)
    try:
        yield seller_id
    finally:
        async with application.database.session() as session:
            await FinReportsRepository(session).purge(seller_id)
            await session.execute(delete(SellerModel).where(SellerModel.id == seller_id))
            await session.commit()


@pytest_asyncio.fixture
async def client(application: Application) -> AsyncIterator[AsyncClient]:
    token = application.token_service.issue_access(uuid.uuid4())
    async with AsyncClient(
        transport=ASGITransport(app=application.get_app()),
        base_url="http://test",
        headers={"Authorization": f"Bearer {token}"},
    ) as http:
        yield http


async def connect(client: AsyncClient, seller: uuid.UUID) -> None:
    response = await client.post(AUTOMATION, json={"seller_id": str(seller)})
    assert response.status_code in (200, 201), response.text


def own(body: dict, seller: uuid.UUID, period: str) -> dict[str, float]:
    """Цифры своего кабинета: в общей тестовой базе могут быть подключены и другие."""
    column = next(item for item in body["periods"] if item["key"] == period)
    return column["by_seller"][str(seller)]


async def upload(client: AsyncClient, rows: list[list[Any]], effective_from: str | None = None):
    content = cost_workbook(WB_HEADER, rows)
    data = {"effective_from": effective_from} if effective_from else None
    return await client.post(f"{API}/costs", files={"workbook": ("costs.xlsx", content)}, data=data)


async def test_weeks_add_both_report_types_and_show_what_is_not_loaded_yet(
    client: AsyncClient, seller: uuid.UUID
) -> None:
    await connect(client, seller)

    response = await client.get(API, params={"year": 2026, "granularity": "week", "seller_id": str(seller)})

    assert response.status_code == 200, response.text
    body = response.json()
    assert [item["key"] for item in body["periods"]] == ["2026-W40", "2026-W39", "2026-W38", "2026-W36", "2026-W31"]
    week = own(body, seller, "2026-W39")
    # Основной отчёт и отчёт «по выкупам» за неделю складываются; возврат вычитается.
    assert (week["revenue_before_spp"], week["revenue_after_spp"], week["for_pay"]) == (5500.0, 4400.0, 4460.0)
    assert (week["planned_commission"], week["commission"]) == (-1040.0, 60.0)
    assert (week["logistics"], week["advertising"]) == (-300.0, -500.0)
    assert week["gross_margin"] == 4400.0 + 60.0 - 300.0 - 500.0
    # Неделя на стыке месяцев — два отчёта WB, одна колонка.
    split = next(item for item in body["periods"] if item["key"] == "2026-W36")
    assert (split["label"], split["by_seller"][str(seller)]["revenue_before_spp"]) == ("36 (31.08-06.09)", 3000.0)
    # Отчёт недели 40 зеркало ещё не дочитало: денег нет, и это сказано.
    pending = next(item for item in body["periods"] if item["key"] == "2026-W40")
    assert pending["pending_sellers"] == [str(seller)] and str(seller) not in pending["by_seller"]
    [state] = body["sellers"]
    assert (state["reports"], state["pending_reports"]) == (7, 1)
    assert body["total_by_seller"][str(seller)]["revenue_before_spp"] == 12600.0
    assert [line["key"] for line in body["lines"]][:4] == [
        "revenue_before_spp",
        "planned_commission",
        "revenue_after_spp",
        "for_pay",
    ]


async def test_months_split_a_week_by_the_date_wb_booked_each_row(client: AsyncClient, seller: uuid.UUID) -> None:
    """Отчёт недели на стыке месяцев чаще один: его строки расходятся по месяцам по дате операции."""
    await connect(client, seller)

    body = (await client.get(API, params={"year": 2026, "granularity": "month", "seller_id": str(seller)})).json()

    assert [(item["key"], item["label"]) for item in body["periods"]] == [
        ("2026-M10", "октябрь"),
        ("2026-M09", "сентябрь"),
        ("2026-M08", "август"),
        ("2026-M07", "июль"),
    ]
    assert own(body, seller, "2026-M07")["revenue_before_spp"] == 400.0
    # Август: 1 августа из недели 27.07–02.08 и отдельный отчёт WB за 31 августа.
    assert own(body, seller, "2026-M08")["revenue_before_spp"] == 700.0 + 1000.0
    assert own(body, seller, "2026-M09")["revenue_before_spp"] == 2000.0 + 3000.0 + 5500.0
    # Недочитанный отчёт 28.09–04.10 делает неполными оба месяца, которых касается.
    pending = {item["key"] for item in body["periods"] if item["pending_sellers"]}
    assert pending == {"2026-M09", "2026-M10"}
    august = next(item for item in body["periods"] if item["key"] == "2026-M08")
    assert (august["date_from"], august["date_to"]) == ("2026-08-01", "2026-08-31")
    assert body["total_by_seller"][str(seller)]["revenue_before_spp"] == 12600.0


async def test_a_cost_file_fills_the_cost_line_and_lists_what_has_no_cost(
    client: AsyncClient, seller: uuid.UUID, application: Application
) -> None:
    await connect(client, seller)
    params: dict[str, str | int] = {"year": 2026, "granularity": "week", "seller_id": str(seller)}
    before = (await client.get(API, params=params)).json()
    assert own(before, seller, "2026-W39")["cost"] == 0.0
    assert {item["nm_id"] for item in before["uncosted"]} == {DRILL, SAW}

    response = await upload(
        client,
        [
            ["Ип Финтест", DRILL, "KARBI - Шуруповерт", "1", "2", "6147", 1500, "50"],
            ["Ип Чужой", 1, "x", "1", "2", "1", 5, "50"],
        ],
    )

    assert response.status_code == 200, response.text
    result = response.json()
    assert (result["marketplace"], result["added"], result["changed"], result["unchanged"]) == ("wb", 1, 0, 0)
    assert result["unknown_cabinets"] == ["Ип Чужой"]
    after = (await client.get(API, params=params)).json()
    week = own(after, seller, "2026-W39")
    # Две продажи шуруповёрта (основной отчёт и «по выкупам») минус возврат; пила без себестоимости.
    assert week["cost"] == -1500.0
    assert week["gross_margin"] == own(before, seller, "2026-W39")["gross_margin"] - 1500.0
    assert [(item["nm_id"], item["vendor_code"], item["revenue"]) for item in after["uncosted"]] == [
        (SAW, "KARBI - Пила", 900.0)
    ]
    # Тот же файл ещё раз — версий не прибавилось; новая цена — новая версия, прошлое не тронуто.
    again = (await upload(client, [["Ип Финтест", DRILL, "KARBI - Шуруповерт", "1", "2", "6147", 1500, "50"]])).json()
    assert (again["added"], again["changed"], again["unchanged"]) == (0, 0, 1)
    async with application.database.session() as session:
        session.add(
            CostPriceModel(
                seller_id=seller,
                marketplace="wb",
                article=str(DRILL),
                effective_from=date(2026, 9, 21),
                vendor_code="KARBI - Шуруповерт",
                cost=money(2000),
            )
        )
        # Первая версия задним числом: до 21 сентября действует она.
        await session.execute(
            update(CostPriceModel)
            .where(CostPriceModel.seller_id == seller, CostPriceModel.cost == money(1500))
            .values(effective_from=date(2026, 1, 1))
        )
        await session.commit()
    repriced = (await client.get(API, params=params)).json()
    assert own(repriced, seller, "2026-W38")["cost"] == -1500.0
    assert own(repriced, seller, "2026-W39")["cost"] == -2000.0


async def test_a_wrong_first_price_is_corrected_by_a_back_dated_upload(client: AsyncClient, seller: uuid.UUID) -> None:
    """Периоды до первой версии считаются по ней — без даты в прошлом опечатку в истории не убрать."""
    await connect(client, seller)
    params: dict[str, str | int] = {"year": 2026, "granularity": "week", "seller_id": str(seller)}
    line = ["Ип Финтест", DRILL, "KARBI - Шуруповерт", "1", "2", "6147"]
    await upload(client, [[*line, 9000, "50"]])
    assert own((await client.get(API, params=params)).json(), seller, "2026-W38")["cost"] == -9000.0

    # Тот же кабинет записан двумя способами — это один селлер и одна версия, а не отказ базы.
    fixed = await upload(
        client,
        [["ИП Финтест Ф.Ф.", DRILL, "x", "1", "2", "1", 1, "50"], [*line, 900, "50"]],
        effective_from="2026-01-01",
    )

    assert fixed.status_code == 200, fixed.text
    assert (fixed.json()["added"], fixed.json()["changed"], fixed.json()["effective_from"]) == (0, 1, "2026-01-01")
    assert own((await client.get(API, params=params)).json(), seller, "2026-W38")["cost"] == -900.0
    future = await upload(client, [[*line, 900, "50"]], effective_from="2999-01-01")
    assert future.status_code == 422


async def test_bad_queries_and_files_are_refused(client: AsyncClient, seller: uuid.UUID) -> None:
    await connect(client, seller)

    assert (await client.get(API, params={"granularity": "day"})).status_code == 422
    assert (await client.get(API, params={"year": 1999})).status_code == 422
    assert (await client.get(API, params={"seller_id": str(uuid.uuid4())})).status_code == 422
    assert (await client.get(f"{API}/export", params={"year": 2026, "period": "вчера"})).status_code == 422
    refused = await client.post(f"{API}/costs", files={"workbook": ("costs.xlsx", b"junk")})
    assert refused.status_code == 422 and "Excel" in refused.json()["detail"]


async def test_the_export_has_the_period_by_cabinet_weeks_months_and_missing_costs(
    client: AsyncClient, seller: uuid.UUID
) -> None:
    await connect(client, seller)

    response = await client.get(f"{API}/export", params={"year": 2026, "period": "2026-W39"})

    assert response.status_code == 200, response.text
    assert "pnl_wb_2026-W39" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == [
        "ОПиУ 39 (21.09-27.09)",
        "Недели 2026",
        "Месяцы 2026",
        "По артикулам 39 нед.",
        "Остатки 39 нед.",
        "Без себестоимости",
    ]
    articles = workbook["По артикулам 39 нед."]
    header = [cell.value for cell in articles[1]]
    assert header[:5] == ["ИП", "Артикул ВБ", "Артикул продавца", "Размер", "Доставки"] and len(header) == 1 + len(
        ARTICLE_COLUMNS
    )
    own_articles = [row for row in articles.iter_rows(min_row=2, values_only=True) if row[0] == "ИП Финтест Ф.Ф."]
    drill = next(row for row in own_articles if row[1] == str(DRILL))
    # Две продажи шуруповёрта минус возврат; выручка до СПП 5000 + 600 − 1000.
    assert (drill[7], drill[13], drill[header.index("Вся стоимость реализованного товара до СПП")]) == (2, 1, 4600)
    assert any("Остатки — из зеркала" in str(row[0]) for row in articles.iter_rows(min_row=2, values_only=True))
    period = workbook.worksheets[0]
    header = [cell.value for cell in period[1]]
    assert header[:2] == ["Статья", "WB итого"] and "ИП Финтест Ф.Ф." in header
    column = header.index("ИП Финтест Ф.Ф.") + 1
    labels = [period.cell(row=index, column=1).value for index in range(2, period.max_row + 1)]
    assert labels[:4] == [
        "Реализация (до СПП)",
        "Плановая комиссия",
        "Реализация (после СПП)",
        "К перечислению за товар",
    ]
    assert period.cell(row=2, column=column).value == 5500
    assert labels[-1] == "Валовая маржа"
    weeks = workbook["Недели 2026"]
    assert [cell.value for cell in weeks[1]][1:3] == ["2026 ИТОГО", "40 (28.09-04.10)"]
    missing = workbook["Без себестоимости"]
    own_rows = [row for row in missing.iter_rows(min_row=2, values_only=True) if row[0] == "ИП Финтест Ф.Ф."]
    assert {row[1] for row in own_rows} == {str(DRILL), str(SAW)}


async def test_a_report_shows_up_only_after_the_worker_has_folded_it(
    client: AsyncClient, seller: uuid.UUID, application: Application
) -> None:
    """Зеркало дочитало отчёт, но воркер ещё не сложил — период «неполный», а не с дырой в цифрах."""
    await connect(client, seller)
    params: dict[str, str | int] = {"year": 2026, "granularity": "week", "seller_id": str(seller)}
    stamp = datetime(2026, 10, 5, 9, 0, tzinfo=UTC)
    async with application.database.session() as session:
        await MirrorRepository(session).advance_sales_report(seller, PENDING, cursor=10, loaded_at=stamp)
        await session.commit()

    before = (await client.get(API, params=params)).json()
    week = next(item for item in before["periods"] if item["key"] == "2026-W40")
    assert week["pending_sellers"] == [str(seller)] and before["sellers"][0]["pending_reports"] == 1

    worker = FinReportsWorker(application.database, SETTINGS, now=lambda: stamp)
    assert await worker.tick() >= 1
    # Сложено — и до следующего интервала воркер зеркало не трогает; сложенное не пересобирается.
    assert await worker.tick() == 0
    interval = timedelta(minutes=SETTINGS.fin_reports.build_interval_minutes)
    later = FinReportsWorker(application.database, SETTINGS, now=lambda: stamp + interval)
    assert await later.tick() == 0

    after = (await client.get(API, params=params)).json()
    week = next(item for item in after["periods"] if item["key"] == "2026-W40")
    assert week["pending_sellers"] == [] and week["by_seller"][str(seller)]["revenue_before_spp"] == 9999.0
    [state] = after["sellers"]
    assert (state["pending_reports"], state["built_at"]) == (0, stamp.isoformat())


async def test_a_long_backlog_is_folded_in_portions_without_waiting_for_the_interval(
    seller: uuid.UUID, application: Application
) -> None:
    stamp = datetime(2026, 10, 5, 9, 0, tzinfo=UTC)
    async with application.database.session() as session:
        repository = FinReportsRepository(session)
        await repository.track(seller)
        await session.execute(delete(WbReportModel).where(WbReportModel.seller_id == seller))
        await session.commit()
    builder = FactsBuilder(application.database)

    first = await builder.build(seller, limit=4, now=stamp)
    second = await builder.build(seller, limit=4, now=stamp)
    third = await builder.build(seller, limit=4, now=stamp)

    # Старшие отчёты первыми; остаток добирается следующим проходом, сложенное не трогается.
    assert [(first.built, first.left), (second.built, second.left), (third.built, third.left)] == [
        (4, 2),
        (2, 0),
        (0, 0),
    ]
    async with application.database.session() as session:
        assert len(await FinReportsRepository(session).built_reports(seller)) == 6


# --- реклама по артикулам из зеркала ----------------------------------------------------


async def test_campaign_spend_is_split_between_articles_by_their_stats(
    application: Application, seller: uuid.UUID
) -> None:
    stamp = datetime(2026, 10, 5, 9, 0, tzinfo=UTC)
    week = (date(2026, 9, 21), date(2026, 9, 27))
    async with application.database.session() as session:
        mirror = MirrorRepository(session)
        await mirror.upsert_advert_campaigns(
            seller,
            [
                AdvertCampaign(1, "пара", 9, "cpm", "unified", (DRILL, SAW), None),
                AdvertCampaign(2, "без статистики", 9, "cpm", "unified", (DRILL, SAW), None),
            ],
            now=stamp,
        )
        await mirror.replace_advert_spend(
            seller,
            *week,
            [
                AdvertSpend(1, date(2026, 9, 22), "Баланс", Decimal("300")),
                AdvertSpend(1, date(2026, 9, 23), "Счет", Decimal("100")),
                AdvertSpend(2, date(2026, 9, 24), "Бонусы", Decimal("50")),
                AdvertSpend(3, date(2026, 9, 24), "Баланс", Decimal("7")),
                AdvertSpend(1, date(2026, 9, 28), "Баланс", Decimal("999")),
            ],
            now=stamp,
        )
        await mirror.replace_advert_nm_stats(
            seller,
            [1],
            *week,
            [
                AdvertNmStat(1, date(2026, 9, 22), DRILL, 0, 0, 0, 0, 0, 0, Decimal("150"), Decimal(0)),
                AdvertNmStat(1, date(2026, 9, 23), SAW, 0, 0, 0, 0, 0, 0, Decimal("50"), Decimal(0)),
            ],
            now=stamp,
        )
        await mirror.save_advert_cursor(seller, date(2026, 9, 26))
        await session.commit()

        ads = await week_ads(session, seller, since=week[0], until=week[1])

    drill, saw = ads.of(DRILL), ads.of(SAW)
    # Кампания 1: 400 ₽ делятся 3:1 по статистике; кампания 2 без статистики — поровну; кампания 3 неизвестна.
    assert (drill.balance, drill.account, drill.bonus, drill.promotion_info) == (
        Decimal("225.00"),
        Decimal("75.00"),
        Decimal("25.00"),
        Decimal("150.00"),
    )
    assert (saw.balance, saw.account, saw.bonus, saw.promotion_info) == (
        Decimal("75.00"),
        Decimal("25.00"),
        Decimal("25.00"),
        Decimal("50.00"),
    )
    assert ads.unallocated.balance == Decimal("7.00") and ads.collected_through == date(2026, 9, 26)
    assert ads.of(999) == AdSpend()
