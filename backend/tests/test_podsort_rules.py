import io
from datetime import UTC, date, datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from backend.modules.wb_core.domain import WarehouseRemain
from backend.modules.wb_podsort.application import CollectionService, PodsortWorkbook, summarize
from backend.modules.wb_podsort.application.calculation import build_rows
from backend.modules.wb_podsort.application.report import ReportContext, SellerSection
from backend.modules.wb_podsort.domain import (
    CENTRAL,
    BarcodeInfo,
    OrderLine,
    Periods,
    PodsortSettings,
    RegionFigures,
    guess_warehouse_region,
    order_region,
)
from backend.modules.wb_podsort.infrastructure.postgres import CountTotals

MOSCOW = ZoneInfo("Europe/Moscow")
DRILL, BINOCULARS, SAW = "2053497104469", "2052163348695", "2055210050862"
SETTINGS = PodsortSettings(window_days=7, cover_days=7, regions=(CENTRAL, "Приволжский"))


def total(barcode: str, region: str, window: int, *, fbs: int = 0, month: int | None = None) -> CountTotals:
    current = window if month is None else month
    return CountTotals(barcode, region, (0, 0, current), window, window, window, fbs, current)


def remain(barcode: str, warehouse: str, quantity: int) -> WarehouseRemain:
    return WarehouseRemain(barcode, "1271611253", "0", "KARBI - Шуруповерт желтый", warehouse, quantity)


def test_orders_fall_into_wb_regions() -> None:
    assert order_region("Россия", "Центральный федеральный округ") == CENTRAL
    # Округа склеены как в отчёте WB.
    assert order_region("Россия", "Северо-Кавказский федеральный округ") == "Южный и Северо-Кавказский"
    assert order_region("Россия", "Сибирский федеральный округ") == "Дальневосточный и Сибирский"
    assert order_region("Беларусь", "") == "Беларусь"
    assert order_region("Россия", "") == "Россия"
    # Страну, которой нет в списке, не прячем в «Россию».
    assert order_region("Монголия", "") == "Монголия"


def test_warehouses_are_placed_by_city() -> None:
    assert guess_warehouse_region("Коледино") == CENTRAL
    assert guess_warehouse_region("Самара (Новосемейкино)") == "Приволжский"
    assert guess_warehouse_region("СПБ Шушары") == "Северо-Западный"
    assert guess_warehouse_region("СЦ Симферополь (Молодежненское)") == "Южный и Северо-Кавказский"
    assert guess_warehouse_region("Нижний Новгород") == "Приволжский"
    assert guess_warehouse_region("Екатеринбург - Перспективный 12") == "Уральский"
    assert guess_warehouse_region("СК Ереван") == "Армения"
    # Не склад и склад без места — не угадываются.
    assert guess_warehouse_region("В пути до получателей") is None
    assert guess_warehouse_region("Склад WB РФ") is None
    assert guess_warehouse_region("Неведомый склад") is None


def test_need_subtracts_the_stock_already_in_the_region() -> None:
    # Шуруповёрт из книги 23.09: 61 шт. за неделю, а в ЦФО уже лежит 580.
    assert RegionFigures(window_orders=61, stock=580).need(7, 7) == 0
    assert RegionFigures(window_orders=61, stock=0).need(7, 7) == 61
    assert RegionFigures(window_orders=61, stock=20).need(7, 7) == 41
    # Две недели продаж на неделю вперёд — половина, половина вверх как в Excel.
    assert RegionFigures(window_orders=15, stock=0).need(14, 7) == 8
    assert RegionFigures(window_orders=0, stock=10).cover(7) is None
    assert RegionFigures(window_orders=14, stock=10).cover(7) == 5.0


def test_periods_are_full_days_up_to_yesterday() -> None:
    periods = Periods(date(2026, 9, 24))

    assert periods.last_day == date(2026, 9, 23)
    assert periods.window_start(7) == date(2026, 9, 17)
    assert periods.months == (date(2026, 7, 1), date(2026, 8, 1), date(2026, 9, 1))
    assert periods.month_end(date(2026, 8, 1)) == date(2026, 8, 31)
    assert periods.month_end(date(2026, 9, 1)) == date(2026, 9, 23)
    # Первое число: вчерашний день — прошлый месяц, он и последний в книге.
    assert Periods(date(2026, 10, 1)).months == (date(2026, 7, 1), date(2026, 8, 1), date(2026, 9, 1))


def test_a_day_is_summed_by_barcode_and_region() -> None:
    line = OrderLine(DRILL, 1271611253, "KARBI - Шуруповерт", "Шуруповерты", "0", CENTRAL, fbs=False)
    fbs = OrderLine(DRILL, 1271611253, "KARBI - Шуруповерт", "Шуруповерты", "0", CENTRAL, fbs=True)
    other = OrderLine(SAW, 1466514852, "KARBI - Пила", "Пилы", "0", "Беларусь", fbs=False)

    counts, infos = summarize([line, fbs, line, other])

    assert [(c.barcode, c.region, c.orders, c.fbs_orders) for c in counts] == [
        (DRILL, CENTRAL, 3, 1),
        (SAW, "Беларусь", 1, 0),
    ]
    assert {info.barcode for info in infos} == {DRILL, SAW}


def collection() -> CollectionService:
    return CollectionService(
        None,  # type: ignore[arg-type]
        None,  # type: ignore[arg-type]
        None,  # type: ignore[arg-type]
        timezone=MOSCOW,
        history_days=5,
        days_per_run=2,
        settle_hours=6,
        refresh_minutes=60,
    )


def test_missing_days_come_newest_first_and_unsettled_days_are_reread() -> None:
    now = datetime(2026, 9, 24, 10, 0, tzinfo=MOSCOW)
    service = collection()
    loaded = {
        # Вчера прочитали в 01:00 — WB ещё дописывал заказы; перечитать можно, прошёл час.
        date(2026, 9, 23): datetime(2026, 9, 24, 1, 0, tzinfo=MOSCOW),
        # Позавчера прочитали после 06:00 следующего дня — сутки устоялись.
        date(2026, 9, 22): datetime(2026, 9, 23, 7, 0, tzinfo=MOSCOW),
        date(2026, 9, 20): datetime(2026, 9, 21, 7, 0, tzinfo=MOSCOW),
    }

    assert service.due_days(loaded, now=now) == [date(2026, 9, 23), date(2026, 9, 21), date(2026, 9, 19)]

    # Только что перечитанный неустоявшийся день ждёт `refresh_minutes`.
    loaded[date(2026, 9, 23)] = datetime(2026, 9, 24, 9, 30, tzinfo=MOSCOW)
    assert service.due_days(loaded, now=now) == [date(2026, 9, 21), date(2026, 9, 19)]


def test_rows_subtract_regional_stock_and_keep_unplaced_stock_aside() -> None:
    totals = [
        total(DRILL, CENTRAL, 61, fbs=40),
        total(DRILL, "Приволжский", 14),
        total(BINOCULARS, CENTRAL, 31),
        total(SAW, CENTRAL, 53),
    ]
    remains = [
        remain(DRILL, "Коледино", 577),
        remain(DRILL, "Котовск", 3),
        remain(DRILL, "Самара (Новосемейкино)", 5),
        remain(DRILL, "Склад WB РФ", 300),
        remain(DRILL, "В пути до получателей", 200),
        remain(DRILL, "Всего находится на складах", 885),
        remain(BINOCULARS, "Тула", 10),
        remain(BINOCULARS, "Неведомый склад", 50),
    ]
    infos = {DRILL: BarcodeInfo(DRILL, 1271611253, "KARBI - Шуруповерт желтый", "Шуруповерты", "0")}

    rows = build_rows("Байбурин", totals, remains, infos, {"Неведомый склад": CENTRAL}, SETTINGS, months=3)

    by_barcode = {row.info.barcode: row for row in rows}
    drill = by_barcode[DRILL]
    assert drill.targets[CENTRAL].stock == 580
    assert drill.targets[CENTRAL].need(7, 7) == 0
    assert drill.targets["Приволжский"].stock == 5
    assert drill.targets["Приволжский"].need(7, 7) == 9
    # «Склад WB РФ» — остаток без места; «в пути» и итог отчёта не считаются вовсе.
    assert drill.unplaced_stock == 300
    assert drill.window_orders == 75 and drill.window_fbs_orders == 40
    assert drill.month_by_region == {CENTRAL: 61, "Приволжский": 14}
    # Склад, которому регион задал человек, вычитается как любой другой.
    assert by_barcode[BINOCULARS].targets[CENTRAL].stock == 60
    assert by_barcode[BINOCULARS].targets[CENTRAL].need(7, 7) == 0
    assert by_barcode[SAW].targets[CENTRAL].need(7, 7) == 53
    # Сначала то, что продаётся быстрее.
    assert [row.info.barcode for row in rows] == [DRILL, SAW, BINOCULARS]


def test_the_workbook_explains_the_subtraction_next_to_the_numbers() -> None:
    rows = build_rows(
        "Байбурин",
        [total(DRILL, CENTRAL, 61), total(SAW, CENTRAL, 53)],
        [remain(DRILL, "Коледино", 580), remain(SAW, "Неведомый склад", 7)],
        {},
        {},
        SETTINGS,
        months=3,
    )
    periods = Periods(date(2026, 9, 24))
    context = ReportContext(
        today=periods.today,
        last_day=periods.last_day,
        window_start=periods.window_start(7),
        months=periods.months,
        settings=SETTINGS,
        unplaced_warehouses=("Неведомый склад",),
    )
    content = PodsortWorkbook(
        context,
        [
            SellerSection("Байбурин", rows, 7, datetime(2026, 9, 24, 6, 0, tzinfo=UTC)),
            SellerSection("ИП Мунаева М.Л.", [], 3, None),
        ],
    ).build()

    book = load_workbook(io.BytesIO(content))
    assert book.sheetnames == ["Подсорт 24.09.2026", "Подсорт Байбурин", "Подсорт ИП Мунаева М Л", "Как читать"]
    summary = book["Подсорт 24.09.2026"]
    notes = [summary.cell(line, 1).value or "" for line in range(1, 10)]
    assert any("Остаток WB в регионе вычтен" in note for note in notes)
    assert any("Остаток нашего склада не учитывается" in note for note in notes)
    assert any("Неведомый склад" in note for note in notes)
    assert any("ИП Мунаева М.Л.: заказы загружены за 3 из 7" in note for note in notes)
    assert any("ИП Мунаева М.Л.: остатки по складам WB ещё не собирались" in note for note in notes)
    values = [[cell.value for cell in row] for row in summary.iter_rows()]
    header = next(row for row in values if row[0] == "Кабинет")
    assert header[5] == "Расчёт системы"
    body = [row for row in values if row[0] == "Байбурин"]
    # Шуруповёрт не везём — в ЦФО его на месяцы вперёд; пилу — везём.
    assert [(row[2], row[5], row[12]) for row in body] == [(SAW, 53, 0)]
    assert summary.cell(values.index(header) + 1, 6).comment is not None
    seller = book["Подсорт Байбурин"]
    assert seller.cell(2, 6).value == "июль 2026"
    assert seller.cell(3, 3).value == DRILL
