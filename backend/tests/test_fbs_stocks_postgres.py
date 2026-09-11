import io
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from backend.modules.wb_core.infrastructure.postgres import SellerRepository
from backend.modules.wb_core.infrastructure.postgres.models import SellerModel
from backend.modules.wb_core.infrastructure.wb import CatalogCard
from backend.modules.wb_fbs_stocks.application import BoardConflictError, CollectionService, FbsStocksService
from backend.modules.wb_fbs_stocks.domain import GROUP_DISTRICT, GROUP_FULFILMENT, GROUP_OWN, SellerWarehouse
from backend.modules.wb_fbs_stocks.infrastructure.postgres import FbsStocksRepository, TrackedSellerModel
from backend.modules.wb_fbs_stocks.infrastructure.wb import WBFbsStocksClient
from backend.shared.settings import load_settings
from backend.storage.pg import Database
from backend.tests.egress_stub import make_gateway
from backend.workers.wb_fbs_stocks.worker import FbsStocksWorker

SETTINGS = load_settings("backend/shared/settings/config.test.yaml")
MOSCOW = ZoneInfo("Europe/Moscow")

PITER = SellerWarehouse(2164974, 245, "4. Питер наш", 1, False)
KAZAN = SellerWarehouse(2197411, 168, "5. Казань наша", 1, False)
SPB_FF = SellerWarehouse(1917658, 10999, "Фулэксперт СПБ", 1, False)
VOLOGDA_FF = SellerWarehouse(2167144, 10128, "Фулэксперт Вологда-2", 1, False)
DBS = SellerWarehouse(1516422, 3049170, "DBS Новокузнецк", 2, False)
WAREHOUSES = [PITER, KAZAN, SPB_FF, VOLOGDA_FF, DBS]

FEN = "2052957817499"
UNKNOWN = "2000000000001"


class FakeClient(WBFbsStocksClient):
    def __init__(self, warehouses: list[SellerWarehouse], stocks: dict[int, dict[str, int]]) -> None:
        super().__init__(make_gateway())
        self.served_warehouses = warehouses
        self.served_stocks = stocks
        self.asked: list[int] = []

    async def warehouses(self, seller_id: str) -> list[SellerWarehouse]:
        return list(self.served_warehouses)

    async def stocks(self, seller_id: str, warehouse_id: int, skus) -> dict[str, int]:
        self.asked.append(warehouse_id)
        return {sku: amount for sku, amount in self.served_stocks.get(warehouse_id, {}).items() if sku in skus}


def fake_client() -> FakeClient:
    return FakeClient(WAREHOUSES, {SPB_FF.warehouse_id: {FEN: 28}, PITER.warehouse_id: {FEN: 0}})


@pytest_asyncio.fixture
async def database() -> AsyncIterator[Database]:
    database = Database()
    await database.connect(SETTINGS.database.url, pool_size=2, max_overflow=0)
    try:
        yield database
    finally:
        await database.disconnect()


@pytest_asyncio.fixture
async def seller(database: Database) -> AsyncIterator[uuid.UUID]:
    async with database.session() as session:
        model = SellerModel(name="ИП Остатки", catalog_sync_status="success")
        session.add(model)
        await session.flush()
        seller_id = model.id
        await FbsStocksRepository(session).track(seller_id)
        # Каталог кабинета: по нему строка получает название и артикул.
        await SellerRepository(session).upsert_catalog(
            seller_id,
            active=[
                CatalogCard(
                    article="1223328520",
                    vendor_code="Фен сенсорный",
                    name="Фен для волос",
                    sizes=[{"chrt_id": 1800164566, "tech_size": "0", "skus": [FEN]}],
                )
            ],
            archived=[],
            archived_available=True,
        )
        await session.commit()
    try:
        yield seller_id
    finally:
        async with database.session() as session:
            await FbsStocksRepository(session).purge_seller(seller_id)
            await session.execute(delete(SellerModel).where(SellerModel.id == seller_id))
            await session.commit()


def service(session: AsyncSession) -> FbsStocksService:
    return FbsStocksService(session, SellerRepository(session), FbsStocksRepository(session), timezone=MOSCOW)


async def seed(database: Database, seller_id: uuid.UUID) -> tuple[uuid.UUID, uuid.UUID]:
    """Зеркало складов, две группы в её порядке и два баркода."""
    async with database.session() as session:
        stocks = FbsStocksRepository(session)
        await stocks.replace_warehouses(seller_id, WAREHOUSES)
        await session.commit()
        board = service(session)
        northwest = await board.add_group(seller_id, "Северо-западный округ", GROUP_DISTRICT)
        own = await board.add_group(seller_id, "Наш склад", GROUP_OWN)
        await board.reorder_groups(seller_id, [own.id, northwest.id])
        await board.set_group_columns(seller_id, own.id, [PITER.warehouse_id, KAZAN.warehouse_id])
        await board.set_group_columns(seller_id, northwest.id, [SPB_FF.warehouse_id, VOLOGDA_FF.warehouse_id])
        assert await board.add_barcodes(seller_id, [FEN, UNKNOWN, FEN], None) == 2
    return own.id, northwest.id


async def test_collection_asks_only_the_columns_and_writes_zero_for_missing_rows(
    database: Database, seller: uuid.UUID
) -> None:
    await seed(database, seller)
    client = fake_client()

    async with database.session() as session:
        result = await CollectionService(session, FbsStocksRepository(session), client).collect(seller)

    # DBS-склад в таблице не стоит — его не спрашивали.
    assert sorted(client.asked) == sorted(
        [PITER.warehouse_id, KAZAN.warehouse_id, SPB_FF.warehouse_id, VOLOGDA_FF.warehouse_id]
    )
    assert (result.warehouses, result.polled, result.barcodes) == (5, 4, 2)
    async with database.session() as session:
        view = await service(session).view(seller)
    fen, unknown = view.rows
    assert (fen.title, fen.article, fen.in_catalog) == ("Фен для волос", "1223328520", True)
    assert unknown.in_catalog is False
    # Строки нет — ноль; сумма группы считается по её складам.
    assert fen.amount(SPB_FF.warehouse_id) == 28 and fen.amount(VOLOGDA_FF.warehouse_id) == 0
    own, northwest = view.groups
    assert [column.name for column in own.columns] == ["4. Питер наш", "5. Казань наша"]
    assert (fen.group_total(own), fen.group_total(northwest)) == (0, 28)
    assert view.collected_at is not None


async def test_a_warehouse_gone_or_being_deleted_leaves_the_table(database: Database, seller: uuid.UUID) -> None:
    await seed(database, seller)
    deleting = SellerWarehouse(KAZAN.warehouse_id, KAZAN.office_id, KAZAN.name, 1, True)
    client = FakeClient([PITER, deleting, SPB_FF, DBS], {})

    async with database.session() as session:
        await CollectionService(session, FbsStocksRepository(session), client).collect(seller)
        view = await service(session).view(seller)
        setup = await service(session).setup(seller)
        # Удаляемый склад остаётся в зеркале с пометкой, но в группу его не поставить.
        with pytest.raises(BoardConflictError):
            await service(session).set_group_columns(seller, view.groups[0].id, [KAZAN.warehouse_id])

    own, northwest = view.groups
    assert [column.name for column in own.columns] == ["4. Питер наш"]
    assert [column.name for column in northwest.columns] == ["Фулэксперт СПБ"]
    placed = {warehouse.name: warehouse.group_id for warehouse in setup.warehouses}
    assert placed["DBS Новокузнецк"] is None and placed["4. Питер наш"] == own.id
    assert placed["5. Казань наша"] is None


async def test_an_empty_warehouse_list_keeps_the_previous_collection(database: Database, seller: uuid.UUID) -> None:
    await seed(database, seller)
    async with database.session() as session:
        await CollectionService(session, FbsStocksRepository(session), fake_client()).collect(seller)

    worker = FbsStocksWorker(database, FakeClient([], {}), SETTINGS)
    assert await worker.collect(seller) is not None
    async with database.session() as session:
        view = await service(session).view(seller)
    # Столбцы и остатки прежнего сбора на месте, ошибка названа.
    assert [column.name for column in view.groups[0].columns] == ["4. Питер наш", "5. Казань наша"]
    assert view.rows[0].amount(SPB_FF.warehouse_id) == 28
    assert view.collection_error is not None


async def test_groups_and_barcodes_keep_the_order_they_were_added_in(database: Database, seller: uuid.UUID) -> None:
    async with database.session() as session:
        board = service(session)
        for title in ("Наш склад", "ФФ Стас", "Приволжский округ"):
            await board.add_group(seller, title, GROUP_DISTRICT)
        await board.add_barcodes(seller, ["3000"], None)
        await board.add_barcodes(seller, ["2000", "1000"], None)
        view = await board.view(seller)
    # Без явной позиции порядок был бы алфавитным, а не её.
    assert [group.title for group in view.groups] == ["Наш склад", "ФФ Стас", "Приволжский округ"]
    assert [row.barcode for row in view.rows] == ["3000", "2000", "1000"]


async def test_notes_and_removal_follow_the_barcode(database: Database, seller: uuid.UUID) -> None:
    await seed(database, seller)
    async with database.session() as session:
        board = service(session)
        await board.set_note(seller, FEN, "стоит в 1С", None)
        with pytest.raises(BoardConflictError):
            await board.set_note(seller, "нет такого", "x", None)
        with pytest.raises(BoardConflictError):
            await board.add_barcodes(seller, ["не баркод!"], None)
        await board.remove_barcode(seller, UNKNOWN)
        view = await board.view(seller)
    assert [(row.barcode, row.note) for row in view.rows] == [(FEN, "стоит в 1С")]


async def test_group_edits_keep_one_place_per_warehouse(database: Database, seller: uuid.UUID) -> None:
    own_id, northwest_id = await seed(database, seller)
    async with database.session() as session:
        board = service(session)
        # Питер переезжает в округ: из «Нашего склада» он уходит сам.
        await board.set_group_columns(seller, northwest_id, [PITER.warehouse_id, SPB_FF.warehouse_id])
        await board.update_group(seller, own_id, title="Свои", kind=GROUP_FULFILMENT)
        with pytest.raises(BoardConflictError):
            await board.set_group_columns(seller, own_id, [999])
        with pytest.raises(BoardConflictError):
            await board.reorder_groups(seller, [own_id])
        view = await board.view(seller)
        await board.delete_group(seller, northwest_id)
        after = await board.view(seller)
    own, northwest = view.groups
    assert (own.title, own.kind, [column.name for column in own.columns]) == ("Свои", "fulfilment", ["5. Казань наша"])
    assert [column.name for column in northwest.columns] == ["4. Питер наш", "Фулэксперт СПБ"]
    assert [group.title for group in after.groups] == ["Свои"]


async def test_export_follows_the_sellers_spreadsheet(database: Database, seller: uuid.UUID) -> None:
    await seed(database, seller)
    async with database.session() as session:
        await CollectionService(session, FbsStocksRepository(session), fake_client()).collect(seller)
        board = service(session)
        await board.set_note(seller, FEN, "дубль карточки", None)
        report = await board.export()

    workbook = load_workbook(io.BytesIO(report.content))
    assert workbook.sheetnames == ["ИП Остатки", "Сравнение"]
    sheet = workbook["ИП Остатки"]
    assert [cell.value for cell in sheet[1]] == [
        "ИП Остатки",
        "Баркод",
        "Наш склад",
        "4. Питер наш",
        "5. Казань наша",
        "Северо-западный округ",
        "Фулэксперт СПБ",
        "Фулэксперт Вологда-2",
    ]
    assert [cell.value for cell in sheet[2]] == ["дубль карточки", FEN, "=SUM(D2:E2)", 0, 0, "=SUM(G2:H2)", 28, 0]
    assert sheet["B2"].number_format == "@"
    # Группы сворачиваются, сводный столбец слева от своих складов.
    assert sheet.column_dimensions["D"].outline_level == 1 and sheet.column_dimensions["C"].outline_level == 0
    assert sheet.sheet_properties.outlinePr.summaryRight is False
    # Нули красятся условным форматом на всей области чисел.
    assert [str(area.sqref) for area in sheet.conditional_formatting] == ["C2:H3"]
    assert sheet.freeze_panes == "C2"

    comparison = workbook["Сравнение"]
    assert [cell.value for cell in comparison[1]] == [
        "ИП Остатки",
        "Баркод",
        "Наш склад",
        "4. Питер наш",
        "5. Казань наша",
        "Северо-западный округ",
    ]
    assert [cell.value for cell in comparison[2]] == ["дубль карточки", FEN, "=SUM(D2:E2)", 0, 0, 28]
    # Красятся только записанные ячейки: разрыв между блоками для Excel равен нулю.
    assert sorted(str(area.sqref) for area in comparison.conditional_formatting) == ["C2:E3", "F2:F3"]
    assert report.filename.startswith("fbs_stocks_")


async def test_worker_polls_on_an_interval_and_serves_the_button(database: Database, seller: uuid.UUID) -> None:
    await seed(database, seller)
    worker = FbsStocksWorker(database, fake_client(), SETTINGS)
    now = datetime.now(UTC)

    assert await worker.collect_due(now) == 1
    # Только что собрали — до истечения интервала кабинет не в очереди.
    assert await worker.collect_due(now) == 0
    assert await worker.collect_due(now + timedelta(minutes=SETTINGS.fbs_stocks.poll_minutes + 1)) == 1

    async with database.session() as session:
        board = service(session)
        first = await board.request_refresh(seller, None)
        again = await board.request_refresh(seller, None)
        assert first.requested_at == again.requested_at
        assert await board.request_refresh_all(None) == 1
    await worker.serve_refresh_requests()
    async with database.session() as session:
        state = await service(session).refresh_state(seller)
        tracked = await session.get(TrackedSellerModel, seller)
    assert state is not None and state.status == "success"
    assert tracked is not None and tracked.collection_error is None


async def test_a_seller_detached_mid_collection_is_not_written(database: Database, seller: uuid.UUID) -> None:
    await seed(database, seller)
    async with database.session() as session:
        await FbsStocksRepository(session).untrack(seller)
        await session.commit()

    async with database.session() as session:
        result = await CollectionService(session, FbsStocksRepository(session), fake_client()).collect(seller)
        assert result.skipped
        assert await FbsStocksRepository(session).facts(seller) == {}
