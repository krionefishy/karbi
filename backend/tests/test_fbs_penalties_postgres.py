import io
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from backend.modules.wb_core.domain import (
    ORDER_SOURCE_ARCHIVE,
    ORDER_SOURCE_LIVE,
    FbsOrder,
    FbsSupply,
    SellerWarehouse,
    WbOffice,
)
from backend.modules.wb_core.infrastructure.postgres import MirrorRepository, SellerRepository
from backend.modules.wb_core.infrastructure.postgres.models import SellerModel
from backend.modules.wb_fbs_penalties.application import (
    TRACE_FOUND,
    TRACE_NO_ORDER,
    TRACE_NO_SUPPLY,
    CollectionService,
    PenaltiesQueryError,
    PenaltiesService,
)
from backend.modules.wb_fbs_penalties.domain import GROUP_PENALTIES, GROUP_STORAGE, ReportRow
from backend.modules.wb_fbs_penalties.infrastructure.postgres import PenaltiesRepository, TrackedSellerModel
from backend.modules.wb_fbs_penalties.infrastructure.wb import WBRealizationClient
from backend.shared.settings import load_settings
from backend.storage.pg import Database
from backend.tests.egress_stub import make_gateway
from backend.workers.wb_fbs_penalties.worker import PenaltiesWorker

SETTINGS = load_settings("backend/shared/settings/config.test.yaml")
MOSCOW = ZoneInfo("Europe/Moscow")
NOW = datetime(2026, 9, 16, 9, 0, tzinfo=UTC)
TODAY = NOW.date()

KAZAN = SellerWarehouse(1917674, "5. Казань наша", 3088703)
SUPPLY = FbsSupply(
    "WB-GI-1",
    "Поставка",
    NOW - timedelta(days=20),
    NOW - timedelta(days=19),
    NOW - timedelta(days=19, hours=-4),
    15,
    True,
    1,
)


def report_row(
    rrd_id: int,
    *,
    srid: str,
    assembly_id: int | None,
    sticker: int | None,
    penalty: float = 0.0,
    storage: float = 0.0,
    day: date | None = None,
) -> ReportRow:
    when = day or (TODAY - timedelta(days=5))
    return ReportRow(
        rrd_id=rrd_id,
        realizationreport_id=1,
        date_from=when - timedelta(days=6),
        date_to=when,
        create_dt=when,
        srid=srid,
        assembly_id=assembly_id,
        sticker_id=sticker,
        order_dt=NOW - timedelta(days=21),
        sale_dt=None,
        rr_dt=when,
        nm_id=1271611253,
        sa_name="karbi - шуруповерт",
        subject_name="Шуруповерты",
        barcode="2053497104469",
        ts_name="0",
        bonus_type_name="Штраф за нарушение срока передачи товара" if penalty else "Платное хранение",
        supplier_oper_name="Штраф",
        delivery_method="FBS, (МГТ)",
        office_name="",
        penalty=penalty,
        deduction=0.0,
        rebill_logistic_cost=0.0,
        storage_fee=storage,
        additional_payment=0.0,
        acceptance=0.0,
    )


def order(order_id: int, rid: str, *, supply: str | None, sticker: int | None = None) -> FbsOrder:
    return FbsOrder(
        order_id=order_id,
        rid=rid,
        order_uid=rid.split(".")[1] if "." in rid else rid,
        created_at=NOW - timedelta(days=21),
        warehouse_id=KAZAN.warehouse_id,
        supply_id=supply,
        office_id=15,
        nm_id=1271611253,
        chrt_id=1,
        sku="2053497104469",
        price_kopecks=10000,
        sticker_id=sticker,
        supplier_status="complete" if sticker else None,
        wb_status="sold" if sticker else None,
        source=ORDER_SOURCE_ARCHIVE if sticker else ORDER_SOURCE_LIVE,
    )


class FakeRealization(WBRealizationClient):
    def __init__(self, rows: list[ReportRow]) -> None:
        super().__init__(make_gateway())
        self.rows_served = rows
        self.calls: list[tuple[date, date]] = []

    async def rows(self, seller_id: str, date_from: date, date_to: date) -> list[ReportRow]:
        self.calls.append((date_from, date_to))
        return list(self.rows_served)


@pytest_asyncio.fixture
async def database() -> AsyncIterator[Database]:
    database = Database()
    await database.connect(SETTINGS.database.url, pool_size=2, max_overflow=0)
    try:
        yield database
    finally:
        await database.disconnect()


@pytest_asyncio.fixture
async def seller(database: Database) -> AsyncIterator[uuid.UUID]:
    async with database.session() as session:
        model = SellerModel(name="ИП Штрафы", catalog_sync_status="success", egress_status="verified")
        session.add(model)
        await session.flush()
        seller_id = model.id
        await PenaltiesRepository(session).track(seller_id)
        mirror = MirrorRepository(session)
        await mirror.replace_seller_warehouses(seller_id, [KAZAN], now=NOW)
        await mirror.upsert_supplies(seller_id, [SUPPLY], now=NOW)
        await mirror.upsert_offices([WbOffice(15, "Коледино", "Подольск", "")], now=NOW)
        await mirror.upsert_orders(
            seller_id,
            [
                order(5551052438, "er.i902a.0.0", supply="WB-GI-1"),
                order(5551052439, "er.i903b.0.0", supply=None),
                order(5551052440, "er.i904c.0.0", supply="WB-GI-1", sticker=33811984302),
            ],
            now=NOW,
        )
        await session.commit()
    try:
        yield seller_id
    finally:
        async with database.session() as session:
            await PenaltiesRepository(session).purge_seller(seller_id)
            await session.execute(delete(SellerModel).where(SellerModel.id == seller_id))
            await session.commit()


def service(session: AsyncSession) -> PenaltiesService:
    return PenaltiesService(
        session, SellerRepository(session), PenaltiesRepository(session), timezone=MOSCOW, orders_history_months=6
    )


ROWS = [
    report_row(1, srid="er.i902a.0.0", assembly_id=5551052438, sticker=57048832986, penalty=198.06),
    report_row(2, srid="er.i903b.0.0", assembly_id=5551052439, sticker=57048832987, penalty=374.98),
    report_row(3, srid="er.i999z.0.0", assembly_id=5551059999, sticker=57048832988, storage=12.5),
    report_row(4, srid="er.i905d.0.0", assembly_id=5551052441, sticker=None),
]


async def collect(
    database: Database, seller_id: uuid.UUID, rows: list[ReportRow], *, now: datetime = NOW
) -> FakeRealization:
    client = FakeRealization(rows)
    async with database.session() as session:
        await CollectionService(
            session, PenaltiesRepository(session), client, window_days=14, backfill_days=90
        ).collect(seller_id, now=now)
    return client


async def test_collection_keeps_only_charged_rows_and_backfills_on_first_run(
    database: Database, seller: uuid.UUID
) -> None:
    client = await collect(database, seller, ROWS)
    assert client.calls == [(TODAY - timedelta(days=90), TODAY)]

    async with database.session() as session:
        stored = await PenaltiesRepository(session).rows_in_period(seller, TODAY - timedelta(days=30), TODAY)
    # Строка без удержаний (4) не хранится; повторный сбор берёт короткое окно и ничего не задваивает.
    assert sorted(row.rrd_id for row in stored) == [1, 2, 3]

    again = await collect(database, seller, ROWS, now=NOW + timedelta(days=1))
    assert again.calls == [(TODAY + timedelta(days=1) - timedelta(days=14), TODAY + timedelta(days=1))]
    async with database.session() as session:
        assert len(await PenaltiesRepository(session).rows_in_period(seller, TODAY - timedelta(days=30), TODAY)) == 3


async def test_the_view_traces_each_row_to_its_warehouse_and_supply(database: Database, seller: uuid.UUID) -> None:
    await collect(database, seller, ROWS)

    async with database.session() as session:
        view = await service(session).view(seller, TODAY - timedelta(days=30), TODAY)

    by_id = {item.row.rrd_id: item for item in view.rows}
    found = by_id[1]
    assert (found.trace, found.warehouse_name, found.supply_id, found.destination_office_name) == (
        TRACE_FOUND,
        KAZAN.name,
        "WB-GI-1",
        "Коледино",
    )
    assert found.supply_scan_dt == SUPPLY.scan_dt
    assert by_id[2].trace == TRACE_NO_SUPPLY and by_id[2].warehouse_name == KAZAN.name
    assert by_id[3].trace == TRACE_NO_ORDER and by_id[3].warehouse_name is None
    totals = {total.group: (total.count, total.amount) for total in view.totals}
    assert totals == {GROUP_PENALTIES: (2, 573.04), GROUP_STORAGE: (1, 12.5)}
    assert [option.name for option in view.warehouses] == [KAZAN.name]

    async with database.session() as session:
        only_penalties = await service(session).view(seller, TODAY - timedelta(days=30), TODAY, group=GROUP_PENALTIES)
        outside = await service(session).view(seller, TODAY - timedelta(days=2), TODAY)
    assert [item.row.rrd_id for item in only_penalties.rows] == [1, 2] or sorted(
        item.row.rrd_id for item in only_penalties.rows
    ) == [1, 2]
    assert outside.rows == ()

    with pytest.raises(PenaltiesQueryError):
        async with database.session() as session:
            await service(session).view(seller, TODAY, TODAY - timedelta(days=1))


async def test_lookup_accepts_stickers_assembly_ids_and_srids(database: Database, seller: uuid.UUID) -> None:
    await collect(database, seller, ROWS)

    async with database.session() as session:
        result = await service(session).lookup(
            seller, ["57048832986", "er.i903b.0.0", "33811984302", "5551052439", "777", "???"]
        )

    by_key = {item.row.srid: item for item in result.rows}
    # Стикер из отчёта → строка со штрафом и складом; srid и номер задания той же строки не задваивают её.
    assert by_key["er.i902a.0.0"].row.amount == 198.06 and by_key["er.i902a.0.0"].trace == TRACE_FOUND
    assert by_key["er.i903b.0.0"].trace == TRACE_NO_SUPPLY
    # Стикер из архива заданий: в отчётах штрафа нет, но склад и поставка известны.
    stub = by_key["er.i904c.0.0"]
    assert (stub.row.rrd_id, stub.row.amount, stub.trace, stub.supply_id) == (0, 0.0, TRACE_FOUND, "WB-GI-1")
    assert {miss.key for miss in result.missing} == {"777", "???"}


async def test_export_writes_numbers_as_text_and_totals_on_top(database: Database, seller: uuid.UUID) -> None:
    await collect(database, seller, ROWS)
    async with database.session() as session:
        report = await service(session).export(seller, TODAY - timedelta(days=30), TODAY)

    assert report.filename.startswith("fbs_penalties_ИП_Штрафы_")
    sheet = load_workbook(io.BytesIO(report.content))["Штрафы"]
    assert sheet["A2"].value == "Штрафы" and sheet["B2"].value == 2 and sheet["C2"].value == 573.04
    header = [cell.value for cell in sheet[5]]
    assert header[:3] == ["Кабинет", "Неделя отчёта", "Баркод"] and "Поставка" in header
    stickers = {sheet.cell(row=index, column=9).value for index in range(6, 9)}
    assert "57048832986" in stickers and sheet.cell(row=6, column=9).number_format == "@"


async def test_worker_collects_once_a_day_and_serves_refresh(database: Database, seller: uuid.UUID) -> None:
    client = FakeRealization(ROWS)
    morning = datetime(2026, 9, 16, 7, 0, tzinfo=MOSCOW)
    worker = PenaltiesWorker(database, client, SETTINGS, now=lambda tz: morning.astimezone(tz))

    assert await worker.collect_due(morning) == 1
    assert await worker.collect_due(morning + timedelta(hours=2)) == 0
    assert len(client.calls) == 1

    async with database.session() as session:
        state = await service(session).request_refresh(seller, None)
    assert state.in_progress
    await worker.serve_refresh_requests()
    async with database.session() as session:
        done = await service(session).refresh_state(seller)
        tracked = await session.get(TrackedSellerModel, seller)
    assert done is not None and done.status == "success"
    assert tracked is not None and tracked.collected_at is not None and len(client.calls) == 2
