import io

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from backend.modules.wb_box_stickers.application import (
    BoxStickerService,
    PackingService,
    StickerBuildError,
    StickerInputError,
    pack,
    parse_assembly,
    parse_scheme,
    plan_packing,
    read_workbook,
)
from backend.modules.wb_box_stickers.domain import Assembly, AssemblyRow, StoredPacking, SupplyGood
from backend.modules.wb_box_stickers.infrastructure.redis import PackingStore
from backend.tests.test_box_stickers import BOXES, DRILL, SAW, FakeSellers, FakeSupplies, excel_rows, workbook_xlsx
from backend.tests.test_box_stickers import SUPPLY as STICKER_SUPPLY

SUPPLY = "41498020"
LAMP = "2050000000035"


def pasted(*rows: str) -> str:
    return "\n".join(rows)


MOSCOW = pasted(
    f"{SUPPLY}\t1304218079\t{DRILL}\t01.06.2026\tЗарядное устройство\t290\t290\t290\t290\t5-50+1-40",
    f"{SUPPLY}\t1179385821\t{SAW}\t7-16\tТриммер\t200\t200\t200\t200\t10-20",
    f"{SUPPLY}\t1304131904\t{LAMP}\t17\tОвощерезка\t9\t9\t9\t9\t01.09.2026",
)


def test_scheme_unfolds_into_boxes() -> None:
    assert parse_scheme("3-60+1-40") == (60, 60, 60, 40)
    assert parse_scheme("12 - 25") == (25,) * 12
    assert parse_scheme("1–16") == (16,)
    with pytest.raises(ValueError):
        parse_scheme("0-5")


def test_assembly_reads_rows_by_content_not_by_column() -> None:
    assembly = parse_assembly(MOSCOW)
    assert assembly.problems == ()
    assert assembly.supply_id == 41498020
    assert [(row.barcode, row.quantity, row.scheme, row.boxes) for row in assembly.rows] == [
        (DRILL, 290, "5-50+1-40", (50, 50, 50, 50, 50, 40)),
        (SAW, 200, "10-20", (20,) * 10),
        (LAMP, 9, "1-9", (9,)),
    ]
    assert assembly.boxes == 17

    shifted = pasted(
        f"{SUPPLY}\t317874626\t{DRILL}\t01.05.2026\tФонарь\t250\t250\t250\t4-60+1-10\t\t1п",
        f"{SUPPLY}\t1304131904\t{LAMP}\t20\tОвощерезка\t10\t10\t10\t01.10.2026",
    )
    assembly = parse_assembly(shifted)
    assert assembly.problems == ()
    assert [(row.quantity, row.boxes) for row in assembly.rows] == [(250, (60, 60, 60, 60, 10)), (10, (10,))]
    assert "Регионы" not in str(assembly)

    spaced = parse_assembly(f"{SUPPLY}  1304218079  {DRILL}  1-6  Зарядное устройство  290  290  290  5-50+1-40")
    assert spaced.problems == () and spaced.rows[0].boxes == (50, 50, 50, 50, 50, 40)


def test_assembly_skips_rows_outside_supply_and_names_the_rest() -> None:
    text = pasted(
        f"{SUPPLY}\t1304218079\t{DRILL}\t1-6\tЗарядное\t290\t290\t290\t290\t5-50+1-40",
        f"\t1248473004\t{LAMP}\t\tЭлектрощётка\t16\t16\t16\tне везем",
    )
    assembly = parse_assembly(text)
    assert assembly.problems == () and [row.barcode for row in assembly.rows] == [DRILL]

    broken = parse_assembly(
        pasted(
            f"{SUPPLY}\t1304218079\t{DRILL}\t1-6\tЗарядное\t290\t290\t290\t290\t5-50+1-40",
            f"{SUPPLY}\t1179385821\t{SAW}\t7-16\tТриммер\t200\t200\t200\t200\t",
            f"{SUPPLY}\t1304131904\t{LAMP}\t17\tОвощерезка\t9\t9\t9\t9\t2-5",
            f"{SUPPLY}\t1304131904\t{LAMP}\t17\tОвощерезка\t9\t9\t9\t9\t1-9",
        )
    )
    assert len(broken.problems) == 3
    assert "строка 2 (баркод" in broken.problems[0]
    assert "строка 3: по схеме «2-5» 10 шт, а в сборке 9 шт" in broken.problems[1]
    assert f"Баркод {LAMP} в строках 3 и 4" in broken.problems[2]

    two = parse_assembly(
        pasted(
            f"41498020\t1\t{DRILL}\t1\tА\t5\t5\t5\t5\t1-5",
            f"41498031\t1\t{SAW}\t1\tБ\t5\t5\t5\t5\t1-5",
        )
    )
    assert two.problems[0].startswith("Вставлены строки разных поставок: 41498020, 41498031")
    assert parse_assembly("Саниев 1561\nСЦ Домодедово-2\t02.10.26").problems == (
        "В вставленном тексте нет строк с баркодом — скопируйте строки поставки из таблицы сборки",
    )
    assert "нет номера поставки" in parse_assembly(f"\t1\t{DRILL}\t1\tА\t5\t5\t5\t5\t1-5").problems[0]


def _assembly(*rows: tuple[str, int, str]) -> Assembly:
    return Assembly(
        supply_id=41498020,
        rows=tuple(
            AssemblyRow(
                line=number,
                supply_id=41498020,
                barcode=barcode,
                quantity=quantity,
                scheme=scheme,
                boxes=parse_scheme(scheme),
            )
            for number, (barcode, quantity, scheme) in enumerate(rows, 1)
        ),
        problems=(),
    )


def test_packing_wants_exactly_as_many_empty_boxes_as_the_scheme() -> None:
    template = read_workbook(workbook_xlsx(excel_rows([0, 1, 2, 3], filled=False)))
    plan = plan_packing(template, _assembly((DRILL, 50, "2-25"), (SAW, 36, "1-36")))
    assert plan.problems == [
        "В шаблоне 4 коробов, а по таблице сборки нужно 3: удалите в кабинете 1 и скачайте шаблон заново"
    ]
    plan = plan_packing(template, _assembly((DRILL, 150, "6-25")))
    assert "нужно 6: добавьте в кабинете 2" in plan.problems[0]

    filled = read_workbook(workbook_xlsx(excel_rows([0, 1], filled=True)))
    plan = plan_packing(filled, _assembly((DRILL, 20, "1-20"), (SAW, 36, "1-36")))
    assert plan.problems[0].startswith("В шаблоне уже заполнены короба 6180474, 6180488")

    plan = plan_packing(template, _assembly((DRILL, 50, "2-25"), (SAW, 56, "1-36+1-20")))
    assert plan.ready
    assert [(box.position, box.shk, box.barcode, box.quantity) for box in plan.boxes] == [
        (1, "6180474", DRILL, 25),
        (2, "6180488", DRILL, 25),
        (3, "6180553", SAW, 36),
        (4, "6180598", SAW, 20),
    ]
    assert plan.boxes[2].package_code == BOXES[2][1]


def test_pack_writes_only_barcode_and_quantity_into_the_same_file() -> None:
    template = workbook_xlsx(excel_rows([3, 0, 2, 1], filled=False))
    plan, data = pack(template, _assembly((DRILL, 50, "2-25"), (SAW, 56, "1-36+1-20")))
    assert plan.ready and data is not None

    sheet = load_workbook(io.BytesIO(data)).active
    assert sheet is not None
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=1, max_row=5)]
    assert rows[0][:3] == ["Баркод товара", "Кол-во товаров", "ШК короба"]
    assert rows[1:] == [
        [DRILL, 25, "6180598", None, BOXES[3][1]],
        [DRILL, 25, "6180474", None, BOXES[0][1]],
        [SAW, 36, "6180553", None, BOXES[2][1]],
        [SAW, 20, "6180488", None, BOXES[1][1]],
    ]
    assert all(isinstance(row[0], str) for row in rows[1:]), "баркод строкой, иначе Excel покажет 2,05E+12"

    boxes = read_workbook(data)
    assert [(box.shk, box.lines[0].barcode, box.lines[0].quantity) for box in boxes] == [
        ("6180598", DRILL, 25),
        ("6180474", DRILL, 25),
        ("6180553", SAW, 36),
        ("6180488", SAW, 20),
    ]

    plan, data = pack(template, _assembly((DRILL, 50, "2-25")))
    assert data is None and "нужно 2" in plan.problems[0]
    with pytest.raises(StickerInputError, match="не открывается"):
        pack(b"not an excel", _assembly((DRILL, 50, "2-25")))


class FakeStore(PackingStore):
    def __init__(self) -> None:
        self.saved: dict[int, StoredPacking] = {}

    async def save(self, packing: StoredPacking) -> None:
        self.saved[packing.supply_id] = packing

    async def load(self, supply_id: int) -> StoredPacking | None:
        return self.saved.get(supply_id)

    async def recent(self) -> list[StoredPacking]:
        return sorted(self.saved.values(), key=lambda item: item.saved_at, reverse=True)


class SuppliesWithGoods(FakeSupplies):
    def __init__(self, goods: list[SupplyGood]) -> None:
        super().__init__()
        self.rows_goods = goods

    async def goods(self, seller_id: str, supply_id: int) -> list[SupplyGood]:
        return list(self.rows_goods)


def goods(*items: tuple[str, int]) -> list[SupplyGood]:
    return [SupplyGood(barcode, quantity, "", "0") for barcode, quantity in items]


def rows(*items: tuple[str, int, str], supply: int = STICKER_SUPPLY) -> str:
    return pasted(
        *(
            f"{supply}\t1\t{barcode}\t\tТовар\t{quantity}\t{quantity}\t{quantity}\t{scheme}"
            for barcode, quantity, scheme in items
        )
    )


async def test_packing_service_checks_rows_against_supply_goods() -> None:
    store = FakeStore()
    supplies = SuppliesWithGoods(goods((DRILL, 50), (SAW, 56)))
    service = PackingService(FakeSellers(), supplies, store)
    template = workbook_xlsx(excel_rows([0, 1, 2, 3], filled=False))

    plan = await service.preview(template, rows((DRILL, 50, "2-25"), (SAW, 56, "1-36+1-20")))
    assert plan.ready and plan.seller_name == "ИП Тестов Т.Т." and len(plan.boxes) == 4
    assert not store.saved, "предпросмотр ничего не запоминает"

    plan, data = await service.build(template, rows((DRILL, 50, "2-25"), (SAW, 56, "1-36+1-20")))
    assert data and store.saved[STICKER_SUPPLY].workbook == data
    assert store.saved[STICKER_SUPPLY].boxes == 4 and store.saved[STICKER_SUPPLY].seller_name == "ИП Тестов Т.Т."
    assert [item.supply_id for item in await service.recent()] == [STICKER_SUPPLY]

    wrong = await service.preview(template, rows((DRILL, 50, "2-25"), (LAMP, 56, "1-36+1-20")))
    assert not wrong.ready and wrong.boxes == []
    assert wrong.problems == [
        f"Этих баркодов нет в поставке в WB: {LAMP} — не те строки или не та поставка?",
        f"В поставке в WB заявлены товары, которых нет в строках сборки: {SAW}",
    ]
    less = await service.preview(template, rows((DRILL, 50, "2-25"), (SAW, 46, "1-36+1-10")))
    assert less.problems == [
        "Количество по коробам не совпадает с заявленным в поставке: "
        f"{SAW}: по коробам 46 шт, в поставке заявлено 56 шт"
    ]

    foreign = await service.preview(template, rows((DRILL, 50, "2-25"), (SAW, 56, "1-36+1-20"), supply=40000999))
    assert any("Поставка 40000999 не нашлась" in problem for problem in foreign.problems)
    with pytest.raises(StickerBuildError):
        await service.build(template, rows((DRILL, 50, "2-25")))


async def test_sticker_service_takes_remembered_workbook_by_supply() -> None:
    store = FakeStore()
    packing = PackingService(FakeSellers(), SuppliesWithGoods(goods((DRILL, 50), (SAW, 56))), store)
    template = workbook_xlsx(excel_rows([0, 1, 2, 3], filled=False))
    _, data = await packing.build(template, rows((DRILL, 50, "2-25"), (SAW, 56, "1-36+1-20")))

    stickers = BoxStickerService(FakeSellers(), FakeSupplies(), store)
    assert await stickers.workbook_for(STICKER_SUPPLY) == data
    with pytest.raises(StickerInputError, match="загрузите Excel файлом"):
        await stickers.workbook_for(40000999)
    assert [box.shk for box in read_workbook(data)] == ["6180474", "6180488", "6180553", "6180598"]


async def test_packing_routes_need_login(client: AsyncClient) -> None:
    files = {"template": ("a.xlsx", b"x")}
    response = await client.post("/api/v1/wb/box-stickers/packing/check", files=files, data={"assembly": "x"})
    assert response.status_code == 401
    response = await client.get("/api/v1/wb/box-stickers/packing/recent")
    assert response.status_code == 401
    response = await client.post("/api/v1/wb/box-stickers/check", files={"stickers": ("a.pdf", b"y")})
    assert response.status_code == 401
