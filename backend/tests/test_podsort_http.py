import io
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, date, datetime, timedelta

import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import delete, func, select

from backend.app.application import Application
from backend.modules.wb_core.application import MirrorService, RemainsMirror
from backend.modules.wb_core.domain import MIRROR_REMAINS, WarehouseRemain
from backend.modules.wb_core.infrastructure.postgres import MirrorRepository, SellerRepository
from backend.modules.wb_core.infrastructure.postgres.models import SellerModel
from backend.modules.wb_core.infrastructure.wb import (
    CatalogCard,
    WBAnalyticsClient,
    WBChatClient,
    WBContentClient,
    WBFeedbackClient,
    WBMarketplaceClient,
    WBWarehouseRemainsClient,
)
from backend.modules.wb_podsort.domain import CENTRAL, OrderLine
from backend.modules.wb_podsort.infrastructure.postgres import (
    OrderCountModel,
    PodsortRepository,
    SettingsModel,
    TrackedSellerModel,
    WarehouseRegionModel,
)
from backend.modules.wb_podsort.infrastructure.wb import WBPodsortStatisticsClient
from backend.shared.settings import load_settings
from backend.tests.egress_stub import make_gateway
from backend.workers.wb_podsort.worker import PodsortWorker

API = "/api/v1/wb/podsort"
AUTOMATION = "/api/v1/automations/wb-podsort/sellers"
SETTINGS = load_settings("backend/shared/settings/config.test.yaml")
DRILL, SAW = "2053497104469", "2055210050862"


class FakeStatistics(WBPodsortStatisticsClient):
    """Каждые сутки — девять шуруповёртов в ЦФО (четыре со склада продавца) и одна пила в Казахстан."""

    def __init__(self) -> None:
        super().__init__(make_gateway())
        self.days: list[date] = []
        self.sellers: set[str] = set()

    async def orders_on(self, seller_id: str, day: date) -> list[OrderLine]:
        self.days.append(day)
        self.sellers.add(seller_id)
        drill = [
            OrderLine(DRILL, 1271611253, "KARBI - Шуруповерт", "Шуруповерты", "0", CENTRAL, fbs=index < 4)
            for index in range(9)
        ]
        return [*drill, OrderLine(SAW, 1466514852, "KARBI - Пила", "Пилы", "0", "Казахстан", fbs=False)]


class FakeRemains(WBWarehouseRemainsClient):
    def __init__(self, remains: list[WarehouseRemain]) -> None:
        super().__init__(make_gateway())
        self.rows = remains

    async def remains(self, seller_id: str) -> list[WarehouseRemain]:
        return list(self.rows)


def own(body: dict) -> list[dict]:
    """Строки своего кабинета: в общей тестовой базе могут быть подключены и другие."""
    return [item for item in body["rows"] if item["seller_name"] == "Байбурин тест"]


def remain(barcode: str, warehouse: str, quantity: int) -> WarehouseRemain:
    return WarehouseRemain(barcode, "1271611253", "0", "KARBI - Шуруповерт желтый", warehouse, quantity)


@pytest_asyncio.fixture
async def application() -> AsyncIterator[Application]:
    application = Application(SETTINGS)
    app = application.get_app()
    async with app.router.lifespan_context(app):
        yield application


async def _reset_module(application: Application) -> None:
    async with application.database.session() as session:
        await session.execute(delete(SettingsModel))
        await session.execute(delete(WarehouseRegionModel))
        await session.commit()


@pytest_asyncio.fixture
async def seller(application: Application) -> AsyncIterator[uuid.UUID]:
    await _reset_module(application)
    async with application.database.session() as session:
        model = SellerModel(name="Байбурин тест", catalog_sync_status="success", egress_status="verified")
        session.add(model)
        await session.flush()
        seller_id = model.id
        await SellerRepository(session).upsert_catalog(
            seller_id,
            active=[
                CatalogCard(
                    article="1271611253",
                    vendor_code="KARBI - Шуруповерт желтый",
                    name="Шуруповерт",
                    subject_name="Шуруповерты",
                    sizes=[{"chrt_id": 1870528015, "tech_size": "0", "skus": [DRILL]}],
                )
            ],
            archived=[],
            archived_available=True,
        )
        await session.commit()
    mirror = MirrorService(
        application.database,
        content=WBContentClient(make_gateway()),
        analytics=WBAnalyticsClient(make_gateway()),
        marketplace=WBMarketplaceClient(make_gateway()),
        feedbacks=WBFeedbackClient(make_gateway()),
        chats=WBChatClient(make_gateway()),
        remains=FakeRemains(
            [
                remain(DRILL, "Коледино", 15),
                remain(DRILL, "Тула", 5),
                remain(DRILL, "Склад WB РФ", 100),
                remain(DRILL, "В пути до получателей", 40),
                remain(DRILL, "Неведомый склад", 7),
            ]
        ),
    )
    outcome = await mirror.collect_remains(seller_id)
    assert (outcome.barcodes, outcome.rows) == (1, 5)
    try:
        yield seller_id
    finally:
        async with application.database.session() as session:
            await PodsortRepository(session).purge_seller(seller_id)
            await session.execute(delete(SellerModel).where(SellerModel.id == seller_id))
            await session.commit()
        await _reset_module(application)


@pytest_asyncio.fixture
async def client(application: Application) -> AsyncIterator[AsyncClient]:
    token = application.token_service.issue_access(uuid.uuid4())
    async with AsyncClient(
        transport=ASGITransport(app=application.get_app()),
        base_url="http://test",
        headers={"Authorization": f"Bearer {token}"},
    ) as test_client:
        yield test_client


async def test_the_remains_mirror_keeps_every_warehouse(application: Application, seller: uuid.UUID) -> None:
    async with application.database.session() as session:
        snapshot = await RemainsMirror(session).snapshot(seller)
        state = await MirrorRepository(session).state(seller, MIRROR_REMAINS)

    assert state is not None and state.collected_at is not None
    assert snapshot.collected_at == state.collected_at
    assert sorted((item.warehouse_name, item.quantity) for item in snapshot.remains) == [
        ("В пути до получателей", 40),
        ("Коледино", 15),
        ("Неведомый склад", 7),
        ("Склад WB РФ", 100),
        ("Тула", 5),
    ]


async def test_podsort_counts_orders_minus_regional_stock(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    assert (await client.post(AUTOMATION, json={"seller_id": str(seller)})).status_code == 201
    statistics = FakeStatistics()
    worker = PodsortWorker(application.database, statistics, SETTINGS)

    # Первый проход — пять свежих суток, второй — ещё пять: окно в семь дней уже полное.
    assert await worker.collect(seller, datetime.now(UTC)) is None
    assert await worker.collect(seller, datetime.now(UTC)) is None
    yesterday = datetime.now(UTC).astimezone(worker.timezone).date() - timedelta(days=1)
    assert statistics.days[:3] == [yesterday, yesterday - timedelta(days=1), yesterday - timedelta(days=2)]
    assert len(statistics.days) == 10

    body = (await client.get(API)).json()
    assert body["region"] == CENTRAL
    [state] = [item for item in body["sellers"] if item["seller_id"] == str(seller)]
    assert state["window_days_loaded"] == 7 and state["history_days_loaded"] == 10
    assert state["remains_at"] is not None
    # 63 за неделю на неделю вперёд минус Коледино и Тула; «Склад WB РФ», в пути и неведомый склад не вычитаются.
    [row] = own(body)
    assert (row["barcode"], row["vendor_code"], row["window_orders"], row["stock"], row["need"]) == (
        DRILL,
        "KARBI - Шуруповерт желтый",
        63,
        20,
        43,
    )
    warehouses = {item["name"]: item for item in body["warehouses"]}
    assert warehouses["Коледино"]["region"] == CENTRAL and warehouses["Коледино"]["source"] == "guess"
    assert (warehouses["Неведомый склад"]["region"], warehouses["Неведомый склад"]["source"]) == (None, "none")
    assert warehouses["Склад WB РФ"]["source"] == "unplaced"
    assert body["warehouses"][0]["source"] == "none"
    assert "В пути до получателей" not in warehouses

    # Человек отнёс неведомый склад к ЦФО — его остаток тоже вычитается.
    response = await client.put(f"{API}/warehouses", json={"name": "Неведомый склад", "region": CENTRAL})
    assert response.status_code == 204
    [row] = own((await client.get(API)).json())
    assert (row["stock"], row["need"]) == (27, 36)
    # Коледино не относится ни к какому региону — и вернуть угадывание.
    await client.put(f"{API}/warehouses", json={"name": "Коледино", "region": None})
    assert own((await client.get(API)).json())[0]["stock"] == 12
    await client.put(f"{API}/warehouses", json={"name": "Коледино", "guess": True})
    assert own((await client.get(API)).json())[0]["stock"] == 27

    assert (await client.put(f"{API}/settings", json={"window_days": 10, "cover_days": 7})).status_code == 422
    assert (
        await client.put(f"{API}/settings", json={"window_days": 7, "cover_days": 7, "regions": ["Марс"]})
    ).status_code == 422
    saved = await client.put(
        f"{API}/settings", json={"window_days": 14, "cover_days": 14, "regions": [CENTRAL, "Казахстан"]}
    )
    assert saved.status_code == 200 and saved.json()["regions"] == [CENTRAL, "Казахстан"]
    # Окно 14 дней, загружено 10: 90 заказов / 14 × 14 − 27.
    [row] = own((await client.get(API)).json())
    assert (row["window_orders"], row["need"]) == (90, 63)
    kazakhstan = own((await client.get(API, params={"region": "Казахстан"})).json())
    assert [(item["barcode"], item["need"]) for item in kazakhstan] == [(SAW, 10)]
    assert (await client.get(API, params={"region": "Уральский"})).status_code == 422

    export = await client.get(f"{API}/export")
    assert export.status_code == 200
    book = load_workbook(io.BytesIO(export.content))
    summary = book.worksheets[0]
    assert summary.title.startswith("Подсорт ")
    notes = " ".join(str(summary.cell(line, 1).value or "") for line in range(1, 12))
    assert "Остаток WB в регионе вычтен" in notes
    assert "Байбурин тест: заказы загружены за 10 из 14" in notes
    assert "Подсорт Байбурин тест" in book.sheetnames


async def test_disconnecting_keeps_the_loaded_days(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    assert (await client.post(AUTOMATION, json={"seller_id": str(seller)})).status_code == 201
    assert (
        await PodsortWorker(application.database, FakeStatistics(), SETTINGS).collect(seller, datetime.now(UTC)) is None
    )

    assert (await client.delete(f"{AUTOMATION}/{seller}")).status_code == 204

    async with application.database.session() as session:
        tracked = await session.get(TrackedSellerModel, seller)
        counts = await session.scalar(
            select(func.count()).select_from(OrderCountModel).where(OrderCountModel.seller_id == seller)
        )
    # Отключение — не удаление: подключили снова — догрузка продолжается, а не начинается заново.
    assert tracked is None
    assert counts == 10
    assert (await client.post(AUTOMATION, json={"seller_id": str(seller)})).status_code == 201
    statistics = FakeStatistics()
    await PodsortWorker(application.database, statistics, SETTINGS).collect(seller, datetime.now(UTC))
    assert len(statistics.days) == 5
    assert min(statistics.days) < datetime.now(UTC).date() - timedelta(days=6)


async def test_a_failing_cabinet_waits_before_the_next_attempt(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    assert (await client.post(AUTOMATION, json={"seller_id": str(seller)})).status_code == 201
    async with application.database.session() as session:
        podsort = PodsortRepository(session)
        await podsort.record_attempt(seller, now=datetime.now(UTC))
        await podsort.fail_collection(seller, "WB Statistics API отвечает HTTP 429")
        await session.commit()
    statistics = FakeStatistics()
    worker = PodsortWorker(application.database, statistics, SETTINGS)

    await worker.tick()
    assert str(seller) not in statistics.sellers

    # Пауза прошла — кабинет снова в очереди, удачный проход снимает ошибку.
    later = PodsortWorker(
        application.database,
        statistics,
        SETTINGS,
        now=lambda: datetime.now(UTC) + timedelta(minutes=SETTINGS.podsort.retry_minutes + 1),
    )
    await later.tick()
    assert str(seller) in statistics.sellers
    async with application.database.session() as session:
        tracked = await session.get(TrackedSellerModel, seller)
    assert tracked is not None and tracked.collection_error is None
