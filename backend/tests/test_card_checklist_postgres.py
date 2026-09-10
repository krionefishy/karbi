import io
import uuid
from collections.abc import AsyncIterator
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.modules.wb_card_checklist.application import (
    REVIEWS_NOT_CONNECTED,
    REVIEWS_OK,
    STOCK_NOT_CONNECTED,
    STOCK_OK,
    STOCK_STALE,
    ArticleNotInChecklistError,
    ChecklistService,
    CollectionService,
)
from backend.modules.wb_card_checklist.domain import CardFacts, PriceFacts, SubjectCharacteristic, Thresholds
from backend.modules.wb_card_checklist.infrastructure.postgres import (
    ChecklistRepository,
    SubjectCharacteristicsModel,
    TrackedSellerModel,
)
from backend.modules.wb_card_checklist.infrastructure.wb import WBCardClient, WBPricesClient
from backend.modules.wb_core.infrastructure.postgres import SellerRepository
from backend.modules.wb_core.infrastructure.postgres.models import SellerModel
from backend.modules.wb_core.infrastructure.wb import WBPermanentError
from backend.modules.wb_reviews.application import ReviewTotalsReader
from backend.modules.wb_reviews.infrastructure.postgres import ReviewSyncRepository
from backend.modules.wb_turnover.application import CurrentStockReader
from backend.modules.wb_turnover.infrastructure.postgres import TurnoverRepository
from backend.shared.settings import load_settings
from backend.storage.pg import Database
from backend.tests.egress_stub import make_gateway
from backend.workers.wb_card_checklist.worker import CardChecklistWorker

SETTINGS = load_settings("backend/shared/settings/config.test.yaml")
MOSCOW = ZoneInfo("Europe/Moscow")
# Предмет, которого нет в настоящем справочнике WB: таблица предметов общая,
# и тест не должен задеть чужие строки.
SUBJECT = 990_001
DIRECTORY = [
    SubjectCharacteristic(1, "Цвет", required=False, popular=True, named_field=False),
    SubjectCharacteristic(2, "Материал", required=False, popular=False, named_field=False),
    SubjectCharacteristic(3, "Мощность", required=False, popular=False, named_field=False),
]


@dataclass(frozen=True)
class Totals:
    total: int
    with_photo: int | None
    with_video: int | None


class FakeStock:
    def __init__(self, value: dict[str, int] | None) -> None:
        self.value = value

    async def stock(self, seller_id: uuid.UUID, today: date) -> dict[str, int] | None:
        return self.value


class FakeReviews:
    def __init__(self, value: dict[str, Totals] | None) -> None:
        self.value = value

    async def totals(self, seller_id: uuid.UUID) -> dict[str, Totals] | None:
        return self.value


class FakeCards(WBCardClient):
    def __init__(self, cards: list[CardFacts]) -> None:
        super().__init__(make_gateway())
        self.served = cards
        self.subjects_asked: list[int] = []

    async def cards(self, seller_id: str) -> list[CardFacts]:
        return list(self.served)

    async def characteristics(self, seller_id: str, subject_id: int) -> list[SubjectCharacteristic]:
        self.subjects_asked.append(subject_id)
        return list(DIRECTORY)


class FakePrices(WBPricesClient):
    def __init__(self, prices: list[PriceFacts] | None = None, error: Exception | None = None) -> None:
        super().__init__(make_gateway())
        self.served = prices or []
        self.error = error

    async def prices(self, seller_id: str) -> list[PriceFacts]:
        if self.error is not None:
            raise self.error
        return list(self.served)


def card(article: str, *, title: str, imt_id: int = 1, photos: int = 5, description: int = 1500) -> CardFacts:
    return CardFacts(
        article=article,
        vendor_code=f"SKU-{article}",
        title=title,
        barcode=f"20{article}",
        imt_id=imt_id,
        subject_id=SUBJECT,
        subject_name="Пилы",
        photo_url="",
        photo_count=photos,
        description_length=description,
        has_video=True,
        characteristic_ids=frozenset({1, 2, 3}),
        card_created_at=datetime(2026, 8, 1, 9, 0, tzinfo=UTC),
    )


CARDS = [
    card("A", title="Бета", imt_id=1),
    card("B", title="Альфа без остатка", imt_id=2),
    card("C", title="Альфа", imt_id=1, photos=1),
]


@pytest_asyncio.fixture
async def database() -> AsyncIterator[Database]:
    database = Database()
    await database.connect(SETTINGS.database.url, pool_size=2, max_overflow=0)
    try:
        yield database
    finally:
        await database.disconnect()


@pytest_asyncio.fixture
async def seller(database: Database) -> AsyncIterator[uuid.UUID]:
    async with database.session() as session:
        model = SellerModel(name="ИП Чек-лист", catalog_sync_status="success")
        session.add(model)
        await session.flush()
        seller_id = model.id
        await ChecklistRepository(session).track(seller_id)
        await session.commit()
    try:
        yield seller_id
    finally:
        async with database.session() as session:
            await ChecklistRepository(session).purge_seller(seller_id)
            await ReviewSyncRepository(session).purge_seller(seller_id)
            await TurnoverRepository(session).purge_seller(seller_id)
            await session.execute(
                delete(SubjectCharacteristicsModel).where(SubjectCharacteristicsModel.subject_id == SUBJECT)
            )
            await session.execute(delete(SellerModel).where(SellerModel.id == seller_id))
            await session.commit()


async def seed(database: Database, seller_id: uuid.UUID) -> None:
    async with database.session() as session:
        checklist = ChecklistRepository(session)
        await checklist.replace_cards(seller_id, CARDS)
        await checklist.replace_prices(seller_id, [PriceFacts("A", 1000, 500, 50, 0)])
        await checklist.save_subject(SUBJECT, DIRECTORY)
        await session.commit()


def service(
    session: AsyncSession,
    stock: dict[str, int] | None = None,
    reviews: dict[str, Totals] | None = None,
) -> ChecklistService:
    return ChecklistService(
        session,
        SellerRepository(session),
        ChecklistRepository(session),
        FakeStock({"A": 15, "B": 5, "C": 10} if stock is None else stock),
        FakeReviews({"A": Totals(3, 1, 0), "C": Totals(2, None, None)} if reviews is None else reviews),
        thresholds=Thresholds(),
        min_stock=10,
        timezone=MOSCOW,
    )


async def test_the_table_takes_cards_with_stock_and_counts_reviews_per_card(
    database: Database, seller: uuid.UUID
) -> None:
    await seed(database, seller)

    async with database.session() as session:
        view = await service(session).view(seller)

    assert (view.stock_state, view.reviews_state, view.min_stock) == (STOCK_OK, REVIEWS_OK, 10)
    # B лежит ниже порога в 10 штук и в таблицу не попадает; порядок — по названию.
    assert [row.article for row in view.rows] == ["C", "A"]
    alpha, beta = view.rows
    items = {state.key: state for state in beta.items}
    assert beta.stock == 15
    assert items["photos"].done
    assert (items["characteristics"].done, items["characteristics"].detail) == (True, "3/3")
    # Отзывы у склейки общие: 3 у A и 2 у C — на карточке покупатель видит 5.
    assert items["reviews_present"].detail == "5 отз."
    assert (items["reviews_with_photo"].done, items["reviews_with_photo"].detail) == (True, "1 с фото")
    assert (items["reviews_with_video"].done, items["reviews_with_video"].detail) == (False, "0 с видео")
    assert (items["discount"].done, items["discount"].detail) == (True, "−50% · 500 ₽")
    # Всё, кроме видео-отзывов: 7 из 8.
    assert (beta.done, beta.total, beta.ready) == (7, 8, False)
    alpha_items = {state.key: state for state in alpha.items}
    # Одного фото по инструкции хватает.
    assert (alpha_items["photos"].done, alpha_items["photos"].detail) == (True, "1 фото")
    assert alpha_items["discount"].done is None


async def test_comments_stick_to_cards_in_the_table(database: Database, seller: uuid.UUID) -> None:
    await seed(database, seller)
    user = uuid.uuid4()

    async with database.session() as session:
        checklist = service(session)
        await checklist.set_comment(seller, "A", "Ждём видео от подрядчика", user)
        with pytest.raises(ArticleNotInChecklistError):
            await checklist.set_comment(seller, "B", "ниже порога", user)

    async with database.session() as session:
        view = await service(session).view(seller)
    assert next(row for row in view.rows if row.article == "A").comment == "Ждём видео от подрядчика"

    async with database.session() as session:
        await service(session).set_comment(seller, "A", "   ", user)
        assert (await ChecklistRepository(session).comments(seller)) == {}


async def test_without_neighbours_the_table_says_why(database: Database, seller: uuid.UUID) -> None:
    await seed(database, seller)

    async with database.session() as session:
        detached = await ChecklistService(
            session,
            SellerRepository(session),
            ChecklistRepository(session),
            FakeStock(None),
            FakeReviews(None),
            thresholds=Thresholds(),
            min_stock=10,
            timezone=MOSCOW,
        ).view(seller)
        stale = await service(session, stock={}).view(seller)

    assert (detached.stock_state, detached.reviews_state, detached.rows) == (
        STOCK_NOT_CONNECTED,
        REVIEWS_NOT_CONNECTED,
        (),
    )
    assert stale.stock_state == STOCK_STALE


async def test_export_follows_the_managers_spreadsheet(database: Database, seller: uuid.UUID) -> None:
    await seed(database, seller)
    async with database.session() as session:
        checklist = service(session)
        await checklist.set_comment(seller, "A", "Проверить рич", None)
        report = await checklist.export(seller)

    workbook = load_workbook(io.BytesIO(report.content))
    assert workbook.sheetnames == ["Чек-лист", "Инструкция"]
    sheet = workbook["Чек-лист"]
    header = [cell.value for cell in sheet[1]]
    assert header[:6] == ["Дата заведения", "Артикул WB", "Артикул продавца", "Баркод", "Наименование товара", "ИП"]
    assert header[6:14] == [
        "Описание (SEO)",
        "Характеристики",
        "Фото",
        "Видео",
        "Скидка / СПП",
        "Отзывы есть",
        "Отзывы с фото",
        "Отзывы с видео",
    ]
    assert header[14:] == ["Готово из 8", "Статус", "Комментарий"]
    second = [cell.value for cell in sheet[3]]
    assert second[1:6] == ["A", "SKU-A", "20A", "Бета", "ИП Чек-лист"]
    # В ячейке — то, что видит WB, результат — цветом.
    assert second[6:14] == ["1500 симв.", "3/3", "5 фото", "есть", "−50% · 500 ₽", "5 отз.", "1 с фото", "0 с видео"]
    fills = [sheet.cell(row=3, column=column).fill.fgColor.rgb[-6:] for column in range(7, 15)]
    assert fills == ["C6EFCE"] * 7 + ["FFC7CE"]
    assert second[14:17] == [7, "НЕ ГОТОВ", "Проверить рич"]
    assert sheet["P3"].fill.fgColor.rgb[-6:] == "FFC7CE"
    assert sheet["A3"].value.date() == date(2026, 8, 1)
    assert sheet.freeze_panes == "G2"
    assert report.filename.startswith("checklist_ИП-Чек-лист_")


async def test_collection_rewrites_cards_and_keeps_prices_a_key_cannot_read(
    database: Database, seller: uuid.UUID
) -> None:
    cards = FakeCards(CARDS)
    async with database.session() as session:
        first = await CollectionService(
            session,
            ChecklistRepository(session),
            cards,
            FakePrices([PriceFacts("A", 1000, 500, 50, 0)]),
            subject_ttl=timedelta(days=7),
        ).collect(seller)

    assert (first.cards, first.prices, first.subjects, first.warning) == (3, 1, 1, None)

    cards.served = [CARDS[0]]
    async with database.session() as session:
        second = await CollectionService(
            session,
            ChecklistRepository(session),
            cards,
            FakePrices(error=WBPermanentError("ключ не имеет доступа к категории «Цены и скидки»")),
            subject_ttl=timedelta(days=7),
        ).collect(seller)

    # Справочник свежий — второй раз его не спрашивают.
    assert cards.subjects_asked == [SUBJECT]
    assert second.subjects == 0
    assert second.warning is not None and second.warning.startswith("Цены не прочитаны")
    async with database.session() as session:
        checklist = ChecklistRepository(session)
        assert [item.article for item in await checklist.card_facts(seller)] == ["A"]
        assert set(await checklist.prices(seller)) == {"A"}
        tracked = await checklist.tracked(seller)
        assert tracked is not None and tracked.collected_at is not None
        assert tracked.collection_error == second.warning


async def test_a_seller_detached_mid_collection_is_not_written(database: Database, seller: uuid.UUID) -> None:
    async with database.session() as session:
        await ChecklistRepository(session).untrack(seller)
        await session.commit()

    async with database.session() as session:
        result = await CollectionService(
            session, ChecklistRepository(session), FakeCards(CARDS), FakePrices(), subject_ttl=timedelta(days=7)
        ).collect(seller)
        assert result.skipped
        assert await ChecklistRepository(session).card_facts(seller) == []


async def test_worker_collects_once_a_day_and_serves_the_button(database: Database, seller: uuid.UUID) -> None:
    # 06:00 МСК — после суточного сбора в 05:00.
    moment = datetime(2026, 9, 10, 6, 0, tzinfo=MOSCOW)
    worker = CardChecklistWorker(database, FakeCards(CARDS), FakePrices(), SETTINGS, now=lambda timezone: moment)

    assert worker.due_since(moment) == datetime(2026, 9, 10, 2, 0, tzinfo=UTC)
    # До пяти утра срок ещё вчерашний.
    assert worker.due_since(moment.replace(hour=4)) == datetime(2026, 9, 9, 2, 0, tzinfo=UTC)

    await worker.collect_due(moment)
    async with database.session() as session:
        tracked = await ChecklistRepository(session).tracked(seller)
        assert tracked is not None and tracked.collected_at is not None
        collected_at = tracked.collected_at
        due = await ChecklistRepository(session).sellers_due(
            worker.due_since(moment), retry_after=moment.astimezone(UTC)
        )
    assert seller not in due

    async with database.session() as session:
        checklist = service(session)
        first = await checklist.request_refresh(seller, None)
        again = await checklist.request_refresh(seller, None)
    assert first.requested_at == again.requested_at

    await worker.serve_refresh_requests()
    async with database.session() as session:
        state = await service(session).refresh_state(seller)
        tracked = await session.scalar(select(TrackedSellerModel).where(TrackedSellerModel.seller_id == seller))
    assert state is not None and state.status == "success"
    assert tracked is not None and tracked.collected_at is not None and tracked.collected_at >= collected_at


async def test_neighbours_hand_over_stock_and_reviews_only_to_their_own(database: Database, seller: uuid.UUID) -> None:
    today = date(2026, 9, 10)
    async with database.session() as session:
        stock = CurrentStockReader(TurnoverRepository(session))
        totals = ReviewTotalsReader(ReviewSyncRepository(session))
        assert await stock.stock(seller, today) is None
        assert await totals.totals(seller) is None

        turnover = TurnoverRepository(session)
        await turnover.track(seller)
        await turnover.upsert_snapshots(seller, today - timedelta(days=5), 0, "fbo", {"A": (99, 99, 0, 0)})
        reviews = ReviewSyncRepository(session)
        await reviews.track(seller)
        await session.commit()
        # Снимок пятидневной давности — это не «текущий остаток».
        assert await stock.stock(seller, today) == {}
        assert await totals.totals(seller) == {}

        await turnover.upsert_snapshots(seller, today, 1, "fbo", {"A": (7, 7, 0, 0)})
        await turnover.upsert_snapshots(seller, today, 1, "fbs", {"A": (5, 5, 0, 0)})
        await reviews.upsert_daily_counts(seller, today - timedelta(days=1), {"A": (0, 0, 0, 1, 2)}, {"A": (1, 1)})
        await reviews.upsert_daily_counts(seller, today, {"A": (0, 0, 1, 1, 2)})
        await session.commit()

        assert await stock.stock(seller, today) == {"A": 12}
        latest = (await totals.totals(seller) or {})["A"]
    # Последний срез снят без подсчёта медиа: «не знаем», а не ноль.
    assert (latest.date, latest.total, latest.with_photo, latest.with_video) == (today, 4, None, None)


async def test_a_refresh_left_running_by_a_dead_worker_is_closed(database: Database, seller: uuid.UUID) -> None:
    """Иначе кнопка «Обновить» у селлера навсегда висит на мёртвом запросе."""
    async with database.session() as session:
        await service(session).request_refresh(seller, None)
    async with database.session() as session:
        claimed = await ChecklistRepository(session).claim_refreshes()
        assert [request.seller_id for request in claimed] == [seller]
        claimed[0].started_at = datetime.now(UTC) - timedelta(hours=2)
        await session.commit()

    worker = CardChecklistWorker(database, FakeCards(CARDS), FakePrices(), SETTINGS)
    await worker.serve_refresh_requests()

    async with database.session() as session:
        state = await service(session).refresh_state(seller)
        assert state is not None and (state.status, state.in_progress) == ("error", False)
        # Место освободилось: новое нажатие создаёт новый запрос.
        fresh = await service(session).request_refresh(seller, None)
    assert fresh.status == "queued"
