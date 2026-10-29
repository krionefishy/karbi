import io
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import delete, update

from backend.app.factory import Application
from backend.modules.wb_core.infrastructure.postgres import SellerRepository
from backend.modules.wb_core.infrastructure.postgres.models import SellerModel
from backend.modules.wb_core.infrastructure.wb import CatalogCard
from backend.modules.wb_turnover.infrastructure.postgres import (
    BoardItemModel,
    RefreshRequestModel,
    TrackedSellerModel,
    TurnoverRepository,
)
from backend.modules.wb_turnover.infrastructure.wb import OrderRow, WBStatisticsClient
from backend.shared.settings import load_settings
from backend.tests.egress_stub import make_gateway
from backend.tests.operators import operator_token

API = "/api/v1/wb/turnover/board"
SETTINGS = load_settings("backend/shared/settings/config.test.yaml")
NAME = "Таблица тест"
DRILL, SAW, GHOST = "2053497104469", "2055210050862", "2000000000001"
DRILL_CHRT, SAW_CHRT = 1870528015, 1870528016


def own(body: dict) -> dict[str, dict]:
    return {row["barcode"]: row for row in body["rows"] if row["seller_name"] == NAME}


def sale(srid: str, barcode: str, article: str, day: date, *, cancelled: bool = False) -> dict:
    return {
        "srid": srid,
        "article": article,
        "barcode": barcode,
        "order_date": day,
        "last_change_date": datetime.combine(day, datetime.min.time(), tzinfo=UTC),
        "is_cancel": cancelled,
        "price": 1000,
        "warehouse_type": "Склад WB",
    }


@pytest_asyncio.fixture
async def application() -> AsyncIterator[Application]:
    application = Application(SETTINGS)
    app = application.get_app()
    async with app.router.lifespan_context(app):
        yield application


@pytest_asyncio.fixture
async def seller(application: Application) -> AsyncIterator[uuid.UUID]:
    async with application.database.session() as session:
        model = SellerModel(name=NAME, catalog_sync_status="success", egress_status="verified")
        session.add(model)
        await session.flush()
        seller_id = model.id
        await SellerRepository(session).upsert_catalog(
            seller_id,
            active=[
                CatalogCard(
                    article="1271611253",
                    vendor_code="KARBI - Шуруповерт",
                    name="Шуруповерт",
                    subject_name="Шуруповерты",
                    sizes=[{"chrt_id": 1870528015, "tech_size": "0", "skus": [DRILL]}],
                ),
                CatalogCard(
                    article="1466514852",
                    vendor_code="KARBI - Пила",
                    name="Пила",
                    subject_name="Пилы",
                    sizes=[{"chrt_id": 1870528016, "tech_size": "0", "skus": [SAW]}],
                ),
            ],
            archived=[],
            archived_available=True,
        )
        session.add(TrackedSellerModel(seller_id=seller_id))
        await session.commit()
    try:
        yield seller_id
    finally:
        async with application.database.session() as session:
            await TurnoverRepository(session).purge_seller(seller_id)
            await session.execute(delete(RefreshRequestModel).where(RefreshRequestModel.seller_id == seller_id))
            await session.execute(delete(SellerModel).where(SellerModel.id == seller_id))
            await session.commit()


@pytest_asyncio.fixture
async def client(application: Application) -> AsyncIterator[AsyncClient]:
    token = await operator_token(application)
    async with AsyncClient(
        transport=ASGITransport(app=application.get_app()),
        base_url="http://test",
        headers={"Authorization": f"Bearer {token}"},
    ) as test_client:
        yield test_client


def today() -> date:
    return datetime.now(ZoneInfo("Europe/Moscow")).date()


def yesterday() -> date:
    return today() - timedelta(days=1)


async def seed_orders(application: Application, seller_id: uuid.UUID, rows: list[dict]) -> None:
    async with application.database.session() as session:
        await TurnoverRepository(session).upsert_orders(seller_id, rows)
        await session.commit()


async def seed_stock(
    application: Application, seller_id: uuid.UUID, by_size: dict[int, int], at: datetime | None = None
) -> None:
    async with application.database.session() as session:
        repository = TurnoverRepository(session)
        await repository.replace_fbo_stocks(seller_id, by_size)
        if at is not None:
            await session.execute(
                update(TrackedSellerModel).where(TrackedSellerModel.seller_id == seller_id).values(fbo_stocks_at=at)
            )
        await session.commit()


async def add(client: AsyncClient, seller_id: uuid.UUID, *barcodes: str):
    return await client.post(f"{API}/sellers/{seller_id}/items", json={"barcodes": list(barcodes)})


async def test_a_row_shows_orders_of_the_week_the_wb_warehouse_stock_and_the_days_it_lasts(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    day = yesterday()
    orders = [sale(f"d{index}", DRILL, "1271611253", day) for index in range(14)]
    orders += [sale("old", DRILL, "1271611253", day - timedelta(days=7))]
    orders += [sale("c1", DRILL, "1271611253", day, cancelled=True)]
    await seed_orders(application, seller, orders)
    await seed_stock(application, seller, {DRILL_CHRT: 30, SAW_CHRT: 500})
    assert (await add(client, seller, DRILL)).json() == {"added": 1, "existing": 0}

    row = own((await client.get(API)).json())[DRILL]

    assert (row["article"], row["vendor_code"], row["name"]) == ("1271611253", "KARBI - Шуруповерт", "Шуруповерт")
    assert (row["orders"], row["stock"], row["rate"], row["days"], row["status"]) == (15, 30, 2.1, 14, "ok")
    assert row["tone"] == "warning"


async def test_orders_outside_the_window_and_today_are_not_counted(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    now = today()
    await seed_orders(
        application,
        seller,
        [
            sale("today", DRILL, "1271611253", now),
            sale("edge", DRILL, "1271611253", now - timedelta(days=7)),
            sale("before", DRILL, "1271611253", now - timedelta(days=8)),
        ],
    )
    await add(client, seller, DRILL)
    body = (await client.get(API)).json()
    assert own(body)[DRILL]["orders"] == 1
    assert (body["window_from"], body["window_to"]) == (
        (now - timedelta(days=7)).isoformat(),
        (now - timedelta(days=1)).isoformat(),
    )


async def test_without_a_stock_snapshot_the_days_are_unknown_not_zero(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    await seed_orders(application, seller, [sale("a", SAW, "1466514852", yesterday())])
    await add(client, seller, SAW)

    row = own((await client.get(API)).json())[SAW]

    assert (row["stock"], row["days"], row["status"]) == (None, None, "no_data")


async def test_an_article_without_orders_is_no_sales(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    await seed_stock(application, seller, {SAW_CHRT: 12})
    await add(client, seller, SAW)

    row = own((await client.get(API)).json())[SAW]

    assert (row["stock"], row["orders"], row["days"], row["status"]) == (12, 0, None, "no_sales")


async def test_fbs_is_yes_no_or_unknown(application: Application, seller: uuid.UUID, client: AsyncClient) -> None:
    await seed_stock(application, seller, {})
    await add(client, seller, DRILL, SAW)
    assert {row["fbs"] for row in own((await client.get(API)).json()).values()} == {"unknown"}

    async with application.database.session() as session:
        repository = TurnoverRepository(session)
        await repository.replace_warehouses(seller, [(7, "4.  Питер наш"), (8, "ФФ Казань")])
        await repository.replace_warehouse_stocks(seller, {("1271611253", 7): 4, ("1466514852", 8): 9})
        await session.commit()
    rows = own((await client.get(API)).json())
    assert (rows[DRILL]["fbs"], rows[SAW]["fbs"]) == ("yes", "no")


async def test_fbs_counts_only_the_chosen_warehouses_and_the_list_can_grow(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    await seed_stock(application, seller, {})
    await add(client, seller, DRILL, SAW)
    async with application.database.session() as session:
        repository = TurnoverRepository(session)
        await repository.replace_warehouses(seller, [(7, "4. Питер наш"), (8, "ФФ Казань")])
        await repository.replace_warehouse_stocks(seller, {("1271611253", 7): 4, ("1466514852", 8): 9})
        await session.commit()

    options = (await client.get(f"{API}/fbs-warehouses")).json()
    by_label = {item["label"]: item for item in options}
    assert by_label["4. Питер наш"]["selected"] is True
    assert by_label["ФФ Казань"]["selected"] is False

    try:
        keys = [item["key"] for item in options if item["selected"]] + [by_label["ФФ Казань"]["key"]]
        assert (await client.put(f"{API}/fbs-warehouses", json={"keys": keys})).status_code == 200
        rows = own((await client.get(API)).json())
        assert (rows[DRILL]["fbs"], rows[SAW]["fbs"]) == ("yes", "yes")
    finally:
        async with application.database.session() as session:
            await TurnoverRepository(session).set_board_fbs_warehouses(
                {item["key"]: item["label"] for item in options if item["selected"]}
            )
            await session.commit()


async def test_a_barcode_missing_from_the_catalog_still_gets_a_row(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    await seed_orders(application, seller, [sale("g", GHOST, "999", yesterday())])
    await add(client, seller, GHOST)

    row = own((await client.get(API)).json())[GHOST]

    assert (row["in_catalog"], row["article"], row["orders"]) == (False, "999", 1)


async def test_orders_collected_before_barcodes_were_stored_are_reported(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    await seed_orders(application, seller, [sale("legacy", "", "1271611253", yesterday())])
    body = (await client.get(API)).json()
    [state] = [item for item in body["sellers"] if item["seller_id"] == str(seller)]
    assert state["orders_without_barcode"] == 1


async def test_a_stale_stock_layer_is_flagged(application: Application, seller: uuid.UUID, client: AsyncClient) -> None:
    await seed_stock(application, seller, {}, at=datetime.now(UTC) - timedelta(hours=30))
    body = (await client.get(API)).json()
    [state] = [item for item in body["sellers"] if item["seller_id"] == str(seller)]
    assert state["stock_stale"] is True


async def test_the_list_is_edited_by_the_analyst(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    assert (await add(client, seller, DRILL, DRILL, SAW)).json() == {"added": 2, "existing": 0}
    assert (await add(client, seller, SAW)).json() == {"added": 0, "existing": 1}
    assert (await add(client, seller, "не баркод!")).status_code == 409

    patched = await client.patch(
        f"{API}/sellers/{seller}/items/{DRILL}", json={"note": "включить FBS", "nomenclature": "Дрель 1С"}
    )
    assert patched.status_code == 204
    row = own((await client.get(API)).json())[DRILL]
    assert (row["note"], row["nomenclature"]) == ("включить FBS", "Дрель 1С")

    assert (await client.put(f"{API}/sellers/{seller}/items/{DRILL}/hidden", json={"hidden": True})).status_code == 204
    body = (await client.get(API)).json()
    assert DRILL not in own(body)
    assert [item["barcode"] for item in body["hidden"] if item["seller_name"] == NAME] == [DRILL]

    await client.put(f"{API}/sellers/{seller}/items/{DRILL}/hidden", json={"hidden": False})
    assert own((await client.get(API)).json())[DRILL]["note"] == "включить FBS"

    assert (await client.delete(f"{API}/sellers/{seller}/items/{DRILL}")).status_code == 204
    assert DRILL not in own((await client.get(API)).json())
    assert (await client.delete(f"{API}/sellers/{seller}/items/{DRILL}")).status_code == 404


async def test_a_seller_outside_the_automation_cannot_get_items(application: Application, client: AsyncClient) -> None:
    async with application.database.session() as session:
        model = SellerModel(name="Не подключён", catalog_sync_status="success", egress_status="verified")
        session.add(model)
        await session.commit()
        stranger = model.id
    try:
        assert (await add(client, stranger, DRILL)).status_code == 404
    finally:
        async with application.database.session() as session:
            await session.execute(delete(SellerModel).where(SellerModel.id == stranger))
            await session.commit()


async def test_refresh_queues_collection_for_sellers_with_rows(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    await add(client, seller, DRILL)
    first = (await client.post(f"{API}/refresh")).json()["queued"]
    second = (await client.post(f"{API}/refresh")).json()["queued"]
    assert first >= 1 and second == first
    async with application.database.session() as session:
        assert (await TurnoverRepository(session).active_refresh(seller)) is not None


async def test_the_workbook_has_the_analysts_columns_and_live_formulas(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    await seed_orders(application, seller, [sale(f"s{i}", DRILL, "1271611253", yesterday()) for i in range(14)])
    await seed_stock(application, seller, {DRILL_CHRT: 30})
    await add(client, seller, DRILL)

    response = await client.get(f"{API}/export")

    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content))["Оборачиваемость"]
    header = [cell.value for cell in sheet[1]]
    assert header == [
        "Хватит на, дн.",
        "Комментарии",
        "Кабинет",
        "Артикул WB",
        "Баркод",
        "Артикул продавца",
        "Номенклатура 1С",
        "Остатки ФБО (Склад ВБ РФ)",
        "Продажи за 7 дней",
        "Средние продажи в день",
        "Остатки ФБС",
    ]
    rows = [dict(zip(header, [cell.value for cell in row], strict=True)) for row in sheet.iter_rows(min_row=2)]
    [row] = [item for item in rows if item["Кабинет"] == NAME]
    assert (row["Баркод"], row["Остатки ФБО (Склад ВБ РФ)"], row["Продажи за 7 дней"]) == (DRILL, 30, 14)
    assert str(row["Средние продажи в день"]).startswith("=I")
    assert "ROUNDDOWN" in str(row["Хватит на, дн."])


async def test_deleting_the_seller_with_purge_drops_its_rows(application: Application, seller: uuid.UUID) -> None:
    async with application.database.session() as session:
        repository = TurnoverRepository(session)
        await repository.add_board_items(seller, [DRILL], None)
        await repository.purge_seller(seller)
        await session.commit()
        assert await repository.board_items([seller]) == []
    async with application.database.session() as session:
        await session.execute(delete(BoardItemModel).where(BoardItemModel.seller_id == seller))


def test_the_statistics_client_reads_the_barcode_of_an_order() -> None:
    row = WBStatisticsClient(make_gateway())._order(
        {
            "nmId": 1271611253,
            "srid": "x",
            "barcode": DRILL,
            "date": "2026-10-21T10:00:00",
            "lastChangeDate": "2026-10-21T11:00:00",
        }
    )
    assert isinstance(row, OrderRow)
    assert row.barcode == DRILL


async def test_each_size_of_a_card_has_its_own_stock(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    """У размеров одной карточки остатки разные; артикульная сумма дала бы всем одну цифру."""
    async with application.database.session() as session:
        await SellerRepository(session).upsert_catalog(
            seller,
            active=[
                CatalogCard(
                    article="1270964388",
                    vendor_code="Дождевики",
                    name="Дождевики для обуви",
                    subject_name="Дождевики",
                    sizes=[
                        {"chrt_id": 500, "tech_size": "M", "skus": ["2053485305847"]},
                        {"chrt_id": 501, "tech_size": "L", "skus": ["2053486133913"]},
                    ],
                )
            ],
            archived=[],
            archived_available=True,
        )
        await session.commit()
    await seed_stock(application, seller, {500: 40, 501: 3})
    await add(client, seller, "2053485305847", "2053486133913")

    rows = own((await client.get(API)).json())

    assert (rows["2053485305847"]["stock"], rows["2053486133913"]["stock"]) == (40, 3)


async def test_a_size_missing_from_the_wb_report_is_an_empty_shelf_once_the_layer_exists(
    application: Application, seller: uuid.UUID, client: AsyncClient
) -> None:
    await seed_stock(application, seller, {DRILL_CHRT: 5})
    await add(client, seller, SAW)

    assert own((await client.get(API)).json())[SAW]["stock"] == 0


async def test_collecting_stocks_stores_the_layer_per_size(application: Application, seller: uuid.UUID) -> None:
    from backend.modules.wb_turnover.application import CollectionService
    from backend.modules.wb_turnover.infrastructure.wb import FBOStockRow, WBAnalyticsClient, WBMarketplaceClient

    class Analytics(WBAnalyticsClient):
        async def stocks(self, seller_id: str) -> list[FBOStockRow]:
            return [
                FBOStockRow("1271611253", DRILL_CHRT, -999999, "Склад WB", "Склад WB", 12, 1, 2),
                FBOStockRow("1466514852", SAW_CHRT, -999999, "Склад WB", "Склад WB", 7, 0, 0),
            ]

    async with application.database.session() as session:
        await CollectionService(
            session,
            SellerRepository(session),
            TurnoverRepository(session),
            WBStatisticsClient(make_gateway()),
            Analytics(make_gateway()),
            WBMarketplaceClient(make_gateway()),
        ).collect_stocks(seller, today(), 0)
    async with application.database.session() as session:
        repository = TurnoverRepository(session)
        assert await repository.fbo_stocks(seller) == {DRILL_CHRT: 12, SAW_CHRT: 7}
        tracked = await repository.tracked(seller)
    assert tracked is not None and tracked.fbo_stocks_at is not None


def workbook(rows: list[tuple[str, str, str]], sheet_title: str = "ОСНОВА ") -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = sheet_title
    sheet.append(["Кабинет", "Артикул WB", "Баркод", "Артикул продавца", "Номенклатура 1С", "Активность"])
    for cabinet, barcode, nomenclature in rows:
        sheet.append([cabinet, 1, int(barcode) if barcode.isdigit() else barcode, "x", nomenclature, "активна"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def upload(client: AsyncClient, content: bytes):
    return await client.post(f"{API}/import", files={"workbook": ("list.xlsx", content)})


async def test_uploading_the_list_connects_the_cabinet_adds_rows_and_queues_collection(
    application: Application, client: AsyncClient
) -> None:
    async with application.database.session() as session:
        model = SellerModel(name="Импорт тест", catalog_sync_status="success", egress_status="verified")
        session.add(model)
        await session.commit()
        seller_id = model.id
    try:
        content = workbook([("Импорт тест", DRILL, "Дрель 1С"), ("Импорт тест", SAW, ""), ("Нет такого", DRILL, "")])

        first = (await upload(client, content)).json()["cabinets"]
        assert first[0] == {
            "name": "Импорт тест",
            "seller_name": "Импорт тест",
            "connected_now": True,
            "added": 2,
            "existing": 0,
            "unknown": 2,
        }
        assert first[1]["seller_name"] is None

        second = (await upload(client, content)).json()["cabinets"][0]
        assert (second["connected_now"], second["added"], second["existing"]) == (False, 0, 2)

        row = [r for r in (await client.get(API)).json()["rows"] if r["seller_name"] == "Импорт тест"]
        assert {item["barcode"]: item["nomenclature"] for item in row} == {DRILL: "Дрель 1С", SAW: ""}
        async with application.database.session() as session:
            assert await TurnoverRepository(session).active_refresh(seller_id) is not None
    finally:
        async with application.database.session() as session:
            await TurnoverRepository(session).purge_seller(seller_id)
            await session.execute(delete(RefreshRequestModel).where(RefreshRequestModel.seller_id == seller_id))
            await session.execute(delete(SellerModel).where(SellerModel.id == seller_id))
            await session.commit()


async def test_a_file_without_the_columns_or_not_excel_is_refused(client: AsyncClient) -> None:
    assert (await upload(client, b"not excel")).status_code == 422
    from openpyxl import Workbook

    book = Workbook()
    book.active.append(["Что-то", "Другое"])
    buffer = io.BytesIO()
    book.save(buffer)
    refused = await upload(client, buffer.getvalue())
    assert refused.status_code == 422 and "нет столбцов" in refused.json()["detail"]
