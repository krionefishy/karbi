import argparse
import asyncio
import re
import sys
from dataclasses import dataclass

from openpyxl import load_workbook

from backend.modules.wb_core.infrastructure.postgres import SellerRepository
from backend.modules.wb_turnover.infrastructure.postgres import TurnoverRepository
from backend.shared.settings import load_settings
from backend.storage.pg import Database

LIST_SHEET = "основа"
HEADERS = {"кабинет": "seller", "баркод": "barcode", "номенклатура 1с": "nomenclature"}


@dataclass(frozen=True, slots=True)
class SheetRow:
    seller: str
    barcode: str
    nomenclature: str


def _norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _barcode(value: object) -> str:
    if value is None:
        return ""
    return str(int(value)) if isinstance(value, float) else str(value).strip()


def read_rows(path: str) -> list[SheetRow]:
    """Строки листа «ОСНОВА»: кабинет, баркод и номенклатура 1С. Остальное берётся из каталога."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = next((item for item in workbook.worksheets if _norm(item.title) == LIST_SHEET), None)
    if sheet is None:
        raise SystemExit(f"В книге нет листа «ОСНОВА»: {[item.title for item in workbook.worksheets]}")
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    columns = {HEADERS[_norm(name)]: index for index, name in enumerate(header) if _norm(name) in HEADERS}
    missing = set(HEADERS.values()) - set(columns)
    if missing:
        raise SystemExit(f"На листе «ОСНОВА» нет столбцов: {sorted(missing)}")
    collected = []
    for row in rows:
        seller = str(row[columns["seller"]] or "").strip()
        barcode = _barcode(row[columns["barcode"]])
        if seller and barcode:
            collected.append(SheetRow(seller, barcode, str(row[columns["nomenclature"]] or "").strip()))
    return collected


async def run(path: str, apply: bool) -> int:
    rows = read_rows(path)
    settings = load_settings()
    database = Database()
    await database.connect(settings.database.url, pool_size=1, max_overflow=0)
    failures = 0
    try:
        async with database.session() as session:
            sellers = await SellerRepository(session).list_sellers()
            turnover = TurnoverRepository(session)
            tracked = await turnover.tracked_seller_ids()
            for name in dict.fromkeys(row.seller for row in rows):
                wanted = [row for row in rows if row.seller == name]
                seller = next((item for item in sellers if _norm(item.name) == _norm(name)), None) or next(
                    (item for item in sellers if _norm(name) in _norm(item.name) or _norm(item.name) in _norm(name)),
                    None,
                )
                if seller is None:
                    print(f"SKIP  «{name}»: селлер не найден ({len(wanted)} баркодов)")
                    failures += 1
                    continue
                if seller.id not in tracked:
                    print(f"SKIP  «{name}» → {seller.name}: не подключён к «Оборачиваемости» ({len(wanted)} баркодов)")
                    failures += 1
                    continue
                cards = await SellerRepository(session).list_barcode_cards(seller.id)
                unknown = [row.barcode for row in wanted if row.barcode not in cards]
                print(f"{name} → {seller.name}: {len(wanted)} баркодов, нет в каталоге: {len(unknown)}")
                for barcode in unknown:
                    print(f"  MISS  {barcode}")
                if not apply:
                    continue
                added = await turnover.add_board_items(seller.id, [row.barcode for row in wanted], None)
                for row in wanted:
                    await turnover.update_board_item(
                        seller.id, row.barcode, note=None, nomenclature=row.nomenclature[:500], updated_by=None
                    )
                await turnover.request_refresh(seller.id, None)
                await session.commit()
                print(f"  OK    добавлено {added}, сбор заказов поставлен в очередь")
    finally:
        await database.disconnect()
    return failures


def main() -> None:
    parser = argparse.ArgumentParser(description="Наполнить таблицу оборачиваемости из листа «ОСНОВА»")
    parser.add_argument("path")
    parser.add_argument("--apply", action="store_true", help="записать в базу, а не только показать")
    args = parser.parse_args()
    sys.exit(1 if asyncio.run(run(args.path, args.apply)) else 0)


if __name__ == "__main__":
    main()
