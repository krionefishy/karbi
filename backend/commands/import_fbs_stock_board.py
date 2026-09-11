"""Первичное наполнение таблицы остатков FBS из книги селлера (Пример.xlsx).

Лист → кабинет по имени (подстрока названия листа в имени селлера), столбцы →
склады по имени из живого списка кабинета, группы → по «плюсикам» книги
(сводный столбец слева от группы), баркоды → колонка B. Порядок групп
переставляется по ТЗ: свои склады, фулфилмент, округа.

Без --apply только показывает, что нашлось и что нет. Кабинет, у которого уже
есть группы, пропускается: команда для первого заполнения, а не для сверки.

Запуск на проде:
  docker compose cp Пример.xlsx api:/tmp/board.xlsx
  docker compose exec api python -m backend.commands.import_fbs_stock_board /tmp/board.xlsx --apply
"""

import argparse
import asyncio
import re
import sys
from dataclasses import dataclass

from openpyxl import load_workbook

from backend.modules.wb_core.infrastructure.postgres import SellerRepository
from backend.modules.wb_core.infrastructure.wb import EgressGateway
from backend.modules.wb_fbs_stocks.domain import GROUP_DISTRICT, GROUP_FULFILMENT, GROUP_OWN
from backend.modules.wb_fbs_stocks.infrastructure.postgres import FbsStocksRepository
from backend.modules.wb_fbs_stocks.infrastructure.wb import WBFbsStocksClient
from backend.shared.settings import load_settings
from backend.storage.pg import Database

SERVICE_SHEETS = {"тз", "сравнение"}
KIND_ORDER = {GROUP_OWN: 0, GROUP_FULFILMENT: 1, GROUP_DISTRICT: 2}


@dataclass
class SheetGroup:
    title: str
    kind: str
    members: list[str]


def _norm(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def _kind(title: str) -> str:
    lowered = _norm(title)
    if "наш" in lowered:
        return GROUP_OWN
    if "округ" in lowered:
        return GROUP_DISTRICT
    return GROUP_FULFILMENT


def read_sheet(sheet) -> tuple[list[SheetGroup], list[str]]:
    """Группы по outline столбцов: сводный — столбец без уровня, за ним столбцы с уровнем 1."""
    headers = {cell.column: str(cell.value).strip() for cell in sheet[1] if cell.value is not None}
    grouped = {
        column
        for dimension in sheet.column_dimensions.values()
        if dimension.outlineLevel
        for column in range(dimension.min, dimension.max + 1)
    }
    groups: list[SheetGroup] = []
    for column in sorted(headers):
        if column <= 2:
            continue
        if column in grouped:
            if groups:
                groups[-1].members.append(headers[column])
            continue
        groups.append(SheetGroup(headers[column], _kind(headers[column]), []))
    barcodes = []
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        value = row[0].value
        if value is None:
            continue
        text = str(int(value)) if isinstance(value, float) else str(value).strip()
        if text:
            barcodes.append(text)
    return groups, barcodes


async def run(path: str, apply: bool) -> int:
    settings = load_settings()
    gateway = EgressGateway(settings.egress)
    client = WBFbsStocksClient(gateway, priority="interactive")
    database = Database()
    await database.connect(settings.database.url, pool_size=1, max_overflow=0)
    failures = 0
    try:
        workbook = load_workbook(path, read_only=False)
        async with database.session() as session:
            sellers = await SellerRepository(session).list_sellers()
            stocks = FbsStocksRepository(session)
            for sheet in workbook.worksheets:
                if _norm(sheet.title) in SERVICE_SHEETS:
                    continue
                seller = next((item for item in sellers if _norm(sheet.title) in _norm(item.name)), None)
                if seller is None:
                    print(f"SKIP  лист «{sheet.title}»: селлер не найден")
                    failures += 1
                    continue
                groups, barcodes = read_sheet(sheet)
                print(f"\n{sheet.title} → {seller.name}: {len(groups)} групп, {len(barcodes)} баркодов")
                if await stocks.groups(seller.id):
                    print("  SKIP  у кабинета уже есть группы")
                    continue
                warehouses = await client.warehouses(str(seller.id))
                by_name: dict[str, list[int]] = {}
                for warehouse in warehouses:
                    if not warehouse.is_deleting:
                        by_name.setdefault(_norm(warehouse.name), []).append(warehouse.warehouse_id)
                resolved: list[tuple[SheetGroup, list[int]]] = []
                for group in sorted(groups, key=lambda item: KIND_ORDER[item.kind]):
                    ids: list[int] = []
                    for name in group.members:
                        candidates = by_name.get(_norm(name)) or []
                        if not candidates:
                            print(f"  MISS  «{name}» нет в кабинете")
                            failures += 1
                            continue
                        # Два одноимённых склада (два «FulFast МОСКВА») — берём по очереди.
                        ids.append(candidates.pop(0))
                    resolved.append((group, ids))
                    print(f"  {group.kind:<10} {group.title}: {len(ids)}/{len(group.members)} складов")
                if not apply:
                    continue
                await stocks.track(seller.id)
                await stocks.replace_warehouses(seller.id, warehouses)
                for group, ids in resolved:
                    created = await stocks.add_group(seller.id, group.title, group.kind)
                    await stocks.set_group_columns(seller.id, created.id, ids)
                added = await stocks.add_barcodes(seller.id, barcodes, None)
                await stocks.request_refresh(seller.id, None)
                await session.commit()
                print(f"  OK    записано, баркодов добавлено {added}, обновление поставлено в очередь")
    finally:
        await database.disconnect()
    return failures


def main() -> None:
    parser = argparse.ArgumentParser(description="Наполнить таблицу остатков FBS из книги селлера")
    parser.add_argument("path")
    parser.add_argument("--apply", action="store_true", help="записать в базу, а не только показать")
    args = parser.parse_args()
    failures = asyncio.run(run(args.path, args.apply))
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
