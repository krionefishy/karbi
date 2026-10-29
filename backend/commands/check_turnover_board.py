import argparse
import asyncio
import re
import sys
from collections import defaultdict
from datetime import date, timedelta

from openpyxl import load_workbook

from backend.modules.wb_core.infrastructure.postgres import SellerRepository
from backend.modules.wb_turnover.application import BoardService
from backend.modules.wb_turnover.infrastructure.postgres import TurnoverRepository
from backend.shared.settings import load_settings
from backend.storage.pg import Database

WB_STOCK_ROW = "Склад WB РФ"
STOCK_TOLERANCE = 0.1


def read_report(path: str) -> tuple[date, dict[str, tuple[int, int]]]:
    """Отчёт кабинета «по данным поставщика»: конец периода и по баркоду (заказано, остаток «Склад WB РФ»).

    Заказано складывается по всем складам баркода, остаток берётся только со
    строки «Склад WB РФ» — так её читает аналитик.
    """
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    title = str(next(rows)[0] or "")
    period = re.search(r"с (\d{2})\.(\d{2})\.(\d{4}) по (\d{2})\.(\d{2})\.(\d{4})", title)
    if period is None:
        raise SystemExit("В первой строке отчёта нет периода «с дд.мм.гггг по дд.мм.гггг»")
    end = date(int(period[6]), int(period[5]), int(period[4]))
    header = [str(value or "").strip() for value in next(rows)]
    try:
        barcode, warehouse = header.index("Баркод"), header.index("Склад")
        orders, stock = header.index("шт."), header.index("Текущий остаток, шт.")
    except ValueError as error:
        raise SystemExit(f"В отчёте нет нужного столбца: {error}") from error
    collected: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for row in rows:
        code = str(row[barcode] or "").strip()
        if not code:
            continue
        collected[code][0] += int(row[orders] or 0)
        if row[warehouse] == WB_STOCK_ROW:
            collected[code][1] += int(row[stock] or 0)
    return end, {code: (values[0], values[1]) for code, values in collected.items()}


async def run(path: str, seller_name: str) -> int:
    end, report = read_report(path)
    settings = load_settings()
    database = Database()
    await database.connect(settings.database.url, pool_size=1, max_overflow=0)
    try:
        async with database.session() as session:
            service = BoardService(
                session,
                SellerRepository(session),
                TurnoverRepository(session),
                window_days=settings.turnover.board_window_days,
                red_days=settings.turnover.board_red_days,
                yellow_days=settings.turnover.board_yellow_days,
                stale_hours=settings.turnover.board_stale_hours,
            )
            view = await service.view(today=end + timedelta(days=1))
    finally:
        await database.disconnect()
    rows = [row for row in view.rows if row.seller_name.lower() == seller_name.lower()]
    if not rows:
        print(f"В таблице нет строк кабинета «{seller_name}»")
        return 1
    failures = 0
    print(f"Период отчёта до {end}: окно таблицы {view.window_from} — {view.window_to}")
    print(f"{'баркод':<16}{'заказы WB':>10}{'заказы у нас':>14}{'остаток WB':>12}{'остаток у нас':>15}")
    for row in rows:
        if row.barcode not in report:
            print(f"{row.barcode:<16}{'нет в отчёте':>10}")
            continue
        reported_orders, reported_stock = report[row.barcode]
        ours = row.stock if row.stock is not None else -1
        bad_orders = reported_orders != row.orders
        bad_stock = abs(ours - reported_stock) > max(2, reported_stock * STOCK_TOLERANCE)
        failures += bad_orders
        print(
            f"{row.barcode:<16}{reported_orders:>10}{row.orders:>14}{reported_stock:>12}{ours:>15}"
            f"{'  ЗАКАЗЫ РАЗОШЛИСЬ' if bad_orders else ''}{'  остаток разошёлся' if bad_stock else ''}"
        )
    print("\nЗаказы сошлись по всем баркодам." if not failures else f"\nРазошлись заказы: {failures}.")
    return 1 if failures else 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Сверить таблицу оборачиваемости с отчётом кабинета WB")
    parser.add_argument("report", help="XLSX «Отчёт по данным поставщика» за 7 дней, период до вчера")
    parser.add_argument("--seller", required=True, help="имя кабинета, как в реестре")
    args = parser.parse_args()
    sys.exit(asyncio.run(run(args.report, args.seller)))


if __name__ == "__main__":
    main()
